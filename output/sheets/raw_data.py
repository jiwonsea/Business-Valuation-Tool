"""Sheet 0: Raw Data — single source of truth for hand-built Excel models.

Leftmost sheet. Dumps every *source* number the engine consumed (financials,
segment data, capital structure, WACC inputs, market consensus, peer multiples)
so the user can rebuild the downstream sheets with their own formulas instead of
reading engine-computed constants.

Derived values (EBITDA, betaL, WACC, EV, ...) are intentionally NOT written here
-- they belong in the sheets that compute them. Engine outputs are shown only as
labelled "engine reference" values for checking one's own formulas.

Named ranges (RD_*) are registered so downstream formulas can read
``=RD_Rf + RD_Bu * RD_ERP`` instead of ``='Raw Data'!$B$31``.
"""

from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from ._ctx import Ctx
from ..excel_styles import (
    BLUE_FILL,
    YELLOW_FILL,
    GREEN_FILL,
    NOTE_FONT,
    SECTION_FONT,
    TITLE_FONT,
    NUM_FMT,
    MULT_FMT,
    style_header_row,
    write_cell,
)

SHEET_NAME = "Raw Data"

# Column widths
_WIDTHS = [30, 18, 16, 16, 16, 16, 14, 14, 46]


def _add_name(ctx: Ctx, name: str, row: int, col: int) -> None:
    """Register a workbook-level named range pointing at a Raw Data cell.

    Defensive: openpyxl's DefinedName API differs across 3.0/3.1. A failure here
    must never break the export -- named ranges are a convenience, not a
    requirement.
    """
    ref = f"'{SHEET_NAME}'!${get_column_letter(col)}${row}"
    try:
        dn = DefinedName(name, attr_text=ref)
        try:
            ctx.wb.defined_names[name] = dn
        except TypeError:  # openpyxl < 3.1 -- DefinedNameList.append()
            ctx.wb.defined_names.append(dn)
    except Exception:  # pragma: no cover -- never block the workbook save
        pass


def _kv(ws, r: int, label: str, value, fmt=None, note: str = "", fill=YELLOW_FILL):
    write_cell(ws, r, 1, label)
    write_cell(ws, r, 2, value, fmt=fmt, fill=fill)
    if note:
        write_cell(ws, r, 3, note)
    return r + 1


def sheet_raw_data(ctx: Ctx):
    """Build the Raw Data sheet. Must be called FIRST so it lands leftmost."""
    ws = ctx.wb.create_sheet(SHEET_NAME, 0)
    ws.sheet_properties.tabColor = "F1C40F"
    for i, w in enumerate(_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    co = ctx.vi.company
    write_cell(ws, 1, 1, f"{co.name} — Raw Data (원천 입력값)", font=TITLE_FONT)
    write_cell(
        ws,
        2,
        1,
        "노란 셀 = 수집된 원천값(직접 수정 가능). 다른 시트의 수식은 이 시트를 참조하도록 연결하세요. "
        "이름정의(RD_*)가 등록되어 있어 =RD_Rf + RD_Bu*RD_ERP 형태로 쓸 수 있습니다.",
        font=NOTE_FONT,
    )

    r = 4

    # ── §1 Company ──
    write_cell(ws, r, 1, "§1. 기업 개요", font=SECTION_FONT)
    r += 1
    r = _kv(ws, r, "회사명", co.name)
    if getattr(co, "ticker", None):
        r = _kv(ws, r, "티커", co.ticker)
    r = _kv(ws, r, "시장", co.market)
    r = _kv(ws, r, "상장구분", co.legal_status)
    r = _kv(ws, r, "분석일", str(co.analysis_date))
    r = _kv(ws, r, "표시단위", ctx.unit, note="아래 모든 금액의 단위")
    unit_row = r
    r = _kv(
        ws,
        r,
        "단위승수 (1단위 = ?)",
        co.unit_multiplier,
        fmt=NUM_FMT,
        note=f"주당가치 = 지분가치 × {co.unit_multiplier:,} ÷ 적용 주식수",
    )
    _add_name(ctx, "RD_UnitMult", unit_row, 2)

    mc = ctx.result.market_comparison
    price = (mc.market_price if mc else None) or ctx.vi.market_price
    if price:
        price_source = mc.price_source if mc else "profile_as_of"
        price_as_of = (
            mc.price_as_of
            if mc and mc.price_as_of
            else co.analysis_date
            if price_source == "profile_as_of"
            else None
        )
        price_row = r
        r = _kv(
            ws,
            r,
            f"기준 주가 ({ctx.currency_sym})",
            price,
            fmt=NUM_FMT,
            note="상대가치·괴리율 계산 기준",
        )
        _add_name(ctx, "RD_Price", price_row, 2)
        source_label = {
            "profile_as_of": "프로필 as-of 가격",
            "profile_snapshot": "프로필 자동 스냅샷 (실시간 조회 실패 폴백)",
            "live": "실시간 조회 가격",
        }.get(price_source, "가격 출처 미지정")
        r = _kv(ws, r, "주가 출처", source_label)
        r = _kv(
            ws,
            r,
            "주가 기준일",
            str(price_as_of) if price_as_of else "미확인",
            note="분석일과 가격 기준일의 일치 여부를 확인",
        )

    # ── §2 Consolidated financials (raw only — no EBITDA) ──
    r += 1
    write_cell(ws, r, 1, "§2. 연결재무제표 (원천)", font=SECTION_FONT)
    write_cell(
        ws,
        r,
        3,
        "EBITDA·D&A합계는 여기 없음 → Financial Summary에서 직접 계산",
        font=NOTE_FONT,
    )
    r += 1
    write_cell(ws, r, 1, f"항목 ({ctx.unit})")
    for i, y in enumerate(ctx.years, 2):
        write_cell(ws, r, i, str(y))
    style_header_row(ws, r, 1 + len(ctx.years))
    hdr_row = r

    raw_rows = [
        ("매출액", "revenue", NUM_FMT),
        ("영업이익", "op", NUM_FMT),
        ("당기순이익", "net_income", NUM_FMT),
        ("총자산", "assets", NUM_FMT),
        ("총부채", "liabilities", NUM_FMT),
        ("총자본", "equity", NUM_FMT),
        ("부채비율 (%)", "de_ratio", "0.0"),
        ("감가상각비 (Dep)", "dep", NUM_FMT),
        ("무형자산상각비 (Amort)", "amort", NUM_FMT),
    ]
    for label, key, fmt in raw_rows:
        r += 1
        write_cell(ws, r, 1, label, bold=True)
        for i, y in enumerate(ctx.years, 2):
            write_cell(ws, r, i, ctx.cons[y].get(key, 0), fmt=fmt, fill=YELLOW_FILL)

    # Base-year column pointer (for named ranges + formula hints)
    by_col = 2 + ctx.years.index(ctx.by) if ctx.by in ctx.years else 2
    by_letter = get_column_letter(by_col)
    write_cell(
        ws,
        hdr_row,
        len(ctx.years) + 3,
        f"기준연도: {ctx.by} ({by_letter}열)",
        font=NOTE_FONT,
    )

    # ── §3 Segment raw data (base year) ──
    if ctx.vi.segment_data.get(ctx.by):
        r += 2
        write_cell(ws, r, 1, f"§3. 부문별 원천 ({ctx.by}년)", font=SECTION_FONT)
        write_cell(
            ws,
            r,
            4,
            "D&A는 유무형자산 비중으로 배분 → 부문 EBITDA 산출",
            font=NOTE_FONT,
        )
        r += 1
        seg_hdr = [
            "부문",
            "코드",
            "매출",
            "영업이익",
            "유무형자산",
            "평가방법",
            "적용 멀티플",
        ]
        for c, h in enumerate(seg_hdr, 1):
            write_cell(ws, r, c, h)
        style_header_row(ws, r, len(seg_hdr))
        for code in ctx.seg_codes:
            s = ctx.vi.segment_data[ctx.by].get(code, {})
            info = ctx.vi.segments.get(code, {})
            r += 1
            write_cell(ws, r, 1, ctx.seg_names.get(code, code))
            write_cell(ws, r, 2, code)
            write_cell(ws, r, 3, s.get("revenue", 0), fmt=NUM_FMT, fill=YELLOW_FILL)
            write_cell(ws, r, 4, s.get("op", 0), fmt=NUM_FMT, fill=YELLOW_FILL)
            write_cell(ws, r, 5, s.get("assets", 0), fmt=NUM_FMT, fill=YELLOW_FILL)
            write_cell(ws, r, 6, info.get("method", "ev_ebitda"))
            write_cell(
                ws,
                r,
                7,
                ctx.vi.multiples.get(code, 0.0),
                fmt=MULT_FMT,
                fill=BLUE_FILL,
                bold=True,
            )

    # ── §4 Capital structure ──
    r += 2
    write_cell(ws, r, 1, "§4. 자본구조 · 주식수", font=SECTION_FONT)
    r += 1
    nd_row = r
    r = _kv(
        ws,
        r,
        f"순차입금 ({ctx.unit})",
        ctx.vi.net_debt,
        fmt=NUM_FMT,
        note="EV → 지분가치 차감항목",
    )
    _add_name(ctx, "RD_NetDebt", nd_row, 2)
    r = _kv(ws, r, "보통주 발행주식수", co.shares_ordinary, fmt=NUM_FMT)
    if co.shares_preferred > 0:
        r = _kv(ws, r, "우선주 발행주식수", co.shares_preferred, fmt=NUM_FMT)
    if co.treasury_shares > 0:
        r = _kv(ws, r, "자사주 (보통주)", co.treasury_shares, fmt=NUM_FMT)
    r = _kv(
        ws,
        r,
        "유통보통주식수",
        co.shares_outstanding,
        fmt=NUM_FMT,
        note="= 보통주 − 자사주 (경제적 유통주식수)",
        fill=BLUE_FILL,
    )
    applied_row = r
    r = _kv(
        ws,
        r,
        "적용 주식수 (주당가치 분모)",
        ctx.vi.valuation_shares,
        fmt=NUM_FMT,
        note="최고확률 기준 시나리오 shares; 없으면 유통보통주식수",
        fill=BLUE_FILL,
    )
    _add_name(ctx, "RD_Shares", applied_row, 2)
    if ctx.vi.cps_principal > 0:
        r = _kv(ws, r, f"CPS 원금 ({ctx.unit})", ctx.vi.cps_principal, fmt=NUM_FMT)
    if ctx.vi.rcps_principal > 0:
        r = _kv(ws, r, f"RCPS 원금 ({ctx.unit})", ctx.vi.rcps_principal, fmt=NUM_FMT)

    # ── §5 WACC inputs (raw params only) ──
    r += 1
    write_cell(ws, r, 1, "§5. WACC 입력값 (원천)", font=SECTION_FONT)
    write_cell(
        ws, r, 3, "βL·Ke·WACC는 계산값 → Assumptions에서 직접 수식으로", font=NOTE_FONT
    )
    r += 1
    wp = ctx.vi.wacc_params
    wacc_inputs = [
        ("무위험이자율 Rf (%)", wp.rf, "RD_Rf", "국고채/UST 10Y"),
        ("주식위험프리미엄 ERP (%)", wp.erp, "RD_ERP", "시장 ERP 가정"),
        ("Unlevered Beta βu", wp.bu, "RD_Bu", "peer 무부채 베타 평균"),
        ("D/E Ratio (%)", wp.de, "RD_DE", "시장가치 기준"),
        ("법인세율 t (%)", wp.tax, "RD_Tax", "실효세율"),
        ("세전 타인자본비용 Kd (%)", wp.kd_pre, "RD_KdPre", "신용등급 스프레드"),
        ("자기자본 비중 E/(D+E) (%)", wp.eq_w, "RD_EqW", "WACC 가중치"),
        ("규모/비상장 프리미엄 (%)", wp.size_premium, "RD_SizePrem", "상장사는 0"),
    ]
    for label, val, name, note in wacc_inputs:
        _add_name(ctx, name, r, 2)
        r = _kv(ws, r, label, val, fmt="0.000", note=note)

    r = _kv(
        ws,
        r,
        "[엔진 참고] WACC (%)",
        round(ctx.result.wacc.wacc, 2),
        fmt="0.00",
        note="본인 수식 결과와 대조용 — 참조하지 말 것",
        fill=GREEN_FILL,
    )

    # ── §6 Market consensus ──
    ms = ctx.vi.market_signals
    if ms and ms.has_any():
        r += 1
        write_cell(ws, r, 1, "§6. 시장 컨센서스 · 매크로", font=SECTION_FONT)
        write_cell(
            ws,
            r,
            3,
            "목표주가 → 역산 멀티플 검증에 사용 (Peer Comparison 참조)",
            font=NOTE_FONT,
        )
        r += 1
        if ms.target_mean:
            _add_name(ctx, "RD_TargetMean", r, 2)
            r = _kv(
                ws,
                r,
                f"애널리스트 목표주가 평균 ({ctx.currency_sym})",
                ms.target_mean,
                fmt=NUM_FMT,
                note=f"N={ms.analyst_count or '?'}",
            )
        if ms.target_high:
            r = _kv(ws, r, "목표주가 최고", ms.target_high, fmt=NUM_FMT)
        if ms.target_low:
            r = _kv(ws, r, "목표주가 최저", ms.target_low, fmt=NUM_FMT)
        if ms.recommendation:
            r = _kv(ws, r, "투자의견 컨센서스", ms.recommendation)
        macro = [
            ("US 10Y 국채금리 (%)", ms.us_10y_yield),
            ("기대인플레이션 (%)", ms.breakeven_inflation),
            ("BAA 신용스프레드 (%)", ms.credit_spread_baa),
            ("VIX", ms.vix),
        ]
        for label, val in macro:
            if val is not None:
                r = _kv(ws, r, label, val, fmt="0.00")

    # ── §7 Peer raw table (XLOOKUP source) ──
    if ctx.vi.peers:
        r += 1
        write_cell(ws, r, 1, "§7. Peer 원천 데이터", font=SECTION_FONT)
        write_cell(
            ws,
            r,
            4,
            "통계·적용 멀티플은 Peer Comparison에서 =MEDIAN/=XLOOKUP으로 직접 산출",
            font=NOTE_FONT,
        )
        r += 1
        ph = [
            "기업명",
            "티커",
            "부문코드",
            "EV/EBITDA",
            "P/E",
            "P/BV",
            "EV/Rev",
            "Beta",
            "출처",
        ]
        for c, h in enumerate(ph, 1):
            write_cell(ws, r, c, h)
        style_header_row(ws, r, len(ph))
        peer_first = r + 1
        for p in ctx.vi.peers:
            r += 1
            write_cell(ws, r, 1, p.name)
            write_cell(ws, r, 2, p.ticker or "-")
            write_cell(ws, r, 3, p.segment_code)
            write_cell(ws, r, 4, p.ev_ebitda, fmt=MULT_FMT, fill=YELLOW_FILL)
            write_cell(ws, r, 5, p.trailing_pe if p.trailing_pe else "-", fmt=MULT_FMT)
            write_cell(ws, r, 6, p.pbv if p.pbv else "-", fmt=MULT_FMT)
            write_cell(ws, r, 7, p.ev_revenue if p.ev_revenue else "-", fmt=MULT_FMT)
            write_cell(ws, r, 8, p.beta if p.beta else "-", fmt="0.00")
            write_cell(ws, r, 9, p.source)
        if r >= peer_first:
            try:
                ctx.wb.defined_names["RD_PeerTable"] = DefinedName(
                    "RD_PeerTable",
                    attr_text=f"'{SHEET_NAME}'!$A${peer_first}:$I${r}",
                )
            except Exception:
                pass
            write_cell(
                ws,
                r + 1,
                1,
                f"이름정의: RD_PeerTable = A{peer_first}:I{r}  "
                f'예) =MEDIAN(IF(INDEX(RD_PeerTable,0,3)="seg",INDEX(RD_PeerTable,0,4)))',
                font=NOTE_FONT,
            )
            r += 1

    # ── §8 Sheet-by-sheet build order ──
    r += 2
    write_cell(ws, r, 1, "§8. 시트별 작성 순서", font=SECTION_FONT)
    r += 1
    write_cell(
        ws,
        r,
        1,
        "연결 방식: 일반 수식 + 이름정의(Named Range)를 권장. "
        "Power Query는 외부 원천(DART/Yahoo CSV) 적재·정제 단계에 쓰고, 밸류에이션 계산 자체는 셀 수식으로 두는 편이 "
        "감사 추적(F2 → 참조 셀 추적)이 쉬움. Power Query는 셀 단위 종속성이 보이지 않음.",
        font=NOTE_FONT,
    )
    r += 2
    guide_hdr = ["순서", "시트", "이 시트에서 할 일", "참조", "대표 수식"]
    for c, h in enumerate(guide_hdr, 1):
        write_cell(ws, r, c, h)
    style_header_row(ws, r, len(guide_hdr))
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 60

    for i, (sheet, todo, src, formula) in enumerate(_build_order(ctx), 1):
        r += 1
        write_cell(ws, r, 1, i)
        write_cell(ws, r, 2, sheet, bold=True)
        write_cell(ws, r, 3, todo)
        write_cell(ws, r, 4, src)
        write_cell(ws, r, 5, formula)

    ws.freeze_panes = "A4"


def _build_order(ctx: Ctx) -> list[tuple[str, str, str, str]]:
    """(sheet, what to do, source, representative formula) — method-aware."""
    rows: list[tuple[str, str, str, str]] = [
        (
            "Raw Data",
            "원천값 확인/수정 (노란 셀만)",
            "DART · Yahoo",
            "입력 전용 — 여기서는 계산하지 않음",
        ),
        (
            "Financial Summary",
            "EBITDA · D&A 산출",
            "Raw Data §2",
            "EBITDA = 영업이익 + 감가상각비 + 무형자산상각비",
        ),
        (
            "Peer Comparison",
            "peer 멀티플 통계 → 적용 멀티플 확정",
            "Raw Data §7",
            "=MEDIAN(범위), 적용 = 중앙값 × (1 ± 프리미엄/할인)",
        ),
        (
            "Assumptions",
            "βL · Ke · WACC 계산, 적용 멀티플 고정",
            "Raw Data §5",
            "βL = RD_Bu*(1+(1-RD_Tax/100)*RD_DE/100) ; "
            "Ke = RD_Rf + βL*RD_ERP ; WACC = Ke*RD_EqW/100 + RD_KdPre*(1-RD_Tax/100)*(1-RD_EqW/100)",
        ),
    ]

    if ctx.method == "sotp":
        rows.append(
            (
                "SOTP Valuation",
                "부문 EBITDA × 멀티플 → EV → 지분가치",
                "Financial Summary + Peer",
                "부문EV = 부문EBITDA × 멀티플 ; EV = SUM(부문EV) ; "
                "지분가치 = EV − RD_NetDebt ; 주당 = 지분가치*RD_UnitMult/RD_Shares",
            )
        )
    elif ctx.method == "dcf_primary":
        rows.append(
            (
                "DCF Valuation",
                "FCFF 예측 → 할인 → TV 합산",
                "Financial Summary + Assumptions",
                "FCFF = EBIT*(1-t) + D&A − Capex − ΔNWC ; "
                "PV = FCFF/(1+WACC)^n ; TV = FCFF_n*(1+g)/(WACC−g)",
            )
        )
    else:
        rows.append(
            (
                f"{ctx.method.upper()} Valuation",
                "주 방법론 계산",
                "Financial Summary + Assumptions",
                "해당 시트의 비고 열 수식 참조",
            )
        )

    rows += [
        (
            "Scenario Analysis",
            "시나리오별 드라이버 반영 → 확률가중",
            "Valuation 시트",
            "=SUMPRODUCT(확률범위, 시나리오가치범위) / 100",
        ),
        (
            "Sensitivity",
            "2변수 표 (예: 멀티플 × WACC)",
            "Valuation 시트",
            "[데이터] → [가상 분석] → [데이터 표] (행/열 입력 셀 지정)",
        ),
        (
            "Relative Valuation",
            "PER/PBR/EV·EBITDA 진단",
            "Raw Data §1·§2",
            "PER = RD_Price / EPS ; EPS = 당기순이익*RD_UnitMult/RD_Shares",
        ),
        (
            "Dashboard",
            "결론 요약 (계산 금지, 참조만)",
            "전 시트",
            "='SOTP Valuation'!B20 형태의 단순 참조",
        ),
    ]
    return rows
