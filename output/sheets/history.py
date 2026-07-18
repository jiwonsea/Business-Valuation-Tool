"""Sheet: persisted valuation history (DB read-only; no engine recalculation)."""

from __future__ import annotations

from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter

from ._ctx import Ctx
from ..excel_styles import TITLE_FONT, style_header_row, write_cell


def sheet_valuation_history(ctx: Ctx) -> None:
    ws = ctx.wb.create_sheet("Valuation History")
    ws.sheet_properties.tabColor = "5B9BD5"
    for col, width in enumerate([14, 18, 18, 14, 12, 12, 18, 18], 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    write_cell(ws, 1, 1, "저장 시점 기준 밸류에이션 이력", font=TITLE_FONT)
    write_cell(
        ws,
        2,
        1,
        "DB에 저장된 당시 결과만 표시합니다. 현재 엔진으로 과거 값을 재계산하지 않습니다.",
    )

    try:
        from db.repository import list_valuation_history

        records = list_valuation_history(
            ticker=ctx.vi.company.ticker,
            market=ctx.vi.company.market,
            company_name=ctx.vi.company.name,
        )
    except Exception:
        records = []

    if not records:
        write_cell(ws, 4, 1, "이력 부족(N=0)", bold=True)
        return

    headers = [
        "분석일",
        "내재가치",
        "시장가격",
        "괴리율(%)",
        "WACC(%)",
        "등급",
        "방법론",
        "버킷",
    ]
    for col, header in enumerate(headers, 1):
        write_cell(ws, 4, col, header)
    style_header_row(ws, 4, len(headers))

    for row_no, record in enumerate(records, 5):
        values = [
            record.analysis_date,
            record.weighted_value,
            record.market_price,
            record.gap_pct,
            record.wacc_pct,
            record.quality_grade,
            record.primary_method,
            record.valuation_bucket,
        ]
        for col, value in enumerate(values, 1):
            fmt = "#,##0" if col in (2, 3) else ("0.00" if col in (4, 5) else None)
            write_cell(ws, row_no, col, value, fmt=fmt)

    chart = LineChart()
    chart.title = "내재가치 vs 시장가격 (저장 이력)"
    chart.y_axis.title = "주당가치"
    chart.x_axis.title = "분석일"
    data = Reference(ws, min_col=2, max_col=3, min_row=4, max_row=4 + len(records))
    categories = Reference(ws, min_col=1, min_row=5, max_row=4 + len(records))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.height = 8
    chart.width = 15
    ws.add_chart(chart, "J4")
