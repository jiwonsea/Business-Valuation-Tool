"""Regression tests for the valuation-engine/reporting boundary."""

from output.excel_builder import export
from valuation_runner import load_profile, run_valuation


def test_ddm_excel_reuses_engine_sensitivity(monkeypatch, tmp_path):
    vi = load_profile("profiles/kb_financial.yaml")
    result = run_valuation(vi)

    def fail_if_recomputed(*args, **kwargs):
        raise AssertionError("Excel renderer must not call calc_ddm")

    monkeypatch.setattr("engine.ddm.calc_ddm", fail_if_recomputed)
    output = export(vi, result, output_dir=str(tmp_path))

    from openpyxl import load_workbook

    workbook = load_workbook(output, data_only=True)
    ddm_text = "\n".join(
        str(cell)
        for row in workbook["DDM Valuation"].iter_rows(values_only=True)
        for cell in row
        if cell is not None
    )
    sensitivity_text = "\n".join(
        str(cell)
        for row in workbook["Sensitivity"].iter_rows(values_only=True)
        for cell in row
        if cell is not None
    )

    assert "DDM 민감도" not in ddm_text
    assert result.sensitivity_primary_label in sensitivity_text
