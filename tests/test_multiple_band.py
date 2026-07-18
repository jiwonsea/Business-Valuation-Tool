"""Phase 2 P1/P2 — historical LTM multiple band contract tests.

Acceptance criteria: HANDOFF_CODEX_phase2_impl_scope_2026-07-18 §4 as amended
by §7.1 (no relaxation). Everything runs OFFLINE: DART data comes from the
pilot v2 snapshot (0 network calls), prices from injected fake providers.
"""

from __future__ import annotations

import io
import json
import re
from contextlib import redirect_stdout
from datetime import date, timedelta
from pathlib import Path

import pytest

from engine.multiple_band import (
    MIN_OBS_FOR_BAND,
    band_verdict,
    build_band,
    percentile_rank,
)
from pipeline.timeseries import (
    NOTE_NO_RCEPT_NO,
    build_band_reports,
    build_observations,
    extract_company_values,
    fetch_price_data,
    resolve_pilot_company,
)
from schemas.point_in_time import (
    EXCLUSION_NOTE_DENOMINATOR_NON_POSITIVE,
    PRICE_SEARCH_WINDOW_DAYS,
    ExcludedYear,
    HistoricalBand,
    MultipleObservation,
    PriceExclusionReason,
    require_allowed_multiple_label,
)

REPO = Path(__file__).resolve().parent.parent
SNAPSHOT = REPO / "research" / "pilot_v2" / "raw_payloads.json"

# Pilot v2 measured facts (pinned — see impl-scope handoff §4 #6):
LG_FY2020_REVENUE_ORIGINAL = 63_262_046  # rcept 20210316000661 (2021-03-16)
LG_FY2020_REVENUE_RESTATED = 58_057_908  # rcept 20220316000886 comparative


def _snapshot() -> dict:
    with open(SNAPSHOT, encoding="utf-8") as f:
        return json.load(f)


def flat_provider(price: float = 60_000.0):
    """Every weekday has the same raw close; no splits (known-empty)."""

    def provider(ticker, start, end):
        closes = {}
        d = start
        while d <= end:
            if d.weekday() < 5:
                closes[d] = price
            d += timedelta(days=1)
        return closes, []

    return provider


def pdata(payload: dict, provider=None):
    """Company-level price_data exactly as build_band_reports assembles it."""
    _, filing_at = extract_company_values(payload)
    return fetch_price_data(provider or flat_provider(), "X.KS", filing_at)


def _mk_obs(fy: int, multiple: float, label: str = "P/B") -> MultipleObservation:
    t = date(fy + 1, 3, 30)
    return MultipleObservation(
        label=label,
        company="X",
        fiscal_year=fy,
        t=t,
        price_date=t,
        price_close_raw_krw=100.0,
        shares_outstanding=1_000,
        market_cap_mkrw=multiple * 1_000.0,
        denominator_mkrw=1_000,
        multiple=multiple,
        rcept_no=f"{fy + 1}0330000001",
    )


# ── Pure band math (engine/multiple_band.py) ──


class TestBandMath:
    def test_quantiles_inclusive_exact(self):
        obs = [_mk_obs(2016 + i, v) for i, v in enumerate([0.8, 1.0, 1.1, 1.2, 1.5, 2.0])]
        b = build_band("P/B", "X", obs)
        assert b.n_obs == 6
        assert b.band_min == 0.8 and b.band_max == 2.0
        assert b.p25 == pytest.approx(1.025)
        assert b.median == pytest.approx(1.15)
        assert b.p75 == pytest.approx(1.425)

    def test_single_observation_collapses(self):
        b = build_band("P/B", "X", [_mk_obs(2020, 1.3)])
        assert (b.band_min, b.p25, b.median, b.p75, b.band_max) == (1.3,) * 5

    def test_empty_band_has_no_stats(self):
        b = build_band("P/B", "X", [])
        assert b.n_obs == 0
        assert b.band_min is None and b.median is None and b.band_max is None

    def test_label_and_company_mixing_rejected(self):
        with pytest.raises(ValueError, match="라벨"):
            build_band("P/S", "X", [_mk_obs(2020, 1.0, label="P/B")])
        with pytest.raises(ValueError, match="회사"):
            build_band("P/B", "Y", [_mk_obs(2020, 1.0)])

    def test_percentile_rank_mid_rank_ties(self):
        obs = [_mk_obs(2016 + i, v) for i, v in enumerate([1.0, 1.0, 2.0, 3.0])]
        b = build_band("P/B", "X", obs)
        assert percentile_rank(b, 1.0) == pytest.approx(0.25)  # (0 + 0.5*2)/4
        assert percentile_rank(b, 10.0) == 1.0
        assert percentile_rank(build_band("P/B", "X", []), 1.0) is None

    def test_verdict_small_n_is_insufficient_history(self):
        obs = [_mk_obs(2016 + i, 1.0 + i / 10) for i in range(MIN_OBS_FOR_BAND - 1)]
        assert band_verdict(build_band("P/B", "X", obs), 1.0).startswith("이력 부족")

    def test_verdict_positions(self):
        obs = [_mk_obs(2016 + i, v) for i, v in enumerate([0.8, 1.0, 1.1, 1.2, 1.5, 2.0])]
        b = build_band("P/B", "X", obs)
        assert band_verdict(b, 0.5) == "역사적 밴드 하단 이탈"
        assert band_verdict(b, 2.5) == "역사적 밴드 상단 이탈"
        assert band_verdict(b, 0.9) == "역사적 하위 25% 구간"
        assert band_verdict(b, 1.9) == "역사적 상위 25% 구간"
        assert band_verdict(b, 1.15) == "역사적 중간 구간"
        assert band_verdict(b, None) == "현재 배수 없음"


# ── Model guards (schemas/point_in_time.py, §7.1) ──


class TestModelGuards:
    def test_forward_labels_forbidden(self):
        for bad in ("12M Forward P/E", "Fwd P/B", "P/E", "EV/EBITDA"):
            with pytest.raises(ValueError):
                _mk_obs(2020, 1.0, label=bad)
            with pytest.raises(ValueError):
                HistoricalBand(label=bad, company="X", n_obs=0)

    def test_price_after_t_is_look_ahead(self):
        ok = _mk_obs(2020, 1.0)
        with pytest.raises(ValueError, match="look-ahead"):
            MultipleObservation(**{**ok.model_dump(), "price_date": ok.t + timedelta(days=1)})

    def test_price_search_window_is_seven_calendar_days(self):
        ok = _mk_obs(2020, 1.0)
        edge = ok.t - timedelta(days=PRICE_SEARCH_WINDOW_DAYS)
        MultipleObservation(**{**ok.model_dump(), "price_date": edge})  # boundary OK
        with pytest.raises(ValueError, match="calendar days"):
            MultipleObservation(
                **{**ok.model_dump(), "price_date": edge - timedelta(days=1)}
            )

    def test_exclusion_requires_mechanical_reason(self):
        with pytest.raises(ValueError, match="기계적 사유"):
            ExcludedYear(fiscal_year=2020)
        ExcludedYear(
            fiscal_year=2020, price_reason=PriceExclusionReason.NO_PRICE_WITHIN_WINDOW
        )
        ExcludedYear(fiscal_year=2020, missing_accounts=("equity",))
        ExcludedYear(fiscal_year=2020, note=EXCLUSION_NOTE_DENOMINATOR_NON_POSITIVE)

    def test_band_stats_cannot_be_fabricated(self):
        with pytest.raises(ValueError, match="날조"):
            HistoricalBand(label="P/B", company="X", n_obs=0, median=1.0)


# ── Point-in-time consumption on the real pilot snapshot (offline) ──


class TestSnapshotPointInTime:
    def test_lg_fy2020_ps_uses_original_never_restated(self):
        """§4 #2 — the FY2020 P/S denominator is the first-disclosure value."""
        lg = _snapshot()["companies"]["LG Electronics"]
        obs, _ = build_observations("P/S", "LG Electronics", lg, pdata(lg))
        fy2020 = next(o for o in obs if o.fiscal_year == 2020)
        assert fy2020.denominator_mkrw == LG_FY2020_REVENUE_ORIGINAL
        assert fy2020.rcept_no == "20210316000661"
        assert fy2020.t == date(2021, 3, 16)
        # The restated comparative exists in the pool but never reaches the band.
        pool, _ = extract_company_values(lg)
        restated = [
            v.value_mkrw
            for v in pool
            if v.account == "revenue" and v.fiscal_year == 2020 and v.basis != "original"
        ]
        assert restated == [LG_FY2020_REVENUE_RESTATED]
        assert LG_FY2020_REVENUE_RESTATED not in [o.denominator_mkrw for o in obs]

    def test_lg_fy2016_excluded_shares_basis_missing(self):
        """Pilot v2 measured fact: LG has exactly one source-missing year."""
        lg = _snapshot()["companies"]["LG Electronics"]
        for label in ("P/B", "P/S"):
            obs, exc = build_observations(label, "LG Electronics", lg, pdata(lg))
            assert len(obs) == 9
            assert [(e.fiscal_year, e.price_reason) for e in exc] == [
                (2016, PriceExclusionReason.SHARES_BASIS_MISSING)
            ]

    def test_pilot_three_companies_full_parity(self):
        """§4 #6 — observation counts match pilot v2 mapping/missing rates."""
        expected = {"삼성전자": 10, "SK하이닉스": 10, "LG전자": 9}
        for name, n in expected.items():
            bands = build_band_reports(company_name=name, price_provider=flat_provider())
            assert bands is not None and len(bands) == 2
            for band in bands:
                assert band.n_obs == n
                assert band.n_obs + len(band.excluded) == 10

    def test_out_of_scope_company_skips_without_network(self):
        assert resolve_pilot_company(name="NAVER", ticker="035420") is None
        assert build_band_reports(company_name="NAVER", ticker="035420") is None


# ── Mechanical exclusion reasons (§7-4: no subjective judgment) ──


def _hynix_payload() -> dict:
    return _snapshot()["companies"]["SK hynix"]


class TestMechanicalExclusions:
    def test_no_price_within_window(self):
        def no_price_provider(ticker, start, end):
            return {}, []

        obs, exc = build_observations("P/B", "SK hynix", _hynix_payload(), pdata(_hynix_payload(), no_price_provider))
        assert obs == []
        assert {e.price_reason for e in exc} == {
            PriceExclusionReason.NO_PRICE_WITHIN_WINDOW
        }

    def test_post_t_prices_never_used(self):
        """Provider offers prices ONLY after t — must not be consumed."""

        def post_t_only_provider(ticker, start, end):
            closes = {end + timedelta(days=i): 60_000.0 for i in range(1, 4)}
            return closes, []

        obs, exc = build_observations("P/B", "SK hynix", _hynix_payload(), pdata(_hynix_payload(), post_t_only_provider))
        assert obs == []
        assert {e.price_reason for e in exc} == {
            PriceExclusionReason.NO_PRICE_WITHIN_WINDOW
        }

    def test_unknown_action_basis_is_mismatch(self):
        def splits_unknown_provider(ticker, start, end):
            closes, _ = flat_provider()(ticker, start, end)
            return closes, None

        obs, exc = build_observations("P/B", "SK hynix", _hynix_payload(), pdata(_hynix_payload(), splits_unknown_provider))
        assert obs == []
        assert {e.price_reason for e in exc} == {
            PriceExclusionReason.SHARES_BASIS_MISMATCH
        }

    def test_split_between_price_date_and_t_detected(self):
        payload = _hynix_payload()
        _, filing_at = extract_company_values(payload)
        t_2019 = filing_at[2019][1]

        def split_provider(ticker, start, end):
            closes = {t_2019 - timedelta(days=2): 60_000.0}
            d = start
            while d <= end:  # other years keep normal weekday closes
                if d.weekday() < 5 and abs((d - t_2019).days) > 10:
                    closes.setdefault(d, 60_000.0)
                d += timedelta(days=1)
            return closes, [t_2019 - timedelta(days=1)]

        obs, exc = build_observations("P/B", "SK hynix", payload, pdata(payload, split_provider))
        assert (2019, PriceExclusionReason.SPLIT_ADJUSTMENT_DETECTED) in [
            (e.fiscal_year, e.price_reason) for e in exc
        ]
        assert 2019 not in [o.fiscal_year for o in obs]

    def test_denominator_non_positive_excluded(self):
        items = [
            {
                "rcept_no": "20210316000661",
                "sj_div": "BS",
                "account_nm": "자본총계",
                "thstrm_amount": "-1,000,000,000",
                "frmtrm_amount": "",
            }
        ]
        payload = {
            "financial": {"2020": {"items": items}},
            "stock": {
                "2020": {
                    "shares": {"shares_ordinary": 1000, "treasury_ordinary": 0}
                }
            },
        }
        obs, exc = build_observations("P/B", "X", payload, pdata(payload))
        assert obs == []
        assert [(e.fiscal_year, e.note) for e in exc] == [
            (2020, EXCLUSION_NOTE_DENOMINATOR_NON_POSITIVE)
        ]

    def test_payload_without_rcept_no_excluded(self):
        payload = {
            "financial": {
                "2020": {
                    "items": [
                        {
                            "sj_div": "BS",
                            "account_nm": "자본총계",
                            "thstrm_amount": "1,000,000,000",
                        }
                    ]
                }
            },
            "stock": {},
        }
        obs, exc = build_observations("P/B", "X", payload, pdata(payload))
        assert obs == []
        assert [(e.fiscal_year, e.note) for e in exc] == [(2020, NOTE_NO_RCEPT_NO)]


# ── Retroactive split adjustment (§11 — Codex 실측 결함 재현) ──


class TestRetroactiveSplitAdjustment:
    def test_samsung_real_yahoo_shape_excluded_fail_closed(self):
        """Codex 실측 재현: 삼성전자 2018-05 50:1 분할이 Yahoo 과거 close에
        소급 반영(auto_adjust=False여도) → FY2016-17은 가격(분할 후 기준) vs
        주식수(분할 전 기준)가 ~50x 어긋남. fail-closed 제외돼야 한다."""
        samsung = _snapshot()["companies"]["Samsung Electronics"]
        base = flat_provider()

        def yahoo_shape(ticker, start, end):
            closes, _ = base(ticker, start, end)
            return closes, [date(2018, 5, 4), date(2018, 5, 16)]

        obs, exc = build_observations(
            "P/B", "Samsung Electronics", samsung, pdata(samsung, yahoo_shape)
        )
        excluded = {(e.fiscal_year, e.price_reason) for e in exc}
        assert (2016, PriceExclusionReason.SPLIT_ADJUSTMENT_DETECTED) in excluded
        assert (2017, PriceExclusionReason.SPLIT_ADJUSTMENT_DETECTED) in excluded
        kept = sorted(o.fiscal_year for o in obs)
        # Split (2018-05) <= FYE2018 -> FY2018+ shares are post-split: consistent.
        assert kept == list(range(2018, 2026))

    def test_split_at_or_before_fye_keeps_the_year(self):
        """분할이 해당 회계연도말 이전이면 주식수도 분할 후 기준 — 관측 유지."""
        samsung = _snapshot()["companies"]["Samsung Electronics"]
        base = flat_provider()

        def split_before_fye(ticker, start, end):
            closes, _ = base(ticker, start, end)
            return closes, [date(2016, 1, 4)]  # FY2016 FYE(2016-12-31) 이전

        obs, exc = build_observations(
            "P/B", "Samsung Electronics", samsung, pdata(samsung, split_before_fye)
        )
        assert len(obs) == 10 and exc == []

    def test_split_after_last_filing_still_detected(self):
        """마지막 접수일 이후의 분할도 과거 close 전체를 소급 조정 — span이
        오늘까지 확장되어 탐지·전 연도 제외돼야 한다."""
        payload = _hynix_payload()
        base = flat_provider()
        recent_split = date.today() - timedelta(days=1)

        def late_split(ticker, start, end):
            assert end >= recent_split  # span-through-today (§11)
            closes, _ = base(ticker, start, end)
            return closes, [recent_split]

        obs, exc = build_observations(
            "P/B", "SK hynix", payload, pdata(payload, late_split)
        )
        assert obs == []
        assert {e.price_reason for e in exc} == {
            PriceExclusionReason.SPLIT_ADJUSTMENT_DETECTED
        }

    def test_fetch_span_extends_through_today(self):
        calls = []
        base = flat_provider()

        def recording(ticker, start, end):
            calls.append((start, end))
            return base(ticker, start, end)

        pdata(_hynix_payload(), recording)
        assert len(calls) == 1
        assert calls[0][1] >= date.today()


# ── Provider contract (§10 — Codex 보류 조건 1·2) ──


class TestProviderContract:
    def test_provider_called_exactly_once_per_company(self):
        """§10-1: one fetch per company, shared by both labels."""
        calls = []
        base = flat_provider()

        def counting(ticker, start, end):
            calls.append((ticker, start, end))
            return base(ticker, start, end)

        bands = build_band_reports(company_name="SK하이닉스", price_provider=counting)
        assert len(bands) == 2 and all(b.n_obs == 10 for b in bands)
        assert len(calls) == 1

    def test_source_failure_distinct_from_known_empty_splits(self):
        """§10-2: ({}, None)=원천 실패 → mismatch, ({}, [])=기준 확립·가격만
        없음 → no_price. 두 사유가 뒤섞이지 않는다."""
        payload = _hynix_payload()
        obs, exc = build_observations("P/B", "SK hynix", payload, ({}, None))
        assert obs == []
        assert {e.price_reason for e in exc} == {
            PriceExclusionReason.SHARES_BASIS_MISMATCH
        }
        obs2, exc2 = build_observations("P/B", "SK hynix", payload, ({}, []))
        assert obs2 == []
        assert {e.price_reason for e in exc2} == {
            PriceExclusionReason.NO_PRICE_WITHIN_WINDOW
        }

    def test_parse_history_frame_contract(self):
        """§10-2: 성공한 history(actions=True) 1회 응답에서 close+split 동시
        추출. 빈 응답=원천 실패(None), 0뿐인 split 열=known-empty([])."""
        pd = pytest.importorskip("pandas")
        from pipeline.timeseries import _parse_history_frame

        start, end = date(2020, 3, 1), date(2020, 3, 31)
        assert _parse_history_frame(None, start, end) == ({}, None)
        assert _parse_history_frame(pd.DataFrame(), start, end) == ({}, None)

        idx = pd.DatetimeIndex([date(2020, 3, 2), date(2020, 3, 3)])
        no_split_col = pd.DataFrame({"Close": [100.0, 101.0]}, index=idx)
        closes, splits = _parse_history_frame(no_split_col, start, end)
        assert closes == {date(2020, 3, 2): 100.0, date(2020, 3, 3): 101.0}
        assert splits is None  # 열 부재 = 기준 수립 불가

        all_zero = pd.DataFrame(
            {"Close": [100.0, 101.0], "Stock Splits": [0.0, 0.0]}, index=idx
        )
        assert _parse_history_frame(all_zero, start, end)[1] == []

        with_split = pd.DataFrame(
            {"Close": [100.0, 101.0], "Stock Splits": [0.0, 2.0]}, index=idx
        )
        assert _parse_history_frame(with_split, start, end)[1] == [date(2020, 3, 3)]


# ── Label contract (§7.1 #4) ──


class TestLabelContract:
    def test_every_generated_label_passes_the_guard(self):
        for name in ("삼성전자", "SK하이닉스", "LG전자"):
            for band in build_band_reports(company_name=name, price_provider=flat_provider()):
                require_allowed_multiple_label(band.label)
                for o in band.observations:
                    require_allowed_multiple_label(o.label)

    def test_out_of_scope_label_fails_closed(self):
        with pytest.raises((KeyError, ValueError)):
            build_observations(
                "P/E", "SK hynix", _hynix_payload(), pdata(_hynix_payload())
            )

    def test_console_output_has_no_forward_labels(self):
        from output.band_report import print_band_reports

        bands = build_band_reports(company_name="SK하이닉스", price_provider=flat_provider())
        buf = io.StringIO()
        with redirect_stdout(buf):
            print_band_reports(bands, {"P/B": 1.2})
        text = buf.getvalue()
        assert not re.search(r"Forward|Fwd", text, re.IGNORECASE)
        assert "참고용" in text and "밸류에이션 입력 아님" in text


# ── Reporting-only + purity + contract-only consumption (§4 #1/#7) ──


class TestConsumptionBoundaries:
    def test_valuation_paths_never_import_multiple_band(self):
        """Engine wiring of P1/P2 was REJECTED — pin it with a source scan."""
        forbidden_consumers = [REPO / "valuation_runner.py", REPO / "orchestrator.py"]
        forbidden_consumers += [
            p for p in (REPO / "engine").glob("*.py") if p.name != "multiple_band.py"
        ]
        for p in forbidden_consumers:
            assert "multiple_band" not in p.read_text(encoding="utf-8", errors="ignore"), (
                f"{p} must not consume the reporting-only band module"
            )

    def test_band_engine_module_is_pure(self):
        src = (REPO / "engine" / "multiple_band.py").read_text(encoding="utf-8")
        for forbidden in ("httpx", "requests", "yfinance", "urllib", "socket", "open("):
            assert forbidden not in src

    def test_timeseries_consumes_contract_api_only(self):
        """§4 #1 — no legacy dict-path consumption in new CODE (docstrings may
        cite the prohibition itself; identifiers may not appear in code)."""
        import ast as ast_mod

        legacy = {"parse_financial_statements", "ACCOUNT_MAP", "CAPEX_MAP"}
        for rel in ("pipeline/timeseries.py", "output/band_report.py", "output/sheets/band.py"):
            tree = ast_mod.parse((REPO / rel).read_text(encoding="utf-8"))
            idents: set[str] = set()
            for node in ast_mod.walk(tree):
                if isinstance(node, ast_mod.Name):
                    idents.add(node.id)
                elif isinstance(node, ast_mod.Attribute):
                    idents.add(node.attr)
                elif isinstance(node, ast_mod.ImportFrom):
                    idents.update(a.name for a in node.names)
                elif isinstance(node, ast_mod.Import):
                    idents.update(a.name for a in node.names)
            assert not (idents & legacy), f"{rel}: legacy path 소비 {idents & legacy}"

    def test_band_source_files_have_no_forward_literals(self):
        for rel in (
            "engine/multiple_band.py",
            "pipeline/timeseries.py",
            "output/band_report.py",
            "output/sheets/band.py",
        ):
            src = (REPO / rel).read_text(encoding="utf-8")
            assert not re.search(r'"12M Forward"|Fwd P/', src), rel


# ── Excel sheet is strictly opt-in (§7.1 #8) ──


class TestExcelOptIn:
    def _ctx(self):
        from openpyxl import Workbook
        from output.sheets._ctx import Ctx

        wb = Workbook()
        return Ctx(
            vi=None, result=None, wb=wb, method="dcf", by=2025,
            seg_names={}, seg_codes=[], cons={}, years=[], unit="백만원",
            currency_sym="원",
        )

    def test_sheet_only_created_when_bands_passed(self):
        from output.sheets.band import sheet_historical_band

        bands = build_band_reports(company_name="SK하이닉스", price_provider=flat_provider())
        ctx = self._ctx()
        assert "Historical Band" not in ctx.wb.sheetnames
        sheet_historical_band(ctx, bands, {"P/B": 1.2})
        assert "Historical Band" in ctx.wb.sheetnames
        ws = ctx.wb["Historical Band"]
        assert ws["A4"].value == "P/B (LTM)"
        assert ws["B4"].value == 10  # n_obs pinned to pilot parity

    def test_export_signature_defaults_keep_workbook_unchanged(self):
        """band_reports defaults to None — the band sheet cannot appear in a
        default (--excel only) workbook."""
        import inspect
        from output.excel_builder import export

        sig = inspect.signature(export)
        assert sig.parameters["band_reports"].default is None
        assert sig.parameters["band_current"].default is None
