"""Fail-closed verification gates for the MU report pipeline."""

from __future__ import annotations

import ast
import json
import re
from collections import Counter
from collections.abc import Iterable
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook
from pypdf import PdfReader

from .charts import CHART_NAMES
from .facts import FORBIDDEN_FIELDS, UNAVAILABLE_STATUSES, Manifest
from .inputs import ConflictConfirmation, EvidenceReader, PACKAGE_DIR, load_pins
from .rle import inventory_days
from .theme import COLORS, contrast_ratio
from .valuation import trailing_pb


class GateError(AssertionError):
    pass


def require(condition: bool, message: str) -> None:
    if not condition:
        raise GateError(message)


def gate_g1_inputs(reader: EvidenceReader) -> None:
    for item in load_pins()["E2-A"]:
        reader.read_bytes(item["path"])


def gate_g2_freeze_display(expected: dict[str, str], actual: dict[str, str]) -> None:
    require(expected == actual, "PREREG_A display strings differ from Freeze A")


def _close(name: str, left: float, right: float, terms: int, differences: list[dict]) -> None:
    delta = left - right
    tolerance = 0.5 * terms
    require(abs(delta) <= tolerance, f"{name}: {left} != {right}; delta {delta} > tolerance {tolerance}")
    if delta:
        differences.append({"identity": name, "difference": delta, "tolerance": tolerance})


def gate_g3_financial_identities(statement: dict[str, float]) -> list[dict]:
    differences: list[dict] = []
    _close("revenue-cogs=gross_profit", statement["revenue"] - statement["cogs"], statement["gross_profit"], 3, differences)
    opex = statement["r_and_d"] + statement["sg_and_a"] + statement["restructuring"] + statement["other_operating"]
    _close("gross_profit-opex=operating_income", statement["gross_profit"] - opex, statement["operating_income"], 6, differences)
    below_op = statement["interest_income"] + statement["interest_expense"] + statement["other_nonoperating"]
    _close("operating_income+below_op=pretax", statement["operating_income"] + below_op, statement["pretax_income"], 5, differences)
    _close("pretax+tax+equity=net", statement["pretax_income"] + statement["tax"] + statement["equity_method"], statement["net_income"], 4, differences)
    eps_tolerance = 0.005 + 0.5 * (abs(statement["diluted_eps"]) + 1.0) / statement["diluted_shares"]
    delta = statement["net_income"] / statement["diluted_shares"] - statement["diluted_eps"]
    require(abs(delta) <= eps_tolerance, f"EPS identity delta {delta} > {eps_tolerance}")
    if delta:
        differences.append({"identity": "net_income/shares=eps", "difference": delta, "tolerance": eps_tolerance})
    return differences


def gate_g3b_balance_sheet(statement: dict[str, float]) -> list[dict]:
    differences: list[dict] = []
    _close("assets=liabilities+equity", statement["total_assets"], statement["total_liabilities"] + statement["total_equity"], 3, differences)
    return differences


def gate_g3c_cash_flow(values: dict[str, float]) -> list[dict]:
    differences: list[dict] = []
    _close("cash flow reconciliation", values["operating"] + values["investing"] + values["financing"] + values["fx"], values["cash_change"], 5, differences)
    return differences


def gate_g3d_company_fcf(values: dict[str, float]) -> None:
    _close("company net capex", values["ppe_expenditures"] - values["ppe_disposal_proceeds"] - values["government_incentives"], values["net_capex"], 4, [])
    _close("company adjusted FCF", values["operating_cash_flow"] - values["net_capex"], values["adjusted_free_cash_flow"], 3, [])


def gate_g7_templates_and_io(paths: Iterable[Path]) -> None:
    for path in paths:
        text = path.read_text(encoding="utf-8")
        if path.suffix == ".py":
            tree = ast.parse(text)
            for node in ast.walk(tree):
                if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id == "open":
                    raise GateError(f"direct open() call: {path}:{node.lineno}")
        if path.suffix in {".yaml", ".yml"} and "i18n" in path.parts:
            gate_g14_template_numbers(text)
            dates = set(re.findall(r"20\d{2}-\d{2}-\d{2}", text))
            require(
                dates <= {"2026-09-25", "2026-09-30", "2026-10-01", "2026-10-04"},
                f"i18n contains date-dependent fact ID or literal: {sorted(dates - {'2026-09-25', '2026-09-30', '2026-10-01', '2026-10-04'})}",
            )


PLACEHOLDER_RE = re.compile(r"\{\{fact:[^}]+}}")


def gate_g9_cutoff(records: list[dict[str, Any]]) -> None:
    for record in records:
        if record.get("section") == "PREREG_A":
            require(record.get("information_class") != "POST_PRINT", "post-print fact in PREREG_A section")
        if record.get("section") == "RLE":
            require(record.get("basis") != "PREREG_A", "RLE consumes PREREG_A")


def gate_g12_manifest(manifest: Manifest) -> None:
    manifest.validate()
    keys = set(manifest.to_dict())
    require(not keys & FORBIDDEN_FIELDS, "forbidden valuation field")


def gate_g12b_disclaimers(texts: dict[str, str]) -> None:
    required = {
        "ko": ("본 문서는 투자 자문이 아닙니다.", "공개 자료 기반 제3자 분석이며 Micron 및 계열사가 작성·검토·승인한 자료가 아닙니다."),
        "en": ("This document is not investment advice.", "This is third-party analysis based on public information and has not been prepared, reviewed, or approved by Micron or its affiliates."),
    }
    for locale, needles in required.items():
        require(locale in texts, f"missing edition: {locale}")
        strings = yaml.safe_load((PACKAGE_DIR / "i18n" / f"{locale}.yaml").read_text(encoding="utf-8"))
        require(strings["disclaimer"]["third_party"] == needles[1], f"{locale} D9 third-party text mismatch")
        for needle in needles:
            require(needle in texts[locale], f"missing {locale} disclaimer text: {needle}")


def gate_inventory_days(manifest: Manifest) -> None:
    require(not any("industry_inventory" in key for key in manifest.facts), "industry inventory fact is forbidden")
    ratios = [fact for key, fact in manifest.facts.items() if key.startswith("ratio.inventory_days.")]
    for fact in ratios:
        if fact.period in {"FY2027", "FY2028"}:
            require(fact.status == "UNAVAILABLE", f"RLE inventory days must be unavailable: {fact.fact_id}")
            continue
        if fact.status == "UNAVAILABLE":
            require(fact.period == "FY2026", f"historical inventory days must be available: {fact.fact_id}")
            continue
        require(fact.lineage is not None, f"inventory-days lineage missing: {fact.fact_id}")
        require(fact.lineage.get("formula") == "inventory_days_v1", f"inventory-days formula mismatch: {fact.fact_id}")
        inputs = fact.lineage.get("inputs", [])
        require(len(inputs) == 4, f"inventory-days lineage must have four inputs: {fact.fact_id}")
        source_facts = [manifest.fact(fact_id) for fact_id in inputs]
        if fact.period == "FY2026":
            require(source_facts[3].raw_value == 53, "FY2026 inventory days must use 53 weeks")
        expected = inventory_days(
            source_facts[0].raw_value,
            source_facts[1].raw_value,
            source_facts[2].raw_value,
            int(source_facts[3].raw_value),
        )
        require(fact.raw_value == expected, f"inventory-days value mismatch: {fact.fact_id}")


def gate_trailing_pb(manifest: Manifest, caption: dict[str, str]) -> None:
    pb_facts = [fact for key, fact in manifest.facts.items() if key.startswith("val.pb_trailing.")]
    require(len(pb_facts) == 1, "exactly one trailing P/B fact is allowed")
    fact = pb_facts[0]
    require(fact.fact_id == "val.pb_trailing.FY2026.A-8K", "RLE or non-FY2026 trailing P/B is forbidden")
    require(fact.lineage is not None and fact.lineage.get("formula") == "trailing_pb_v1", "trailing P/B lineage is invalid")
    inputs = fact.lineage.get("inputs", [])
    require(len(inputs) == 3, "trailing P/B lineage must have three inputs")
    expected = trailing_pb(*(manifest.fact(fact_id).raw_value for fact_id in inputs))
    require(fact.raw_value == expected, "trailing P/B value mismatch")
    for key in ("price_date", "equity_date", "shares_date"):
        require(bool(caption.get(key)), f"trailing P/B caption lacks {key}")


def gate_market_data_box(config: dict[str, Any]) -> None:
    rows = config.get("rows", [])
    require(len(rows) == 4, "market-data box must contain exactly four rows")
    fact_ids = []
    for row in rows:
        fact_ids.extend(re.findall(r"\{\{fact:([^}]+)}}", row.get("value", "")))
    require(
        fact_ids == [
            "market.price.2026-10-01.CITED",
            "market.shares_outstanding.CITED",
            "market.market_cap.2026-10-01.CALCULATED",
            "meta.fiscal_year_end.FY2026.A-8K",
        ],
        "market-data box fact set or order is invalid",
    )
    serialized = json.dumps(config, ensure_ascii=False).lower()
    for forbidden in ("52-week", "52주", "volume", "거래량", "current ratio", "유동비율"):
        require(forbidden not in serialized, f"forbidden market-data field: {forbidden}")


def gate_g12c_conflict(
    record: ConflictConfirmation,
    edition: str,
    now_kst: datetime,
    rendered: dict[str, dict[str, Any]],
) -> None:
    require(record.status == "not_held", "author holding status is not not_held")
    age = now_kst - record.confirmed_at_kst
    require(timedelta(0) <= age <= timedelta(hours=24), "conflict confirmation is future-dated or expired")
    require(record.edition == edition, "conflict confirmation edition mismatch")
    require(set(rendered) == {"ko", "en"}, "conflict disclosure requires both locales")
    for locale in ("ko", "en"):
        strings = yaml.safe_load((PACKAGE_DIR / "i18n" / f"{locale}.yaml").read_text(encoding="utf-8"))
        stamp = record.confirmed_at_kst.isoformat()
        expected_full = strings["disclaimer"]["conflict"].replace("{{conflict.confirmed_at_kst}}", stamp)
        expected_short = strings["disclaimer"]["conflict_short"]
        edition_render = rendered[locale]
        require(edition_render.get("record_blob_sha") == record.git_blob_sha, f"{locale} uses a different confirmation record")
        require(expected_full in edition_render.get("cover", ""), f"{locale} cover conflict disclosure missing")
        require(expected_full in edition_render.get("ending", ""), f"{locale} ending conflict disclosure missing")
        pages = edition_render.get("pdf_pages", [])
        require(bool(pages), f"{locale} conflict footer evidence missing")
        require(all(expected_short in page for page in pages), f"{locale} conflict_short absent from a PDF page")


DRYRUN_CONFLICT = {
    "ko": "[DRY RUN — 발행 전 보유 확인 미실시]",
    "en": "[DRY RUN — pre-publication holding check not performed]",
}


def gate_g12c_dryrun(rendered: dict[str, dict[str, Any]], confirmation: ConflictConfirmation | None = None) -> None:
    require(confirmation is None, "DRYRUN must not receive a conflict confirmation record")
    require(set(rendered) == {"ko", "en"}, "DRYRUN disclosure requires both locales")
    for locale in ("ko", "en"):
        strings = yaml.safe_load((PACKAGE_DIR / "i18n" / f"{locale}.yaml").read_text(encoding="utf-8"))
        expected = DRYRUN_CONFLICT[locale]
        edition = rendered[locale]
        pages = edition.get("pdf_pages", [])
        require(expected in edition.get("cover", ""), f"{locale} DRYRUN cover disclosure missing")
        require(expected in edition.get("ending", ""), f"{locale} DRYRUN ending disclosure missing")
        normalized_pages = [" ".join(page.split()) for page in pages]
        require(bool(pages) and all(expected in page for page in normalized_pages), f"{locale} DRYRUN disclosure absent from a PDF page")
        all_text = "\n".join([edition.get("cover", ""), edition.get("ending", ""), *pages])
        require(strings["disclaimer"]["conflict_short"] not in all_text, f"{locale} real conflict_short rendered in DRYRUN")
        prefix, suffix = strings["disclaimer"]["conflict"].split("{{conflict.confirmed_at_kst}}")
        require(not (prefix in all_text and suffix in all_text), f"{locale} real conflict disclosure rendered in DRYRUN")


def gate_dryrun_boundary(output: Path, manifest: Manifest, audit_payload: bytes) -> None:
    resolved = output.resolve()
    root = (Path(__file__).resolve().parents[3] / "logs" / "_mu_report_runs").resolve()
    require(root in resolved.parents and resolved.name.startswith("dryrun_"), "DRYRUN output escaped its allowed directory")
    require(manifest.metadata.get("layer") == "DRYRUN", "DRYRUN manifest layer missing")
    require(manifest.metadata.get("fixture_values") is False, "fixture values are forbidden in DRYRUN")
    audit = json.loads(audit_payload)
    allowed = {item["path"] for item in load_pins()["E2-A"]}
    reads = {item["path"] for item in audit["reads"]}
    require(reads <= allowed, "DRYRUN read a non-E2-A input")


def gate_dryrun_watermark(texts: list[str], expected: str) -> None:
    normalized = [" ".join(text.split()) for text in texts]
    require(bool(texts) and all(expected in text for text in normalized), "DRYRUN watermark absent from a PDF page")


def gate_dryrun_placeholders(manifest: Manifest, chart_manifest: dict[str, Any]) -> None:
    for fact_id in (
        "is.revenue.FY2026Q4.A-8K",
        "market.price.2026-10-01.CITED",
        "ratio.inventory_days.FY2027.RLE",
        "ratio.inventory_days.FY2028.RLE",
    ):
        fact = manifest.fact(fact_id)
        require(fact.status == "UNAVAILABLE" and fact.raw_value is None, f"DRYRUN placeholder contains a value: {fact_id}")
    heatmap = chart_manifest.get("07_valuation_heatmap", {})
    require(heatmap.get("fact_ids") == [] and heatmap.get("values") == [], "DRYRUN valuation heatmap contains a value")


def gate_theme(html_text: str, i18n_texts: Iterable[str]) -> dict[str, float]:
    ratios = {
        "ink_on_paper": contrast_ratio(COLORS["ink"], COLORS["paper"]),
        "primary_on_paper": contrast_ratio(COLORS["primary"], COLORS["paper"]),
        "gray_dark_on_paper": contrast_ratio(COLORS["gray_dark"], COLORS["paper"]),
        "warning_on_paper": contrast_ratio(COLORS["warning"], COLORS["paper"]),
    }
    require(all(value >= 4.5 for value in ratios.values()), f"theme contrast below 4.5: {ratios}")
    require(not any(re.search(r"#[0-9A-Fa-f]{6}", text) for text in i18n_texts), "i18n contains a color value")
    require(re.search(r"<img[^>]+(?:logo|micron)", html_text, re.I) is None, "logo image is forbidden")
    return ratios


def gate_g13_charts(chart_manifest: dict[str, Any]) -> None:
    require(set(CHART_NAMES) <= set(chart_manifest), "required chart set is incomplete")
    from .revision import REVISION_CHARTS

    if set(chart_manifest) & set(REVISION_CHARTS):
        require(set(chart_manifest) == set(CHART_NAMES) | set(REVISION_CHARTS), "revision requires sixteen charts")


def _annual_year(period: str) -> int | None:
    match = re.fullmatch(r"FY(\d{4})(?:[AE])?", period)
    return int(match.group(1)) if match else None


def gate_g13c_chart_semantics(manifest: Manifest, chart_manifest: dict[str, Any]) -> None:
    contract = {
        "01_quarterly_revenue_margin": ("bar_line", 3, ("is.revenue.", "ratio.gross_margin.", "ratio.operating_margin."), ("FQ4-26",)),
        "02_business_unit_mix": ("stacked_bar", 4, ("bu.revenue.",), ("FQ4-26",)),
        "03b_guidance_beat_history": ("guidance_actual", 2, ("guidance.gm_gaap.", "actual.gm_gaap."), ("FQ4-26",)),
        "04_beat_history": ("two_panel", 2, ("beat.revenue_pct.", "beat.eps_gaap_usd."), ("FQ4-26",)),
        "05_scenario_fan": ("fan", 4, ("is.revenue.", "prereg.revenue."), ("FY2027Q1", "FY2027Q2", "FY2027Q3", "FY2027Q4")),
        "06_annual_income": ("grouped_bar", 3, ("is.revenue.", "is.operating_income.", "is.net_income."), ("FY2026A-8K", "FY2027E", "FY2028E")),
        "07_valuation_heatmap": ("heatmap", 1, (), ("PRICE_UNAVAILABLE",)),
        "08_cash_flow_capex_net_cash": ("multi_series", 3, ("cf.fcf_adjusted.", "cf.net_capex.", "bs.net_cash_unadjusted."), ("SCA_UNAVAILABLE",)),
    }
    if manifest.metadata.get("layer") == "E2-B":
        contract.update({
            "05_scenario_fan": ("fan", 4, ("is.revenue.", "rle.revenue."), ()),
            "06_annual_income": ("grouped_bar", 3, ("is.revenue.", "is.operating_income.", "is.net_income."), ()),
            "07_valuation_heatmap": ("heatmap", 1, ("val.pe_sensitivity.",), ()),
            "08_cash_flow_capex_net_cash": ("multi_series", 3, ("cf.fcf_adjusted.", "cf.net_capex.", "bs.net_cash_unadjusted."), ()),
        })
        for chart_id in ("01_quarterly_revenue_margin", "02_business_unit_mix", "03b_guidance_beat_history", "04_beat_history"):
            chart_type, count, families, _ = contract[chart_id]
            contract[chart_id] = (chart_type, count, families, ())
    if manifest.metadata.get("revision") == 1:
        from .revision import REVISION_CHARTS

        contract.update({key: (*value, ()) for key, value in REVISION_CHARTS.items()})
    require(set(chart_manifest) == set(contract), "G-13c chart set differs from the phase chart contract")
    for chart_id, (chart_type, count, families, slots) in contract.items():
        chart = chart_manifest[chart_id]
        require(chart.get("chart_type") == chart_type, f"G-13c chart type mismatch: {chart_id}")
        require(chart.get("series_count") == count, f"G-13c series count mismatch: {chart_id}")
        require(tuple(chart.get("allowed_families", ())) == families, f"G-13c fact-family contract mismatch: {chart_id}")
        require(tuple(chart.get("placeholder_slots", ())) == slots, f"G-13c placeholder contract mismatch: {chart_id}")
        for fact_id in chart.get("fact_ids", []):
            require(any(fact_id.startswith(family) for family in families), f"G-13c forbidden fact family: {chart_id}/{fact_id}")
            if manifest.metadata.get("revision") == 1:
                fact = manifest.fact(fact_id)
                require(fact.status == "AVAILABLE" and isinstance(fact.raw_value, (int, float)), f"G-13c nonnumeric or unavailable point: {fact_id}")
        if manifest.metadata.get("revision") == 1:
            require(chart.get("values") == [float(manifest.fact(key).raw_value) for key in chart["fact_ids"]], f"G-13c point values differ: {chart_id}")
    if manifest.metadata.get("revision") == 1:
        from .revision import RANGE_J

        ranges = chart_manifest["10_price_bit_ranges"]
        require("J" in ranges["caption"]["basis"] and "J" in ranges.get("footnote", ""), "C10 must disclose judgement grade J")
        for low, high in RANGE_J.values():
            require(f"{low}–{high}%" in ranges.get("footnote", ""), "C10 conversion table is incomplete")
        for key in ranges["fact_ids"]:
            fact = manifest.fact(key)
            require(fact.basis == "J" and fact.lineage.get("grade") == "J", "C10 range fact lacks J provenance")
            phrase = next((phrase for phrase in RANGE_J if phrase in fact.lineage["company_wording"]), None)
            expected = RANGE_J[phrase][0 if key.endswith(".low") else 1] if phrase else None
            require(fact.raw_value == expected, "C10 endpoint differs from approved J conversion")
        costs = chart_manifest["14_opex_net_capex_trend"]
        missing = {(row["period"], row["series"]) for row in manifest.metadata["c14_missing"]}
        for year in (2024, 2025, 2026):
            for quarter in range(1, 5):
                period = f"FY{year}Q{quarter}"
                for metric in ("opex", "net_capex"):
                    plotted = [key for key in costs["fact_ids"] if manifest.fact(key).period == period and f".{metric}." in key]
                    require(bool(plotted) != ((period, metric) in missing), f"C14 missing-quarter contract differs: {period}/{metric}")
                    if (period, metric) in missing:
                        require(f"{period}/{metric}" in costs.get("footnote", ""), "C14 omission is absent from footnote")

    cash_chart = chart_manifest["08_cash_flow_capex_net_cash"]
    periods = cash_chart.get("periods", [])
    require(periods and all(_annual_year(period) is not None for period in periods), "G-13c cash chart mixes annual and quarterly periods")
    for family in contract["08_cash_flow_capex_net_cash"][2]:
        series_periods = cash_chart.get("series_periods", {}).get(family, [])
        years = [_annual_year(period) for period in series_periods]
        require(all(year is not None for year in years), f"G-13c nonannual period in cash series: {family}")
        require(years == sorted(years), f"G-13c cash series is not chronological: {family}")
        require(len(years) == len(set(years)), f"G-13c duplicate cash-series period: {family}")
    chart_years = [_annual_year(period) for period in periods]
    first_year, last_year = min(chart_years), max(chart_years)
    plotted = set(cash_chart.get("fact_ids", []))
    for family in contract["08_cash_flow_capex_net_cash"][2]:
        expected = {
            fact.fact_id
            for fact in manifest.facts.values()
            if fact.fact_id.startswith(family)
            and isinstance(fact.raw_value, (int, float))
            and (year := _annual_year(fact.period)) is not None
            and first_year <= year <= last_year
            and (manifest.metadata.get("layer") != "E2-B" or not re.search(r"\.(?:bear|bull)$", fact.fact_id))
        }
        require(expected <= plotted, f"G-13c cash series omits manifest values: {family}/{sorted(expected - plotted)}")


def gate_categorical_verbatim(values: Iterable[str], normalized_source_text: str) -> None:
    for value in values:
        normalized = " ".join(value.split())
        require(normalized in normalized_source_text, f"categorical value is not an exact source substring: {value}")


def gate_g13b_chart_identity(manifest: Manifest, chart_manifest: dict[str, Any], base_pe: float | None = None, heatmap_base: float | None = None) -> None:
    for chart_id, chart in chart_manifest.items():
        fact_ids = chart["fact_ids"]
        values = chart["values"]
        require(len(fact_ids) == len(values), f"chart point shape mismatch: {chart_id}")
        for fact_id, value in zip(fact_ids, values, strict=True):
            fact = manifest.fact(fact_id)
            require(float(fact.raw_value) == float(value), f"chart/table mismatch: {chart_id}/{fact_id}")
    if base_pe is not None:
        require(heatmap_base == base_pe, "heatmap base cell differs from primary table")


def gate_g14_provenance(manifest: Manifest) -> None:
    manifest.validate()
    for fact in manifest.facts.values():
        source = manifest.sources[fact.source_id]
        require(bool(source.sha256 and source.as_of and source.path), f"incomplete provenance: {fact.fact_id}")
        if fact.basis == "CALCULATED":
            require(bool(fact.lineage), f"missing lineage: {fact.fact_id}")


def gate_g14_template_numbers(template: str) -> None:
    scrubbed = PLACEHOLDER_RE.sub("", template)
    scrubbed = re.sub(r"Edition 1, Revision 1|제1판 개정 1", "", scrubbed)
    scrubbed = re.sub(r"EX-99\.1|\b4-lever\b|(?:SCORED|FROZEN)\s*§(?:\([^)]*\)|\d+)|p\.\s*\d+(?:[–-]\d+)?|A\d+[–-]A\d+", "", scrubbed)
    scrubbed = re.sub(r"(?:Plan|계획)\s*§\d+(?:[-–]\d+)*", "", scrubbed, flags=re.I)
    scrubbed = re.sub(r"FY20\d{2}\s*Q[1-4]|FY20\d{2}E?|FY\d{2}(?:A|E)?|FQ[1-4]-\d{2}|20\d{2}-\d{2}-\d{2}(?:\s+23:59\s+KST)?|\b(?:14|53) weeks?\b|(?:14|53)주|A-\d+K|\b8-K\b|10-[QK]|\bQ[1-4]\b|(?:Figure|그림)\s*(?:1[0-6]|[1-9])\b\.?", "", scrubbed)
    scrubbed = re.sub(r"(?m)^\s{2}\d{2}[a-z]?_[^:]+:", "", scrubbed)
    tokens = re.findall(r"(?<!\w)[-+]?\$?\d+(?:[.,]\d+)?%?", scrubbed)
    require(not tokens, f"template contains unreferenced numeric/date literals: {tokens}")


def _reference_map(entries: list[dict[str, str]], locale: str) -> dict[str, dict[str, str]]:
    mapped: dict[str, dict[str, str]] = {}
    for entry in entries:
        position = entry.get("position", "")
        require(bool(position), f"missing {locale} reference position")
        require(position not in mapped, f"duplicate {locale} reference position: {position}")
        mapped[position] = entry
    return mapped


def _number_multiset(text: str) -> Counter[str]:
    # These equivalent edition labels tokenize Korean/English digits differently.
    text = text.replace("Edition 1, Revision 1", "").replace("제1판 개정 1", "")
    tokens = re.findall(r"(?<!\w)[-+]?\$?\d[\d,]*(?:\.\d+)?%?", text)
    return Counter(token.replace(",", "").replace("$", "").lstrip("+") for token in tokens)


def gate_g15_parity(
    reference_logs: dict[str, list[dict[str, str]]],
    manifest: Manifest,
    rendered_texts: dict[str, str],
) -> None:
    require(set(reference_logs) == {"ko", "en"}, "both locale reference logs are required")
    require(set(rendered_texts) == {"ko", "en"}, "both rendered locale texts are required")
    by_locale = {locale: _reference_map(entries, locale) for locale, entries in reference_logs.items()}
    require(set(by_locale["ko"]) == set(by_locale["en"]), "KO/EN reference positions differ")
    for position in sorted(by_locale["ko"]):
        ko = by_locale["ko"][position]
        en = by_locale["en"][position]
        for field in ("kind", "role", "fact_id", "raw_value", "unit", "period", "basis", "label", "status"):
            require(ko.get(field) == en.get(field), f"KO/EN {field} differs at {position}")
        fact = manifest.fact(ko["fact_id"])
        for field in ("raw_value", "unit", "period", "basis", "label", "status"):
            require(ko.get(field) == getattr(fact, field), f"reference differs from manifest at {position}: {field}")
    require(
        _number_multiset(rendered_texts["ko"]) == _number_multiset(rendered_texts["en"]),
        "KO/EN normalized number multisets differ",
    )


def gate_g15_toc(entries: dict[str, list[dict[str, str]]]) -> None:
    require(set(entries) == {"ko", "en"}, "both locale TOCs are required")
    ko_ids = [entry.get("id") for entry in entries["ko"]]
    en_ids = [entry.get("id") for entry in entries["en"]]
    require(len(ko_ids) == 12 and len(set(ko_ids)) == 12, "TOC must contain twelve unique sections")
    require(ko_ids == en_ids, "KO/EN TOC section order or anchors differ")


def gate_narrative_contract(narrative: dict[str, Any], manifest: Manifest) -> None:
    """Require the approved KO/EN narrative shape and fact-token order to match."""
    from .revision import NEW_BINDINGS, NEW_SECTIONS

    expected = {"thesis", "scenarios", "risks", "catalysts", *NEW_SECTIONS}
    require(set(narrative.get("meta", {}).get("sections", [])) == expected, "narrative meta sections differ")
    bindings = narrative.get("fact_bindings", [])
    binding_ids = [item.get("token") for item in bindings]
    require(len(binding_ids) == 12 + len(NEW_BINDINGS) and len(set(binding_ids)) == len(binding_ids), "narrative revision bindings must be unique and complete")
    require(set(binding_ids) <= set(manifest.facts), "narrative binding is absent from manifest")

    def shape(value: Any) -> Any:
        if isinstance(value, dict):
            return {key: shape(child) for key, child in value.items() if key != "src"}
        if isinstance(value, list):
            return [shape(child) for child in value]
        return type(value).__name__

    def tokens(value: Any) -> list[str]:
        if isinstance(value, dict):
            return [token for child in value.values() for token in tokens(child)]
        if isinstance(value, list):
            return [token for child in value for token in tokens(child)]
        return PLACEHOLDER_RE.findall(value) if isinstance(value, str) else []

    for section in expected:
        localized = narrative.get(section, {})
        require(set(localized) == {"ko", "en"}, f"narrative locale missing: {section}")
        if section in {"business_structure", "methodology"}:
            count = 4 if section == "business_structure" else 5
            require(all(len(localized[locale]) == count for locale in ("ko", "en")), f"narrative paragraph count differs: {section}")
            require(_number_multiset(" ".join(localized["ko"])) == _number_multiset(" ".join(localized["en"])), f"narrative numeric parity differs: {section}")
        require(shape(localized["ko"]) == shape(localized["en"]), f"narrative KO/EN structure differs: {section}")
        ko_ids = [match[7:-2] for match in tokens(localized["ko"])]
        en_ids = [match[7:-2] for match in tokens(localized["en"])]
        require(ko_ids == en_ids, f"narrative KO/EN fact-token order differs: {section}")
    used = {
        match[7:-2]
        for section in expected
        for locale in ("ko", "en")
        for match in tokens(narrative[section][locale])
    }
    require(used == set(binding_ids), "narrative binding list and used tokens differ")


def gate_ed1_rendered(texts: dict[str, str]) -> None:
    require(set(texts) == {"ko", "en"}, "both ed1 locales are required")
    for locale, text in texts.items():
        require("section_stub" not in text, f"section_stub remains in {locale} ed1")
        require("FIXTURE" not in text, f"fixture prose remains in {locale} ed1")
        require(PLACEHOLDER_RE.search(text) is None, f"unresolved fact token remains in {locale} ed1")


def gate_g20_phase_audit(reader: EvidenceReader, phases: tuple[str, ...]) -> None:
    payload_1 = reader.audit_payload()
    payload_2 = reader.audit_payload()
    require(payload_1 == payload_2, "audit manifest is non-deterministic")
    audit = json.loads(payload_1)
    expected = {
        item["path"]
        for phase in phases
        for item in load_pins()[phase]
        if item["sha256"] is not None
    }
    require({row["path"] for row in audit["reads"]} == expected, "phase audit does not cover every populated pin")


def gate_g15b_localized_ui(chart_manifests: dict[str, dict[str, Any]], rendered_texts: dict[str, str]) -> None:
    require(set(chart_manifests) == {"ko", "en"}, "G-15b requires both chart locales")
    require(set(rendered_texts) == {"ko", "en"}, "G-15b requires both rendered locales")
    require(set(chart_manifests["ko"]) == set(chart_manifests["en"]), "G-15b chart sets differ")
    prohibited = re.compile(r"\b(?:Figure|Revenue|Margin|Guidance|Actual|Cash Flow|Net Cash|Business Unit|Metric|Value|Basis|Period)\b", re.I)
    whitelist = re.compile(r"\b(?:GAAP|PREREG_A|RLE|A-8K|FQ\d(?:-\d{2})?|FY\d{2,4}[A-Z0-9-]*|CMBU|CDBU|MCBU|AEBU|DRAM|NAND|FCF|SCA|ROE|EPS|USD|Micron|UNAVAILABLE)\b")
    for chart_id, chart in chart_manifests["ko"].items():
        text = " ".join(chart.get("ui_text", []))
        scrubbed = whitelist.sub("", text)
        require(prohibited.search(scrubbed) is None, f"G-15b English UI remains in KO chart: {chart_id}")
        require(chart.get("path", "").endswith("_ko.png"), f"G-15b KO chart asset is not locale-specific: {chart_id}")
        require(chart_manifests["en"][chart_id].get("path", "").endswith("_en.png"), f"G-15b EN chart asset is not locale-specific: {chart_id}")
    require("| Metric | Value | Basis |" not in rendered_texts["ko"], "G-15b English table UI remains in KO edition")
    raw_codes = re.compile(r"\b(?:ABOVE_HIGH|IN_RANGE|BELOW_LOW|NO_LABEL)\b")
    raw_status_codes = re.compile(r"\b(?:AVAILABLE|AVAILABLE_BACKSOLVED|UNAVAILABLE_WITHOUT_ASSUMPTIONS|NOT_ACTIVATED|PASS_GAAP)\b")
    for locale, text in rendered_texts.items():
        require(raw_codes.search(text) is None, f"G-15b raw verdict code remains in {locale} edition")
        require(raw_status_codes.search(text) is None, f"G-15b raw status code remains in {locale} edition")
    allowed_terms = re.compile(
        r"\b(?:Micron|Technology|NASDAQ|MU|GAAP|PREREG_A|RLE|A-8K|FQ\d(?:-\d{2})?|FY\d{2,4}[A-Z0-9-]*|"
        r"CMBU|CDBU|MCBU|AEBU|DRAM|NAND|HBM|FCF|SCA|ROE|EPS|USD|PP&E|NetCash|source_id|NOT_IN_SOURCE|"
        r"UNAVAILABLE(?:_WITHOUT_ASSUMPTIONS)?|PASS|CITED|CONFIRMED|ESTIMATED)\b",
        re.I,
    )
    ko_body = allowed_terms.sub("", rendered_texts["ko"])
    ko_body = re.sub(r"\bSRC-[A-Za-z0-9_-]+\b", "", ko_body)
    english_cell_run = re.compile(r"\b[A-Za-z][A-Za-z'-]*(?:[\s,;:()/*+=.-]+[A-Za-z][A-Za-z'-]*){3,}\b")
    for line_number, line in enumerate(ko_body.splitlines(), start=1):
        if not line.strip().startswith("|"):
            continue
        for column, cell in enumerate(line.strip().strip("|").split("|"), start=1):
            match = english_cell_run.search(cell)
            require(match is None, f"G-15b English sentence remains in KO table cell: line {line_number}, column {column}: {match.group(0)[:120] if match else ''}")
    english_run = re.compile(r"\b[A-Za-z][A-Za-z'-]*(?:[\s,;:()/.]+[A-Za-z][A-Za-z'-]*){4,}\b")
    match = english_run.search(ko_body)
    require(match is None, f"G-15b English sentence remains in KO edition: {match.group(0)[:120] if match else ''}")


G23_REQUIRED_AVAILABLE = {
    "market.price.2026-10-01.CITED",
    "market.shares_outstanding.CITED",
    "market.market_cap.2026-10-01.CALCULATED",
    "meta.fiscal_year_end.FY2026.A-8K",
    "meta.price_date.2026-10-01.CITED",
    "meta.shares_date.CITED",
    "meta.equity_date.FY2026.A-8K",
    "bs.total_equity.FY2026.A-8K",
    "val.pb_trailing.FY2026.A-8K",
    "is.revenue.FY2026Q4.A-8K",
    "is.gross_margin.FY2026Q4.A-8K",
    "is.operating_margin.FY2026Q4.A-8K",
    "is.revenue.FY2026A.A-8K",
    "is.net_income.FY2026A.A-8K",
    "is.diluted_eps.FY2026A.A-8K",
    "is.revenue.FY2027E.base",
    "is.revenue.FY2028E.base",
    "is.pretax.FY2027E.base",
    "is.pretax.FY2028E.base",
    "is.net_income.FY2027E.base",
    "is.net_income.FY2028E.base",
    "cf.fcf_adjusted.FY2027E.base",
    "cf.fcf_adjusted.FY2028E.base",
    "bs.net_cash_unadjusted.FY2027E.base",
    "bs.net_cash_unadjusted.FY2028E.base",
}


def gate_g23_availability(
    manifest: Manifest,
    chart_manifests: dict[str, dict[str, Any]],
    rendered_texts: dict[str, str] | None = None,
) -> dict[str, Any]:
    """Enforce ed1 availability and return every unavailable fact/chart slot for audit."""
    unavailable_facts = sorted(
        fact.fact_id for fact in manifest.facts.values() if fact.status in UNAVAILABLE_STATUSES
    )
    chart_slots = sorted(
        f"{locale}/{chart_id}/{slot}"
        for locale, charts in chart_manifests.items()
        for chart_id, chart in charts.items()
        for slot in chart.get("placeholder_slots", [])
    )
    rendered_locations = sorted(
        f"{locale}:line {line_number}: {line.strip()[:160]}"
        for locale, text in (rendered_texts or {}).items()
        for line_number, line in enumerate(text.splitlines(), start=1)
        if re.search(r"\b(?:UNAVAILABLE|NOT_IN_SOURCE|UNAVAILABLE_WITHOUT_ASSUMPTIONS)\b", line)
    )
    missing = sorted(
        fact_id
        for fact_id in G23_REQUIRED_AVAILABLE
        if fact_id not in manifest.facts or manifest.fact(fact_id).status != "AVAILABLE"
    )
    audit = {
        "unavailable_count": len(unavailable_facts) + len(chart_slots) + len(rendered_locations),
        "unavailable_locations": [*unavailable_facts, *chart_slots, *rendered_locations],
        "required_available_failures": missing,
    }
    require(not missing, f"G-23 required available locations are unavailable: {missing}")
    for locale, text in (rendered_texts or {}).items():
        rows = [line for line in text.splitlines() if line.startswith("| ")]
        for label, metric, occurrence in (
            ("세전이익" if locale == "ko" else "Pretax income", "pretax", 0),
            ("순이익" if locale == "ko" else "Net income", "net_income", 1),
        ):
            matches = [row for row in rows if row.split("|")[1].strip() == label]
            require(len(matches) > occurrence, f"G-23 required rendered row missing: {locale}/{label}/{occurrence}")
            cells = [cell.strip() for cell in matches[occurrence].strip("|").split("|")]
            require(len(cells) == 8, f"G-23 required rendered row shape mismatch: {locale}/{label}/{occurrence}")
            for year, cell in zip((2027, 2028), cells[-2:]):
                fact_id = f"is.{metric}.FY{year}E.base"
                require(cell == manifest.fact(fact_id).display[locale], f"G-23 required rendered cell unavailable or mismatched: {locale}/{fact_id}: {cell}")
    require(not chart_slots, f"G-23 chart placeholders remain: {chart_slots}")
    return audit


def gate_g3f_scored_source(manifest: Manifest, scored_text: str) -> None:
    """Require every displayed scoring value to be parsed from the pinned SCORED source."""
    from .narrative import _scored_values

    parsed = _scored_values(scored_text)
    fact_map = {
        "revenue_actual": "scored.revenue_actual.FQ4FY26", "revenue_forecast": "scored.revenue_forecast.FQ4FY26",
        "revenue_error": "scored.revenue_error.FQ4FY26", "revenue_ape": "scored.revenue_ape.FQ4FY26",
        "eps_actual": "scored.eps_actual.FQ4FY26", "eps_forecast": "scored.eps_forecast.FQ4FY26",
        "eps_error": "scored.eps_error.FQ4FY26", "eps_ape": "scored.eps_ape.FQ4FY26",
        "nongaap_eps_actual": "scored.nongaap_eps_actual.FQ4FY26", "nongaap_eps_forecast": "scored.nongaap_eps_forecast.FQ4FY26",
        "nongaap_eps_error": "scored.nongaap_eps_error.FQ4FY26", "nongaap_eps_ape": "scored.nongaap_eps_ape.FQ4FY26",
        "opex_actual": "scored.opex_actual.FQ4FY26", "opex_forecast": "scored.opex_forecast.FQ4FY26",
        "opex_error": "scored.opex_error.FQ4FY26", "opex_ape": "scored.opex_ape.FQ4FY26",
        "revenue_band_low": "scored.revenue_band_low.FQ4FY26", "revenue_band_high": "scored.revenue_band_high.FQ4FY26",
        "eps_band_low": "scored.eps_band_low.FQ4FY26", "eps_band_high": "scored.eps_band_high.FQ4FY26",
        "lever_revenue": "scored.lever.revenue.FQ4FY26", "lever_op_margin": "scored.lever.op_margin.FQ4FY26",
        "lever_op_to_ni": "scored.lever.op_to_ni.FQ4FY26", "lever_shares": "scored.lever.shares.FQ4FY26",
        "lever_total": "scored.lever.total.FQ4FY26", "labels_hit": "scored.labels_hit.FQ4FY26",
        "band_coverage": "scored.band_coverage.FQ4FY26", "sf7": "scored.sf7.FQ1FY27",
    }
    for key, fact_id in fact_map.items():
        fact = manifest.fact(fact_id)
        require(fact.source_id == "SRC-SCORED-FQ4FY26", f"G-3f non-SCORED provenance: {fact_id}")
        expected, actual = parsed[key], fact.raw_value
        if isinstance(expected, (int, float)):
            require(isinstance(actual, (int, float)) and abs(float(actual) - float(expected)) < 1e-9, f"G-3f value mismatch: {fact_id}")
        else:
            require(actual == expected, f"G-3f value mismatch: {fact_id}")


def gate_g24_xlsx_labels(xlsx_path: str | Path) -> list[str]:
    pattern = re.compile(r"FIXTURE|DRY\s+RUN|E2-A\s+fixture", re.I)
    findings = []
    workbook = load_workbook(xlsx_path, read_only=True, data_only=False)
    try:
        for sheet in workbook:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and (match := pattern.search(cell.value)):
                        findings.append(f"{sheet.title}!{cell.coordinate}: {match.group(0)}")
    finally:
        workbook.close()
    require(not findings, f"G-24 XLSX rehearsal labels remain: {findings}")
    return findings


def gate_g24_raw_markup(pdf_texts: dict[str, list[str]], xlsx_path: str | Path | None = None) -> dict[str, list[str]]:
    require(set(pdf_texts) == {"ko", "en"}, "G-24 requires both PDF locales")
    patterns = {
        "heading": re.compile(r"####"),
        "bold": re.compile(r"\*\*"),
        "backtick": re.compile(r"`"),
        "brace": re.compile(r"[{}]"),
        "source_id": re.compile(r"source_id:"),
        "python_tuple": re.compile(r"\(\s*[-+]?\d[\d,.]*\s*,\s*\)"),
    }
    findings: dict[str, list[str]] = {"ko": [], "en": []}
    for locale, pages in pdf_texts.items():
        for page_number, text in enumerate(pages, start=1):
            for label, pattern in patterns.items():
                if pattern.search(text):
                    findings[locale].append(f"page {page_number}: {label}")
    require(not findings["ko"] and not findings["en"], f"G-24 raw markup remains: {findings}")
    if xlsx_path is not None:
        findings["xlsx"] = gate_g24_xlsx_labels(xlsx_path)
    return findings


def gate_g25_cover_dates(rendered_texts: dict[str, str]) -> None:
    require(set(rendered_texts) == {"ko", "en"}, "G-25 requires both rendered locales")
    expected = {
        "ko": (
            r"발행일:\s*\d{4}-\d{2}-\d{2}",
            r"자료 기준일:\s*2026-10-01\(주가\)\s*·\s*2026-09-30\(실적\)",
            r"정보 컷오프:\s*2026-10-04\s+KST",
        ),
        "en": (
            r"Issue date:\s*\d{4}-\d{2}-\d{2}",
            r"Data as of:\s*2026-10-01 \(price\)\s*·\s*2026-09-30 \(results\)",
            r"Information cutoff:\s*2026-10-04\s+KST",
        ),
    }
    for locale, patterns in expected.items():
        cover = rendered_texts[locale].split('<a id="section-company"></a>', 1)[0]
        for pattern in patterns:
            require(re.search(pattern, cover) is not None, f"G-25 cover date missing or non-date: {locale}: {pattern}")


def gate_g16_caption(caption: dict[str, str]) -> None:
    for key in ("number", "title", "unit", "source", "as_of", "basis"):
        require(bool(caption.get(key)), f"caption lacks {key}")
    require(re.fullmatch(r"(?:Figure|그림)\s+(?:1[0-6]|[1-9])\.", caption["number"]) is not None, "caption number is not a concrete Figure 1-16 label")
    require(caption["unit"] not in {"per axis", "per chart axis", "various", "각 축", "차트 축 기준"}, "caption unit is generic")
    require(any(token in caption["unit"] for token in ("USD", "%")), "caption unit is not concrete")
    require(
        any(token in caption["source"] for token in ("SRC-", "Micron", "SEC", "MU", "고정 테스트", "Fixed test", "본 리포트", "This report", "저자 승인", "Author-approved"))
        or "기준 주가" in caption["source"]
        or "reference price" in caption["source"].lower(),
        "caption source lacks a concrete source identifier",
    )
    require(re.fullmatch(r"\d{4}-\d{2}-\d{2}", caption["as_of"]) is not None, "caption as_of is not an ISO date")
    require(any(token in caption["basis"] for token in ("GAAP", "UNAVAILABLE", "CITED", "회사 공시", "Company statement")), "caption accounting basis or company-statement basis is not concrete")


def gate_missing_glyphs(messages: list[str]) -> None:
    failures = [message for message in messages if re.search(r"glyph.*missing|missing.?glyph", message, re.I)]
    require(not failures, "missing glyph warnings: " + "; ".join(failures))


def gate_pdf_bold_runs(path: Path, texts: list[str]) -> dict[str, int]:
    import pdfplumber

    locations = {}
    with pdfplumber.open(path) as document:
        for target in texts:
            normalized = re.sub(r"\s+", "", target)
            for number, page in enumerate(document.pages, 1):
                chars = [(letter, char["fontname"]) for char in page.chars
                         for letter in char["text"] if not letter.isspace()]
                joined = "".join(letter for letter, _ in chars)
                start = joined.find(normalized)
                if start >= 0 and all("bold" in font.lower() for _, font in chars[start:start + len(normalized)]):
                    locations[target] = number
                    break
            require(target in locations, f"PDF text lacks a real Bold font run: {target}")
    return locations


def gate_r22_presentation(outputs, chart_manifests, narrative, manifest) -> dict:
    from .presentation import FIGURE_ORDER
    from .render import _locale

    audit = {}
    for locale in ("ko", "en"):
        text = outputs[f"md_{locale}"].read_text(encoding="utf-8")
        images = re.findall(r"!\[([^]]+)]", text)
        require(tuple(images) == FIGURE_ORDER, f"R22 figure order/duplication: {locale}")
        require("SRC-" not in text, f"R22 internal source ID exposed: {locale}")
        for number, chart_id in enumerate(images, 1):
            chart = chart_manifests[locale][chart_id]
            require(re.search(r"\d+", chart["caption"]["number"]).group() == str(number), "R22 figure numbering mismatch")
            gate_missing_glyphs(chart["render_warnings"])
            layout = chart["layout"]
            width, height = layout["figure_size"]
            for item in layout["text_bounds"]:
                x, y, w, h = item["bbox"]
                require(min(x, y) >= -1 and x + w <= width + 1 and y + h <= height + 1, "R22 text outside bitmap")
        config = _locale(locale)
        targets = [config["ed1"]["title"], config["narrative_ui"]["bull"], config["narrative_ui"]["bear"]]
        targets += [item["claim"] for side in ("bull", "bear") for item in narrative["thesis"][locale][side]]
        targets += [item["title"] for item in narrative["risks"][locale]]
        targets = [PLACEHOLDER_RE.sub(lambda match: manifest.fact(match.group(1)).display[locale], target) for target in targets]
        audit[locale] = gate_pdf_bold_runs(outputs[f"pdf_{locale}"], targets)
    return audit


def gate_r23_table_layout(html_text: str, base_url: str | None = None) -> dict:
    """Check actual paginated header text and appendix table geometry."""
    from weasyprint import HTML

    document = HTML(string=html_text, base_url=base_url).render()
    headers, appendix_tables = 0, []
    for number, page in enumerate(document.pages, 1):
        for box in page._page_box.descendants():
            if type(box).__name__ == "TableCellBox" and box.element_tag == "th":
                headers += 1
                for child in box.descendants():
                    if type(child).__name__ == "TextBox":
                        require(child.position_x >= box.position_x - .1 and child.position_x + child.width <= box.position_x + box.border_width() + .1,
                                f"R23 header outside its own cell: page {number}: {child.text}")
            if type(box).__name__ == "TableBox" and box.element is not None and box.element.get("class") in {"appendix-rules", "post-print-changes"}:
                right = (box.position_x + box.border_width()) * .75
                require(right <= 549.9, f"R23 appendix table outside content: page {number}: {right:.2f}pt")
                appendix_tables.append({"page": number, "class": box.element.get("class"), "right_pt": right})
    require(headers > 0, "R23 table header audit found no headers")
    return {"headers": headers, "appendix_tables": appendix_tables}


def gate_r23_presentation(outputs, chart_manifests) -> dict:
    audit = {}
    for locale in ("ko", "en"):
        html_path = outputs[f"html_{locale}"]
        audit[locale] = gate_r23_table_layout(html_path.read_text(encoding="utf-8"), str(html_path.parent))
        for chart in chart_manifests[locale].values():
            require(chart["layout"]["minimum_legend_tick_pdf_pt"] >= 7, "R23 legend/tick below PDF 7pt")
    return audit


def gate_g17_labels(text: str, locale: str) -> None:
    require(locale in {"ko", "en"}, f"unsupported locale: {locale}")
    week_tokens = ("14주", "53주") if locale == "ko" else ("14 weeks", "53 weeks")
    for token in ("FY2026Q4", *week_tokens, "PREREG_A", "A-8K", "RLE", "UNAVAILABLE"):
        require(token in text, f"required label missing: {token}")


def gate_g17_net_cash_label(configs: dict[str, dict[str, Any]]) -> None:
    expected = {
        "ko": "순현금(SCA 예치금 미조정)",
        "en": "Net cash (not adjusted for SCA deposits)",
    }
    forbidden = {"ko": "조정 전 순현금", "en": "Unadjusted net cash"}
    require(set(configs) == {"ko", "en"}, "both locale label configs are required")
    for locale, config in configs.items():
        chart = config["charts"]["08_cash_flow_capex_net_cash"]
        text = " ".join([chart["title"], *chart["series"]])
        require(expected[locale] in text, f"corrected net-cash label missing: {locale}")
        require(forbidden[locale] not in text, f"obsolete net-cash label remains: {locale}")


def gate_g17b_consensus(reference_log: list[dict[str, str]], manifest: Manifest) -> None:
    unavailable = {
        fact.fact_id
        for fact in manifest.facts.values()
        if fact.status == "UNAVAILABLE_FOR_COMPARISON"
    }
    for entry in reference_log:
        if entry["fact_id"] in unavailable:
            require(entry.get("role") == "display", f"unavailable consensus used as comparison: {entry['position']}")
    for fact in manifest.facts.values():
        inputs = set((fact.lineage or {}).get("inputs", []))
        require(not (inputs & unavailable), f"calculation consumes unavailable consensus: {fact.fact_id}")


def gate_g18_hygiene(paths: Iterable[Path]) -> None:
    for path in paths:
        payload = path.read_bytes()
        require(b"\x00" not in payload, f"NUL byte: {path}")
        if path.suffix.lower() in {".py", ".yaml", ".yml", ".json", ".md", ".html", ".mjs"}:
            require(b"\r\n" not in payload and b"\r" not in payload, f"non-LF newline: {path}")
            text = payload.decode("utf-8")
            require(not any(line.endswith((" ", "\t")) for line in text.splitlines()), f"trailing whitespace: {path}")


def gate_g19_bridge(manifest: Manifest) -> None:
    manifest.validate()
    bridge = [fact for key, fact in manifest.facts.items() if key.startswith("bridge.nongaap_fixed_027")]
    require(len(bridge) == 1, "exactly one fixed +$0.27 bridge fact is required")


def gate_g20_audit(reader: EvidenceReader) -> None:
    payload_1 = reader.audit_payload()
    payload_2 = reader.audit_payload()
    require(payload_1 == payload_2, "audit manifest is non-deterministic")
    audit = json.loads(payload_1)
    reserved = {item["path"] for phase in ("E2-B", "E2-C") for item in load_pins()[phase]}
    require(not reserved & {row["path"] for row in audit["reads"]}, "E2-A audit contains reserved input")


def qa_pdf(
    path: Path,
    locale: str,
    expected_total: int | None = None,
    toc_entries: list[dict[str, str]] | None = None,
    conflict_short: str | None = None,
) -> None:
    payload = path.read_bytes()
    require(payload.rstrip().endswith(b"%%EOF"), f"PDF lacks EOF: {path}")
    pdf = PdfReader(path)
    require(len(pdf.pages) > 0, f"empty PDF: {path}")
    advice = "본 문서는 투자 자문이 아닙니다." if locale == "ko" else "This document is not investment advice."
    independent = "Micron 미승인" if locale == "ko" else "not approved by Micron"
    texts = [page.extract_text() or "" for page in pdf.pages]
    qa_pdf_texts(texts, advice, path, independent)
    conflict_short = conflict_short or (
        "작성자 MU 주식 미보유 - 발행 직전 재확인."
        if locale == "ko"
        else "Author does not hold MU shares - reconfirmed before publication."
    )
    qa_pdf_page_metadata(texts, expected_total or len(texts), conflict_short, toc_entries)
    font_names = []
    embedded = []
    for page in pdf.pages:
        fonts = page.get("/Resources", {}).get("/Font", {})
        for font_ref in fonts.values():
            font = font_ref.get_object()
            font_names.append(str(font.get("/BaseFont", "")))
            descendants = font.get("/DescendantFonts", [])
            for descendant in descendants:
                descriptor = descendant.get_object().get("/FontDescriptor")
                if descriptor:
                    descriptor = descriptor.get_object()
                    embedded.append(any(key in descriptor for key in ("/FontFile", "/FontFile2", "/FontFile3")))
    require(any("Noto" in name for name in font_names), f"Noto font not present in {path}")
    require(bool(embedded) and all(embedded), f"font is not embedded in {path}")


def qa_pdf_texts(texts: list[str], advice: str, path: Path | str = "PDF", independent: str | None = None) -> None:
    for text in texts:
        require("file:///" not in text, f"browser file footer in {path}")
        require(advice in text, f"disclaimer absent from PDF page in {path}")
        if independent:
            require(independent in text, f"independence disclaimer absent from PDF page in {path}")


def qa_pdf_page_metadata(
    texts: list[str],
    manifest_total: int,
    conflict_short: str,
    toc_entries: list[dict[str, str]] | None = None,
) -> None:
    total = len(texts)
    require(total == manifest_total, "manifest total_pages differs from PDF page count")
    normalized = [" ".join(text.split()) for text in texts]
    for current, text in enumerate(normalized, start=1):
        require(re.search(rf"(?<!\d){current}\s*/\s*{total}(?!\d)", text) is not None, f"PDF page footer number mismatch at page {current}")
        require(conflict_short in text, f"conflict_short absent from PDF page {current}")
    if toc_entries:
        labels = [f"{index}. {entry['title']}" for index, entry in enumerate(toc_entries, start=1)]
        toc_candidates = [
            index
            for index, text in enumerate(normalized)
            if re.search(r"\bContents\b|목차", text)
            and sum(label in text for label in labels) >= len(labels) - 1
        ]
        require(len(toc_candidates) == 1, "PDF TOC page could not be identified uniquely")
        toc_index = toc_candidates[0]
        toc_text = normalized[toc_index]
        for index, entry in enumerate(toc_entries, start=1):
            label = f"{index}. {entry['title']}"
            start = toc_text.find(label)
            require(start >= 0, f"TOC entry missing: {label}")
            next_start = len(toc_text)
            if index < len(toc_entries):
                next_label = f"{index + 1}. {toc_entries[index]['title']}"
                candidate = toc_text.find(next_label, start + len(label))
                if candidate >= 0:
                    next_start = candidate
            segment = toc_text[start + len(label) : next_start]
            page_numbers = re.findall(r"\d+", segment)
            require(bool(page_numbers), f"TOC page number missing: {label}")
            listed_page = int(page_numbers[0])
            actual_pages = [
                page_index + 1
                for page_index, text in enumerate(normalized)
                if page_index != toc_index and label in text
            ]
            require(bool(actual_pages) and listed_page == actual_pages[0], f"TOC target page mismatch: {label}")


def qa_html(path: Path) -> None:
    text = path.read_text(encoding="utf-8")
    require(path.stat().st_size < 5 * 1024 * 1024, "HTML is 5MB or larger")
    require(not re.search(r"<(?:script|link|img)[^>]+(?:src|href)=[\"']https?://", text, re.I), "external HTML dependency")
    anchors = set(re.findall(r"\bid=[\"']([^\"']+)", text))
    targets = set(re.findall(r"\bhref=[\"']#([^\"']+)", text))
    require(targets <= anchors, "HTML contains an unresolved local anchor")


def qa_markdown(path: Path) -> None:
    text = path.read_text(encoding="utf-8")
    for line in text.splitlines():
        if line.startswith("|"):
            require(line.count("|") >= 3, f"malformed Markdown table row: {line}")
    anchors = set(re.findall(r'<a id="([^"]+)"></a>', text))
    targets = set(re.findall(r"\]\(#([^)]+)\)", text))
    require(targets <= anchors, "Markdown contains an unresolved local anchor")


def qa_xlsx(path: Path) -> None:
    workbook = load_workbook(path, data_only=False, read_only=True)
    try:
        rows_by_sheet: dict[str, list[str]] = {}
        for sheet in workbook.worksheets:
            flattened = []
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        flattened.append(str(cell.value))
            rows_by_sheet[sheet.title] = flattened
    finally:
        workbook.close()
    qa_xlsx_rows(rows_by_sheet)


def qa_xlsx_rows(rows_by_sheet: dict[str, list[str]]) -> None:
    errors = {"#REF!", "#DIV/0!", "#VALUE!", "#N/A"}
    for title, flattened in rows_by_sheet.items():
        for value in flattened:
            require(value not in errors, f"XLSX error in {title}")
        joined = " ".join(flattened).lower()
        for token in ("source", "as of", "unit"):
            require(token in joined, f"{title} lacks {token} metadata")


def gate_g21_formats(outputs: dict[str, Path], dryrun: bool = False) -> None:
    manifest = json.loads(outputs["manifest"].read_text(encoding="utf-8"))
    page_counts = manifest["metadata"]["document"]["total_pages"]
    for locale in ("ko", "en"):
        strings = yaml.safe_load((PACKAGE_DIR / "i18n" / f"{locale}.yaml").read_text(encoding="utf-8"))
        conflict_short = DRYRUN_CONFLICT[locale] if dryrun else None
        qa_pdf(outputs[f"pdf_{locale}"], locale, page_counts[locale], strings["sections"], conflict_short)
        qa_html(outputs[f"html_{locale}"])
        qa_markdown(outputs[f"md_{locale}"])
    qa_xlsx(outputs["xlsx"])
