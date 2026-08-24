"""Sheet: Historical Band — LTM P/B·P/S point-in-time band (reporting-only).

Only built when band reports are explicitly passed (--band --excel); the
default workbook is untouched, keeping --excel-only output semantically
identical (§7.1 #8). Not a valuation input.
"""

from __future__ import annotations

from openpyxl.utils import get_column_letter

from engine.multiple_band import MIN_OBS_FOR_BAND, band_verdict, percentile_rank
from schemas.point_in_time import HistoricalBand
from ._ctx import Ctx
from ..excel_styles import (
    GREEN_FILL,
    MULT_FMT,
    SECTION_FONT,
    TITLE_FONT,
    YELLOW_FILL,
    style_header_row,
    write_cell,
)


def sheet_historical_band(
    ctx: Ctx, bands: list[HistoricalBand], current: dict[str, float] | None = None
) -> None:
    current = current or {}
    ws = ctx.wb.create_sheet("Historical Band")
    ws.sheet_properties.tabColor = "2C6E49"
    for c, w in enumerate([14, 10, 10, 10, 10, 10, 12, 26], 1):
        ws.column_dimensions[get_column_letter(c)].width = w

    write_cell(
        ws, 1, 1, "역사적 배수 밴드 (LTM, point-in-time — 참고용)", font=TITLE_FONT
    )
    r = 3

    # ── Summary table ──
    for c, h in enumerate(
        ["지표", "N", "min", "p25", "median", "p75", "max", "현재 위치"], 1
    ):
        write_cell(ws, r, c, h)
    style_header_row(ws, r, 8)
    for band in bands:
        r += 1
        write_cell(ws, r, 1, f"{band.label} (LTM)")
        write_cell(ws, r, 2, band.n_obs)
        for c, v in enumerate(
            [band.band_min, band.p25, band.median, band.p75, band.band_max], 3
        ):
            if v is None:
                write_cell(ws, r, c, "—")
            else:
                write_cell(ws, r, c, round(v, 3), fmt=MULT_FMT)
        cur = current.get(band.label)
        verdict = band_verdict(band, cur)
        if band.n_obs < MIN_OBS_FOR_BAND:
            write_cell(ws, r, 8, f"이력 부족(N={band.n_obs})", fill=YELLOW_FILL)
        elif cur is not None:
            rank = percentile_rank(band, cur)
            rank_txt = f" ({rank * 100:.0f}%ile)" if rank is not None else ""
            write_cell(ws, r, 8, f"{cur:.2f}x — {verdict}{rank_txt}")
        else:
            write_cell(ws, r, 8, verdict)

    # ── Per-year observations ──
    for band in bands:
        r += 2
        write_cell(ws, r, 1, f"{band.label} 연도별 관측", font=SECTION_FONT)
        r += 1
        for c, h in enumerate(
            [
                "FY",
                "접수일(t)",
                "가격일",
                "raw close",
                "유통주식수",
                "분모(백만원)",
                "배수",
                "rcept_no",
            ],
            1,
        ):
            write_cell(ws, r, c, h)
        style_header_row(ws, r, 8)
        for o in band.observations:
            r += 1
            write_cell(ws, r, 1, o.fiscal_year)
            write_cell(ws, r, 2, o.t.isoformat())
            write_cell(ws, r, 3, o.price_date.isoformat())
            write_cell(ws, r, 4, o.price_close_raw_krw)
            write_cell(ws, r, 5, o.shares_outstanding)
            write_cell(ws, r, 6, o.denominator_mkrw)
            write_cell(ws, r, 7, round(o.multiple, 4), fmt=MULT_FMT, fill=GREEN_FILL)
            write_cell(ws, r, 8, o.rcept_no)
        for exc in band.excluded:
            r += 1
            reason = (
                exc.price_reason.value
                if exc.price_reason
                else (
                    f"missing:{','.join(exc.missing_accounts)}"
                    if exc.missing_accounts
                    else exc.note
                )
            )
            write_cell(ws, r, 1, exc.fiscal_year)
            write_cell(ws, r, 2, f"제외 — {reason}", fill=YELLOW_FILL)

    r += 2
    write_cell(
        ws,
        r,
        1,
        "* 최초 공시값(basis=original)·접수일 raw close(auto_adjust=False)·당시 유통주식수 기준. "
        "보간·소급·재작성값 소비 없음. 밸류에이션 입력 아님 (참고용).",
    )
