"""Render locale editions and a shared workbook from one manifest."""

from __future__ import annotations

import json
import html as html_module
import os
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pypdf import PdfReader
from weasyprint import HTML

from .charts import dryrun_specs, e2b_specs, fixture_specs, render_charts
from .facts import Manifest
from .inputs import ConflictConfirmation, PACKAGE_DIR, atomic_write, load_conflict_confirmation
from .narrative import appendix_markdown, narrative_section_markdown
from .rle import ReportLayerAssumptions
from .theme import COLORS, FONT_STACK

PLACEHOLDER = re.compile(r"\{\{fact:([^}]+)}}")
DISCLAIMER_KEYS = ("disclaimer.not_advice", "disclaimer.third_party")
FIXTURE_CONFLICT = Path("forecast/tests/fixtures/mu_report/conflict_confirmation_ok.yaml")
FIXTURE_NOW_KST = datetime(2099, 1, 1, 12, 30, tzinfo=timezone(timedelta(hours=9)))


def _locale(locale: str) -> dict:
    return yaml.safe_load((PACKAGE_DIR / "i18n" / f"{locale}.yaml").read_text(encoding="utf-8"))


def _reference(fact_id: str, position: str, kind: str, role: str, manifest: Manifest) -> dict:
    fact = manifest.fact(fact_id)
    return {
        "position": position,
        "kind": kind,
        "role": role,
        "fact_id": fact_id,
        "raw_value": fact.raw_value,
        "unit": fact.unit,
        "period": fact.period,
        "basis": fact.basis,
        "label": fact.label,
        "status": fact.status,
    }


def _replace(
    text: str,
    manifest: Manifest,
    locale: str,
    position: str,
    kind: str,
    refs: list[dict],
    role: str = "display",
) -> str:
    occurrence = 0

    def replacement(match: re.Match[str]) -> str:
        nonlocal occurrence
        fact_id = match.group(1)
        refs.append(_reference(fact_id, f"{position}:{occurrence}", kind, role, manifest))
        occurrence += 1
        return manifest.fact(fact_id).display[locale]
    return PLACEHOLDER.sub(replacement, text)


def _conflict_text(strings: dict, record: ConflictConfirmation, key: str) -> str:
    return strings["disclaimer"][key].replace(
        "{{conflict.confirmed_at_kst}}",
        record.confirmed_at_kst.isoformat(),
    )


def _toc_entries(locale: str) -> list[dict[str, str]]:
    return [dict(section) for section in _locale(locale)["sections"]]


def _fact_cell(manifest: Manifest, fact_id: str, locale: str, position: str, refs: list[dict]) -> str:
    if fact_id not in manifest.facts:
        return "UNAVAILABLE"
    return _replace(f"{{{{fact:{fact_id}}}}}", manifest, locale, position, "table_cell", refs)


def _financial_tables(manifest: Manifest, locale: str, strings: dict, refs: list[dict]) -> list[str]:
    labels = {
        "revenue": ("매출", "Revenue"), "cogs": ("매출원가", "Cost of revenue"),
        "gross_profit": ("매출총이익", "Gross profit"), "r_and_d": ("연구개발비", "R&D"),
        "sg_and_a": ("판매관리비", "SG&A"), "restructuring": ("구조조정", "Restructuring"),
        "other_operating": ("기타 영업", "Other operating"), "operating_income": ("영업이익", "Operating income"),
        "interest_income": ("이자수익", "Interest income"), "interest_expense": ("이자비용", "Interest expense"),
        "other_nonoperating": ("기타 영업외", "Other non-operating"), "pretax_income": ("세전이익", "Pretax income"),
        "tax": ("법인세", "Income tax"), "equity_method": ("지분법", "Equity-method result"),
        "net_income": ("순이익", "Net income"), "diluted_shares": ("희석주식수", "Diluted shares"), "diluted_eps": ("희석 EPS", "Diluted EPS"),
        "cash": ("현금", "Cash"), "short_term_investments": ("단기투자", "Short-term investments"),
        "receivables": ("매출채권", "Receivables"), "inventories": ("재고", "Inventories"),
        "total_current_assets": ("유동자산", "Current assets"), "long_term_investments": ("장기투자", "Long-term investments"),
        "ppe": ("유형자산", "Property, plant and equipment"), "total_assets": ("자산총계", "Total assets"),
        "accounts_payable_accrued": ("매입채무·미지급", "Accounts payable and accrued expenses"),
        "current_debt": ("유동성 차입금", "Current debt"), "total_current_liabilities": ("유동부채", "Current liabilities"),
        "long_term_debt": ("장기차입금", "Long-term debt"), "total_liabilities": ("부채총계", "Total liabilities"),
        "total_equity": ("자본총계", "Total equity"), "d_and_a": ("감가상각", "Depreciation and amortization"),
        "sbc": ("주식보상", "Stock-based compensation"), "receivables_change": ("매출채권 변동", "Change in receivables"),
        "inventories_change": ("재고 변동", "Change in inventories"), "accounts_payable_change": ("매입채무 변동", "Change in accounts payable"),
        "other_current_liabilities_change": ("기타 유동부채 변동", "Change in other current liabilities"),
        "operating_cash_flow": ("영업현금흐름", "Operating cash flow"), "ppe_expenditures": ("유형자산 취득", "PPE expenditures"),
        "government_incentives": ("정부보조금", "Government incentives"), "debt_repayments": ("차입금 상환", "Debt repayments"),
        "debt_issuance": ("차입금 조달", "Debt issuance"), "dividends": ("배당", "Dividends"),
        "ppe_disposal_proceeds": ("유형자산 처분", "PPE disposal proceeds"), "investing_cash_flow": ("투자현금흐름", "Investing cash flow"),
        "financing_cash_flow": ("재무현금흐름", "Financing cash flow"), "fx_effect": ("환율 효과", "FX effect"),
        "cash_change": ("현금 증감", "Change in cash"), "share_repurchases": ("자사주", "Share repurchases"),
        "gross_margin": ("매출총이익률", "Gross margin"), "operating_margin": ("영업이익률", "Operating margin"),
        "etr": ("유효세율", "Effective tax rate"), "inventory_days": ("재고일수", "Inventory days"),
        "net_cash_unadjusted": ("순현금(SCA 예치금 미조정)", "Net cash (not adjusted for SCA deposits)"), "fcf_adjusted": ("조정 FCF", "Adjusted FCF"),
        "net_capex": ("순설비투자", "Net capex"), "roe": ("ROE", "ROE"),
        "net_capex_revenue": ("순설비투자/매출", "Net capex / revenue"), "opex": ("영업비용", "Operating expenses"),
        "below_op": ("영업이익 아래 항목", "Below-operating items"), "working_capital_investment": ("운전자본 투자", "Working-capital investment"),
    }
    locale_index = 0 if locale == "ko" else 1
    columns = strings["financial_tables"]["columns"]
    not_preregistered = strings["financial_tables"]["not_preregistered"]
    out: list[str] = []

    def cell(fact_id: str, position: str, fallback: str = "UNAVAILABLE") -> str:
        return _fact_cell(manifest, fact_id, locale, position, refs) if fact_id in manifest.facts else fallback

    def income_values(metric: str) -> list[str]:
        history = [cell(f"is.{metric}.FY{year}A", f"table:is:{metric}:{year}", "—") for year in (2023, 2024, 2025)]
        prereg_id = f"is.{metric}.FY2026E.PREREG_A"
        prereg = cell(prereg_id, f"table:is:{metric}:prereg", not_preregistered) if metric in {"revenue", "net_income", "diluted_eps"} else not_preregistered
        actual = cell(f"is.{metric}.FY2026A.A-8K", f"table:is:{metric}:actual")
        rle_metric = "pretax" if metric == "pretax_income" else metric
        fy27 = cell(f"is.{rle_metric}.FY2027E.base", f"table:is:{metric}:fy27") if metric not in {"r_and_d", "sg_and_a", "restructuring", "other_operating", "interest_income", "interest_expense", "other_nonoperating", "equity_method"} else "UNAVAILABLE"
        fy28 = cell(f"is.{rle_metric}.FY2028E.base", f"table:is:{metric}:fy28") if metric not in {"r_and_d", "sg_and_a", "restructuring", "other_operating", "interest_income", "interest_expense", "other_nonoperating", "equity_method"} else "UNAVAILABLE"
        return [*history, prereg, actual, fy27, fy28]

    def balance_values(metric: str) -> list[str]:
        history = [cell(f"bs.{metric}.FY{year}A", f"table:bs:{metric}:{year}", "—") for year in (2023, 2024, 2025)]
        actual = cell(f"bs.{metric}.FY2026.A-8K", f"table:bs:{metric}:actual")
        return [*history, not_preregistered, actual, "UNAVAILABLE", "UNAVAILABLE"]

    def cash_values(metric: str) -> list[str]:
        history = [cell(f"cf.{metric}.FY{year}A", f"table:cf:{metric}:{year}", "—") for year in (2023, 2024, 2025)]
        actual = cell(f"cf.{metric}.FY2026A", f"table:cf:{metric}:actual")
        rle_map = {
            "net_income": "net_income", "d_and_a": "d_and_a", "sbc": "sbc", "operating_cash_flow": "operating_cash_flow",
            "net_capex": "net_capex", "fcf_adjusted": "fcf_adjusted", "dividends": "dividends",
            "working_capital_investment": "working_capital_investment",
        }
        rle_metric = rle_map.get(metric)
        rle_family = "is" if metric == "net_income" else "cf"
        fy27 = cell(f"{rle_family}.{rle_metric}.FY2027E.base", f"table:cf:{metric}:fy27") if rle_metric else "UNAVAILABLE"
        fy28 = cell(f"{rle_family}.{rle_metric}.FY2028E.base", f"table:cf:{metric}:fy28") if rle_metric else "UNAVAILABLE"
        return [*history, not_preregistered, actual, fy27, fy28]

    def add_table(title: str, family: str, metrics: tuple[str, ...]) -> None:
        out.extend([f"### {title}", "", "| " + " | ".join(columns) + " |", "|---|---:|---:|---:|---:|---:|---:|---:|"])
        for metric in metrics:
            values = income_values(metric) if family == "is" else balance_values(metric) if family == "bs" else cash_values(metric)
            out.append(f"| {labels[metric][locale_index]} | " + " | ".join(values) + " |")
        out.append("")
    add_table(strings["financial_tables"]["income_title"], "is", (
        "revenue", "cogs", "gross_profit", "opex", "r_and_d", "sg_and_a", "restructuring", "other_operating",
        "operating_income", "interest_income", "interest_expense", "other_nonoperating", "pretax_income",
        "below_op", "tax", "equity_method", "net_income", "diluted_shares", "diluted_eps",
    ))
    add_table(strings["financial_tables"]["balance_title"], "bs", (
        "cash", "short_term_investments", "receivables", "inventories", "total_current_assets", "long_term_investments",
        "ppe", "total_assets", "accounts_payable_accrued", "current_debt", "total_current_liabilities",
        "long_term_debt", "total_liabilities", "total_equity",
    ))
    add_table(strings["financial_tables"]["cashflow_title"], "cf", (
        "net_income", "d_and_a", "sbc", "working_capital_investment", "receivables_change", "inventories_change", "accounts_payable_change",
        "other_current_liabilities_change", "operating_cash_flow", "ppe_expenditures", "government_incentives",
        "ppe_disposal_proceeds", "net_capex", "fcf_adjusted", "investing_cash_flow", "debt_repayments", "debt_issuance", "share_repurchases",
        "dividends", "financing_cash_flow", "fx_effect", "cash_change",
    ))
    out.extend([f"### {strings['financial_tables']['ratio_title']}", "", "| " + " | ".join(columns) + " |", "|---|---:|---:|---:|---:|---:|---:|---:|"])
    for metric, family in (("gross_margin", "ratio"), ("operating_margin", "ratio"), ("etr", "ratio"), ("roe", "ratio"), ("inventory_days", "ratio"), ("net_capex_revenue", "ratio"), ("net_cash_unadjusted", "bs")):
        history = [cell(f"{family}.{metric}.FY{year}.A", f"table:ratio:{metric}:{year}", "—") for year in (2023, 2024, 2025)]
        if metric == "net_cash_unadjusted":
            actual = cell("bs.net_cash_unadjusted.FY2026.A-8K", "table:ratio:net_cash:actual")
            fy27 = cell("bs.net_cash_unadjusted.FY2027E.base", "table:ratio:net_cash:fy27")
            fy28 = cell("bs.net_cash_unadjusted.FY2028E.base", "table:ratio:net_cash:fy28")
        else:
            actual = cell(f"ratio.{metric}.FY2026.A-8K", f"table:ratio:{metric}:actual")
            fy27 = cell(f"ratio.{metric}.FY2027E.base", f"table:ratio:{metric}:fy27") if metric in {"gross_margin", "operating_margin", "etr", "net_capex_revenue"} else "UNAVAILABLE"
            fy28 = cell(f"ratio.{metric}.FY2028E.base", f"table:ratio:{metric}:fy28") if metric in {"gross_margin", "operating_margin", "etr", "net_capex_revenue"} else "UNAVAILABLE"
        out.append(f"| {labels[metric][locale_index]} | " + " | ".join([*history, not_preregistered, actual, fy27, fy28]) + " |")
    out.extend(["", strings["financial_tables"]["not_preregistered_note"], ""])
    return out


def _scored_table(manifest: Manifest, locale: str, strings: dict, refs: list[dict]) -> list[str]:
    ui = strings["scored"]
    rows = []
    definitions = (
        (ui["revenue"], "scored.revenue_actual.FQ4FY26", "scored.revenue_forecast.FQ4FY26", "scored.revenue_error.FQ4FY26", "scored.revenue_ape.FQ4FY26"),
        (ui["gaap_eps"], "scored.eps_actual.FQ4FY26", "scored.eps_forecast.FQ4FY26", "scored.eps_error.FQ4FY26", "scored.eps_ape.FQ4FY26"),
        (ui["nongaap_eps"], "scored.nongaap_eps_actual.FQ4FY26", "scored.nongaap_eps_forecast.FQ4FY26", "scored.nongaap_eps_error.FQ4FY26", "scored.nongaap_eps_ape.FQ4FY26"),
        (ui["gaap_opex"], "scored.opex_actual.FQ4FY26", "scored.opex_forecast.FQ4FY26", "scored.opex_error.FQ4FY26", "scored.opex_ape.FQ4FY26"),
    )
    for index, (label, actual, forecast, error, ape) in enumerate(definitions):
        values = [_fact_cell(manifest, fact_id, locale, f"table:scored:{index}:{column}", refs) for column, fact_id in enumerate((actual, forecast, error, ape))]
        rows.append(f"| {label} | " + " | ".join(values) + " |")
    summary = [
        f"{ui['labels']}: " + _fact_cell(manifest, "scored.labels_hit.FQ4FY26", locale, "text:scored:labels", refs) + f"/4 {ui['labels_suffix']}",
        f"{ui['bands']}: " + _fact_cell(manifest, "scored.band_coverage.FQ4FY26", locale, "text:scored:bands", refs)
        + " · " + _fact_cell(manifest, "scored.revenue_band_low.FQ4FY26", locale, "text:scored:revenue_band_low", refs)
        + "–" + _fact_cell(manifest, "scored.revenue_band_high.FQ4FY26", locale, "text:scored:revenue_band_high", refs)
        + " · " + _fact_cell(manifest, "scored.eps_band_low.FQ4FY26", locale, "text:scored:eps_band_low", refs)
        + "–" + _fact_cell(manifest, "scored.eps_band_high.FQ4FY26", locale, "text:scored:eps_band_high", refs),
        f"{ui['lever']}: " + " / ".join(
            _fact_cell(manifest, f"scored.lever.{name}.FQ4FY26", locale, f"text:scored:lever:{name}", refs)
            for name in ("revenue", "op_margin", "op_to_ni", "shares")
        ) + " = " + _fact_cell(manifest, "scored.lever.total.FQ4FY26", locale, "text:scored:lever:total", refs),
        f"{ui['sf7']}: " + _fact_cell(manifest, "scored.sf7.FQ1FY27", locale, "text:scored:sf7", refs),
    ]
    return [
        f"### {ui['title']}", "",
        f"| {ui['metric']} | {ui['actual']} | {ui['forecast']} | {ui['error']} | {ui['ape']} |",
        "|---|---:|---:|---:|---:|", *rows, "", *summary, "",
    ]


def _guidance_table(manifest: Manifest, locale: str, strings: dict, refs: list[dict]) -> list[str]:
    ui = strings["guidance_table"]
    rows = []
    definitions = (
        ("revenue", "guidance.revenue_gaap.FQ1FY27", "guidance.revenue_nongaap.FQ1FY27"),
        ("gm", "guidance.gm_gaap.FQ1FY27", "guidance.gm_nongaap.FQ1FY27"),
        ("opex", "guidance.opex_gaap.FQ1FY27", "guidance.opex_nongaap.FQ1FY27"),
        ("eps", "guidance.eps_gaap.FQ1FY27", "guidance.eps_nongaap.FQ1FY27"),
        ("shares", "guidance.diluted_shares.FQ1FY27", "guidance.diluted_shares.FQ1FY27"),
    )
    for index, ((name, gaap_id, nongaap_id), label, basis) in enumerate(zip(definitions, ui["rows"], ui["basis"], strict=True)):
        gaap = _fact_cell(manifest, gaap_id, locale, f"table:guidance:{name}:gaap", refs)
        nongaap = _fact_cell(manifest, nongaap_id, locale, f"table:guidance:{name}:nongaap", refs)
        if name == "revenue":
            tolerance = _fact_cell(manifest, "guidance.revenue_tolerance.FQ1FY27", locale, "table:guidance:revenue:tolerance", refs)
            gaap = f"{gaap} ± {tolerance}"
            nongaap = f"{nongaap} ± {tolerance}"
        if name == "eps":
            tolerance = _fact_cell(manifest, "guidance.eps_tolerance.FQ1FY27", locale, "table:guidance:eps:tolerance", refs)
            gaap = f"{gaap} ± {tolerance}"
            nongaap = f"{nongaap} ± {tolerance}"
        rows.append(f"| {label} | {gaap} | {nongaap} | {basis} |")
    return [f"### {ui['title']}", "", "| " + " | ".join(ui["headers"]) + " |", "|---|---:|---:|---|", *rows, ""]


def _margin_bridge_table(manifest: Manifest, locale: str, strings: dict, refs: list[dict]) -> list[str]:
    ui = strings["margin_bridge_table"]
    definitions = (
        (ui["periods"][0], ui["metrics"][0], "bridge.gm.FY2026Q3.A", "bridge.gm.FY2026Q4.A", "bridge.gm.delta.FY2026Q3_to_Q4"),
        (ui["periods"][0], ui["metrics"][1], "bridge.opex.FY2026Q3.A", "bridge.opex.FY2026Q4.A", "bridge.opex.delta.FY2026Q3_to_Q4"),
        (ui["periods"][0], ui["metrics"][2], "bridge.opex_ratio.FY2026Q3.A", "bridge.opex_ratio.FY2026Q4.A", "bridge.opex_ratio.contribution.FY2026Q3_to_Q4"),
        (ui["periods"][1], ui["metrics"][0], "bridge.gm.FY2026A", "bridge.gm.FY2027E.base", "bridge.gm.delta.FY2026A_to_FY2027E.base"),
        (ui["periods"][1], ui["metrics"][1], "bridge.opex.FY2026A", "bridge.opex.FY2027E.base", "bridge.opex.delta.FY2026A_to_FY2027E.base"),
        (ui["periods"][1], ui["metrics"][2], "bridge.opex_ratio.FY2026A", "bridge.opex_ratio.FY2027E.base", "bridge.opex_ratio.contribution.FY2026A_to_FY2027E.base"),
    )
    rows = []
    for index, (period, metric, start_id, end_id, change_id) in enumerate(definitions):
        values = [_fact_cell(manifest, fact_id, locale, f"table:margin_bridge:{index}:{column}", refs) for column, fact_id in enumerate((start_id, end_id, change_id))]
        rows.append(f"| {period} | {metric} | " + " | ".join(values) + " |")
    return [f"### {ui['title']}", "", "| " + " | ".join(ui["headers"]) + " |", "|---|---|---:|---:|---:|", *rows, "", ui["note"], ""]


def _valuation_table(manifest: Manifest, locale: str, strings: dict, refs: list[dict]) -> list[str]:
    ui = strings["valuation"]
    periods = ["FY2025A", "FY2026A", "FY2027E.bear", "FY2027E.base", "FY2027E.bull", "FY2028E.bear", "FY2028E.base", "FY2028E.bull"]
    rows = []
    for index, period in enumerate(periods):
        values = [
            _fact_cell(manifest, f"val.{metric}.{period}", locale, f"table:valuation:{index}:{metric}", refs)
            for metric in ("pe", "ev_ebitda", "ev_sales", "fcf_yield")
        ]
        rows.append(f"| {period} | " + " | ".join(values) + " |")
    return [
        f"### {ui['title']}", "",
        f"| {ui['period']} | {ui['pe']} | {ui['ev_ebitda']} | {ui['ev_sales']} | {ui['fcf_yield']} |",
        "|---|---:|---:|---:|---:|", *rows, "",
    ]


def render_markdown(
    manifest: Manifest,
    locale: str,
    conflict: ConflictConfirmation | None = None,
    dryrun: bool = False,
    chart_manifest: dict[str, dict] | None = None,
    narrative: dict | None = None,
    assumptions: ReportLayerAssumptions | None = None,
    asset_prefix: str = "assets",
    render_date: str | None = None,
) -> tuple[str, list[dict]]:
    strings = _locale(locale)
    ed1 = narrative is not None and not dryrun
    cover = strings["ed1"]["cover"] if ed1 else strings["cover"]
    refs: list[dict] = []
    rows = []
    summary_rows = strings["dryrun"]["summary_rows"] if dryrun else strings["summary_rows"]
    if manifest.metadata.get("fixture") and not dryrun:
        summary_rows = [
            row for row in summary_rows
            if all(fact_id in manifest.facts for fact_id in PLACEHOLDER.findall(row["value"]))
        ]
    for row_index, row in enumerate(summary_rows):
        value = _replace(
            row["value"],
            manifest,
            locale,
            f"table:summary:{row_index}:value",
            "table_cell",
            refs,
        )
        rows.append(f"| {row['label']} | {value} | {row['basis']} |")
    thesis = _replace(strings["thesis"], manifest, locale, "text:thesis", "text_placeholder", refs)
    methodology = _replace(
        strings["methodology"],
        manifest,
        locale,
        "text:methodology",
        "text_placeholder",
        refs,
    )
    conflict_full = strings["dryrun"]["conflict"] if dryrun else (_conflict_text(strings, conflict, "conflict") if conflict else "")
    market_rows = []
    for row_index, row in enumerate(strings["market_data"]["rows"]):
        value = _replace(row["value"], manifest, locale, f"table:market:{row_index}:value", "table_cell", refs)
        basis = _replace(row["basis"], manifest, locale, f"table:market:{row_index}:basis", "table_cell", refs)
        if dryrun and value == "UNAVAILABLE":
            basis = strings["dryrun"]["unavailable_price"] if row_index in (0, 2) else strings["dryrun"]["unavailable_input"]
        market_rows.append(f"| {row['label']} | {value} | {basis} |")
    lines = [
        f"# {strings['dryrun']['title'] if dryrun else (strings['ed1']['title'] if ed1 else strings['title'])}",
        "",
        f"> {strings['dryrun']['ribbon'] if dryrun else (strings['ed1']['ribbon'] if ed1 else strings['edition_label'])}",
        "",
        strings["disclaimer"]["third_party"],
        "",
        cover["author"],
        cover["issued"].replace("{{render_date}}", render_date or "UNAVAILABLE"),
        cover["as_of"],
        cover["cutoff"],
        cover["completeness"],
        *[f"- {item}" for item in cover["observations"]],
        "",
        conflict_full,
        "",
        thesis,
        "",
        f"### {strings['market_data']['title']}",
        "",
        f"| {strings['table']['metric']} | {strings['table']['value']} | {strings['table']['basis']} |",
        "|---|---:|---|",
        *market_rows,
        "",
        *([strings["market_data"]["note"], ""] if ed1 else []),
        f"### {strings['dryrun']['summary_title'] if dryrun else (strings['ed1']['summary_title'] if ed1 else strings['title'])}",
        "",
        f"| {strings['table']['metric']} | {strings['table']['value']} | {strings['table']['basis']} |",
        "|---|---:|---|",
        *rows,
        "",
        "<!-- TOC_START -->",
        f"## {strings['toc_title']}",
        "",
    ]
    for index, section in enumerate(strings["sections"], start=1):
        lines.append(f"- [{index}. {section['title']}](#section-{section['id']})")
    lines.extend(["", "<!-- TOC_END -->", ""])
    chart_sections = {
        "company": ["02_business_unit_mix"],
        "fq4": ["04_beat_history"],
        "outlook": ["05_scenario_fan"],
        "business": ["03b_guidance_beat_history"],
        "financials": ["01_quarterly_revenue_margin", "06_annual_income", "08_cash_flow_capex_net_cash"],
        "valuation": ["07_valuation_heatmap"],
    }
    for index, section in enumerate(strings["sections"], start=1):
        section_id = section["id"]
        lines.extend([f"<a id=\"section-{section_id}\"></a>", f"## {index}. {section['title']}", ""])
        if dryrun:
            lines.extend([strings["dryrun"]["section_stub"], ""])
        elif narrative is None:
            lines.extend([strings["section_stub"], ""])
        elif section_id in {"thesis", "scenarios", "risks", "catalysts"}:
            def replace_narrative_fact(value: str, position: str) -> str:
                return _replace(value, manifest, locale, f"text:{section_id}:{position}", "text_placeholder", refs)

            lines.extend(narrative_section_markdown(section_id, locale, narrative, replace_narrative_fact, strings))
        else:
            lines.extend([strings["section_connectors"][section_id], ""])
            if section_id == "appendix":
                if assumptions is None:
                    raise ValueError("ed1 appendix requires report-layer assumptions")
                lines.extend(appendix_markdown(assumptions, locale, strings))
        if chart_manifest:
            for chart_id in chart_sections.get(section["id"], []):
                chart = chart_manifest[chart_id]
                lines.extend([
                    f"![{chart_id}]({asset_prefix}/{chart['path']})",
                    "",
                    "CAPTION: "
                    + f"{chart['caption']['number']} {chart['caption']['title']} — "
                    + " · ".join(chart["caption"][key] for key in ("unit", "source", "as_of", "basis")),
                    "",
                ])
        if section_id == "fq4" and narrative is not None and not dryrun:
            lines.extend(_scored_table(manifest, locale, strings, refs))
        if section_id == "outlook" and ed1:
            lines.extend(_guidance_table(manifest, locale, strings, refs))
        if section_id == "margin_bridge" and ed1:
            lines.extend(_margin_bridge_table(manifest, locale, strings, refs))
        if section["id"] == "business" and (dryrun or narrative is not None):
            qualitative_rows = []
            for product in ("DRAM", "NAND"):
                period = "FQ3-26" if dryrun else "FQ4-26"
                pricing = _fact_cell(manifest, f"qual.{product.lower()}.pricing.{period}.CITED", locale, f"table:qual:{product}:pricing", refs)
                bits = _fact_cell(manifest, f"qual.{product.lower()}.bit_shipments.{period}.CITED", locale, f"table:qual:{product}:bits", refs)
                qualitative_rows.append([product, period, pricing, bits, strings["financial_tables"]["verbatim_treatment"]])
            lines.extend([
                f"### {strings['financial_tables']['dram_nand_title']}", "",
                "| " + " | ".join(strings["financial_tables"]["dram_nand_headers"]) + " |",
                "|---|---|---|---|---|",
                *["| " + " | ".join(row) + " |" for row in qualitative_rows], "",
            ])
        if section["id"] == "financials":
            if dryrun or narrative is not None:
                lines.extend(_financial_tables(manifest, locale, strings, refs))
        if section["id"] == "valuation":
            pb_value = _replace("{{fact:val.pb_trailing.FY2026.A-8K}}", manifest, locale, "table:valuation:pb", "table_cell", refs)
            pb_caption = _replace(strings["valuation"]["pb_caption"], manifest, locale, "caption:valuation:pb", "text_placeholder", refs)
            lines.extend([
                *_valuation_table(manifest, locale, strings, refs),
                f"| {strings['table']['metric']} | {strings['table']['value']} | {strings['table']['basis']} |",
                "|---|---:|---|",
                f"| {strings['valuation']['trailing_pb']} | {pb_value} | A-8K |",
                "",
                pb_caption,
                "",
            ])
            if dryrun:
                lines.extend([strings["dryrun"]["unavailable_price"], ""])
        if section["id"] == "appendix":
            lines.extend([f"### {strings['methodology_title']}", "", methodology, ""])
    lines.extend([
        strings["disclaimer"]["not_advice"],
        "",
        strings["disclaimer"]["third_party"],
        "",
        conflict_full,
        "",
    ])
    return "\n".join(lines), refs


def _html_from_markdown(markdown: str, locale: str, conflict_short: str | None = None, dryrun: bool = False) -> str:
    lines = markdown.splitlines()
    has_cover = "<!-- TOC_START -->" in lines
    body = ["<section class='cover'>"] if has_cover else []
    in_table = False
    table_header_pending = False
    in_toc = False
    pending_anchor: str | None = None

    def inline(value: str) -> str:
        rendered = html_module.escape(value, quote=False)
        rendered = re.sub(r"`([^`]+)`", r"<code>\1</code>", rendered)
        rendered = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", rendered)
        rendered = re.sub(r"\*([^*]+)\*", r"<em>\1</em>", rendered)
        return rendered

    def close_table() -> None:
        nonlocal in_table, table_header_pending
        if in_table:
            body.append("</table>")
            in_table = False
            table_header_pending = False
    for line in lines:
        if line == "<!-- TOC_START -->":
            close_table()
            body.append("</section>")
            body.append("<nav class='toc'>")
            in_toc = True
        elif line == "<!-- TOC_END -->":
            close_table()
            body.append("</nav>")
            in_toc = False
        elif re.fullmatch(r'<a id="[a-z0-9_-]+"></a>', line):
            close_table()
            pending_anchor = re.search(r'id="([^"]+)"', line).group(1)
        elif line.startswith("- [") and "](#" in line:
            match = re.fullmatch(r"- \[([^]]+)]\(#([^)]+)\)", line)
            if not match:
                raise ValueError(f"invalid table-of-contents entry: {line}")
            body.append(f"<div class='toc-entry'><a href='#{match.group(2)}'>{inline(match.group(1))}</a></div>")
        elif line.startswith("# "):
            close_table()
            body.append(f"<h1>{inline(line[2:])}</h1>")
        elif line.startswith("#### "):
            close_table()
            body.append(f"<h4>{inline(line[5:])}</h4>")
        elif line.startswith("### "):
            close_table()
            body.append(f"<h3>{inline(line[4:])}</h3>")
        elif line.startswith("## "):
            close_table()
            anchor = f" id='{pending_anchor}'" if pending_anchor else ""
            body.append(f"<h2{anchor}>{inline(line[3:])}</h2>")
            pending_anchor = None
        elif line.startswith("> "):
            close_table()
            body.append(f"<aside>{inline(line[2:])}</aside>")
        elif line.startswith("!["):
            match = re.fullmatch(r"!\[([^]]+)]\(([^)]+)\)", line)
            if not match:
                raise ValueError(f"invalid image entry: {line}")
            close_table()
            body.append(f"<figure><img src='{match.group(2)}' alt='{match.group(1)}'></figure>")
        elif line.startswith("|"):
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            if set("".join(cells)) <= {"-", ":"}:
                continue
            if not in_table:
                body.append("<table>")
                in_table = True
                table_header_pending = True
            tag = "th" if table_header_pending else "td"
            rendered_cells = []
            for index, cell in enumerate(cells):
                is_numeric = re.fullmatch(r"(?:UNAVAILABLE|N/M|—|[-+()$€£¥\d,.% ]+)", cell) is not None
                css_class = " class='wide-number'" if tag == "td" and index > 0 and is_numeric else ""
                rendered_cells.append(f"<{tag}{css_class}>{inline(cell)}</{tag}>")
            body.append("<tr>" + "".join(rendered_cells) + "</tr>")
            table_header_pending = False
        elif line:
            close_table()
            if line.startswith("CAPTION: ") and body and body[-1].startswith("<figure>"):
                body[-1] = body[-1].replace("</figure>", f"<figcaption>{line.removeprefix('CAPTION: ')}</figcaption></figure>")
            else:
                body.append(f"<p>{inline(line)}</p>")
    close_table()
    if in_toc:
        body.append("</nav>")
    strings = _locale(locale)
    footer_parts = [strings["disclaimer"]["not_advice"], strings["disclaimer"]["third_party_short"]]
    if conflict_short:
        footer_parts.append(conflict_short)
    footer = " · ".join(footer_parts).replace("'", "\\'")
    section_break = "h2[id] { break-before:page; }" if dryrun else ""
    css = f"@page {{ size: A4; margin: 20mm 16mm 19mm; @bottom-center {{ content: '{footer} · ' counter(page) '/' counter(pages); color:{COLORS['gray_dark']}; font-family:{FONT_STACK}; font-size:5.7pt; white-space:nowrap; }} }}\n" + f"""
body {{ font-family: {FONT_STACK}; color:{COLORS['ink']}; background:{COLORS['paper']}; font-size:10pt; line-height:1.5; }}
p {{ margin:2mm 0; }}
h1 {{ color:{COLORS['ink']}; border-left:8px solid {COLORS['primary']}; padding-left:12px; font-size:21pt; line-height:1.18; margin:0 0 5mm; }}
h2 {{ color:{COLORS['primary']}; border-bottom:1px solid {COLORS['gray_mid']}; padding-bottom:2mm; margin-top:9mm; }}
h4 {{ font-weight:700; font-size:10pt; margin:4mm 0 1mm; }}
{section_break}
aside {{ background:{COLORS['blue_pale']}; border-left:3px solid {COLORS['warning']}; padding:8px 12px; color:{COLORS['gray_dark']}; }}
table {{ width:100%; table-layout:fixed; border-collapse:collapse; margin:4mm 0; background:white; font-size:8pt; }}
th {{ background:{COLORS['primary']}; color:white; text-align:left; padding:5px; overflow-wrap:anywhere; }}
td {{ border-bottom:1px solid {COLORS['gray_mid']}; padding:5px; overflow-wrap:anywhere; }}
th:first-child, td:first-child {{ width:18%; }}
tr:nth-child(even) td {{ background:{COLORS['gray_pale']}; }}
td:not(:first-child) {{ font-size:7.2pt; }}
td.wide-number {{ text-align:right; font-variant-numeric:tabular-nums; white-space:nowrap; }}
.cover {{ font-size:9pt; line-height:1.35; }}
.cover p {{ margin:1.5mm 0; }}
.cover table {{ margin:2.5mm 0; }}
.cover th, .cover td {{ padding:4px; }}
p:last-child, p:nth-last-child(2) {{ color:{COLORS['gray_dark']}; font-size:8pt; }}
.toc {{ page-break-before:always; page-break-after:always; }}
.toc-entry {{ margin:2.5mm 0; }}
.toc-entry a {{ color:{COLORS['ink']}; text-decoration:none; }}
.toc-entry a::after {{ content: leader('.') target-counter(attr(href), page); }}
figure {{ margin:5mm 0; break-inside:avoid; }}
figure img {{ display:block; width:100%; max-height:105mm; object-fit:contain; }}
figcaption {{ color:{COLORS['gray_dark']}; font-size:8pt; margin-top:2mm; }}
.dryrun-header {{ position:fixed; top:-14mm; left:0; right:0; text-align:center; color:{COLORS['warning']}; font-weight:bold; font-size:7pt; }}
.dryrun-watermark {{ position:fixed; top:42%; left:8%; transform:rotate(-32deg); color:rgba(22,78,135,.10); font-size:30pt; font-weight:bold; z-index:-1; }}
"""
    dryrun_markup = ""
    if dryrun:
        watermark = strings["dryrun"]["watermark"]
        dryrun_markup = f"<div class='dryrun-header'>{watermark}</div><div class='dryrun-watermark'>{watermark}</div>"
    return f"<!doctype html><html lang='{locale}'><head><meta charset='utf-8'><style>{css}</style></head><body>{dryrun_markup}{''.join(body)}</body></html>"


def _normalize_xlsx_archive(source: Path, target: Path) -> None:
    with ZipFile(source, "r") as input_archive, ZipFile(target, "w", compression=ZIP_DEFLATED) as output_archive:
        for member in sorted(input_archive.infolist(), key=lambda item: item.filename):
            normalized = ZipInfo(member.filename, date_time=(2000, 1, 1, 0, 0, 0))
            normalized.compress_type = ZIP_DEFLATED
            normalized.create_system = 0
            normalized.external_attr = member.external_attr
            content = input_archive.read(member.filename)
            if member.filename == "docProps/core.xml":
                # openpyxl replaces the modified timestamp during save.
                content = re.sub(
                    rb"(<dcterms:modified\b[^>]*>)[^<]*(</dcterms:modified>)",
                    rb"\g<1>2000-01-01T00:00:00Z\g<2>",
                    content,
                )
            output_archive.writestr(normalized, content)


def _write_ed1_xlsx_summary(summary, conflict: ConflictConfirmation, render_date: str) -> None:
    locales = {locale: _locale(locale) for locale in ("ko", "en")}
    for row in summary.iter_rows(min_row=3, max_row=13, min_col=1, max_col=7):
        for cell in row:
            cell.value = None
    summary["A1"] = locales["ko"]["ed1"]["title"]
    summary["A3"], summary["B3"], summary["E3"] = "Field", "KO", "EN"
    for column in "BCDEFG":
        summary.column_dimensions[column].width = 20
    fields = (
        ("Edition", lambda strings: strings["ed1"]["ribbon"]),
        ("Author", lambda strings: strings["ed1"]["cover"]["author"]),
        ("Issue date", lambda strings: strings["ed1"]["cover"]["issued"].replace("{{render_date}}", render_date)),
        ("Data as of", lambda strings: strings["ed1"]["cover"]["as_of"]),
        ("Information cutoff", lambda strings: strings["ed1"]["cover"]["cutoff"]),
        ("Boundary", lambda strings: strings["ed1"]["cover"]["completeness"]),
        ("Not advice", lambda strings: strings["disclaimer"]["not_advice"]),
        ("Third party", lambda strings: strings["disclaimer"]["third_party"]),
        ("Conflict", lambda strings: _conflict_text(strings, conflict, "conflict")),
        ("Source / unit", lambda strings: "Canonical fact manifest / Unit: per fact"),
    )
    for row, (label, value_for_locale) in enumerate(fields, start=4):
        summary.cell(row, 1, label)
        summary.cell(row, 1).fill = PatternFill("solid", fgColor=COLORS["blue_pale"].lstrip("#"))
        summary.cell(row, 1).font = Font(name="Arial", size=10, bold=True, color=COLORS["primary"].lstrip("#"))
        widths = []
        for locale, first_column, last_column in (("ko", 2, 4), ("en", 5, 7)):
            summary.merge_cells(start_row=row, start_column=first_column, end_row=row, end_column=last_column)
            value = value_for_locale(locales[locale])
            cell = summary.cell(row, first_column, value)
            cell.font = Font(name="Arial", size=10, color=COLORS["ink"].lstrip("#"))
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            widths.append(sum(2 if ord(char) > 127 else 1 for char in value))
        summary.row_dimensions[row].height = max(30, ((max(widths) + 69) // 70) * 15 + 6)
    summary.merge_cells("B3:D3")
    summary.merge_cells("E3:G3")
    for address in ("A3", "B3", "E3"):
        summary[address].font = Font(name="Arial", size=10, bold=True, color=COLORS["primary"].lstrip("#"))


def _render_xlsx(
    manifest: Manifest, output_path: Path, dryrun: bool = False,
    conflict: ConflictConfirmation | None = None, render_date: str | None = None,
) -> None:
    ed1 = manifest.metadata.get("layer") == "E2-B" and not dryrun
    if ed1 and (conflict is None or render_date is None):
        raise ValueError("ed1 XLSX requires the confirmed disclosure and render date")
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    facts = workbook.create_sheet("Facts")
    summary.sheet_view.showGridLines = False
    facts.sheet_view.showGridLines = False

    ink = COLORS["ink"].lstrip("#")
    blue = COLORS["primary"].lstrip("#")
    pale_blue = COLORS["blue_pale"].lstrip("#")
    white = "FFFFFF"
    line = Side(style="thin", color="D8DEE6")

    summary.merge_cells("A1:G1")
    summary["A1"] = "MU FY2026 Q4 — PRE-PRINT DRY RUN / NOT FOR DISTRIBUTION" if dryrun else "MU FY2026 Q4 — FIXTURE / NOT REAL DATA"
    summary["A1"].fill = PatternFill("solid", fgColor=ink)
    summary["A1"].font = Font(name="Arial", size=16, bold=True, color=white)
    for row, values in enumerate(
        (
            ("Boundary", "E2-A actual history + PREREG_A only" if dryrun else "E2-A fixture only"),
            ("Week basis", "FY2026 Q4 14 weeks / FY2026 53 weeks"),
            ("Advice", "This document is not investment advice."),
            ("Branding", "Independent format; no broker or bank branding."),
        ),
        start=3,
    ):
        summary.cell(row, 1, values[0])
        summary.cell(row, 2, values[1])
        summary.cell(row, 1).fill = PatternFill("solid", fgColor=pale_blue)
        summary.cell(row, 1).font = Font(name="Arial", bold=True, color=blue)
        summary.cell(row, 2).alignment = Alignment(wrap_text=True)
    for column, value in enumerate(
        ("Source / as of", "Canonical fact manifest / 2026-09-25 DRYRUN" if dryrun else "Canonical fact manifest / 2099-01-01 FIXTURE", "Unit: per fact"),
        start=1,
    ):
        summary.cell(8, column, value)
    summary.column_dimensions["A"].width = 20
    summary.column_dimensions["B"].width = 48
    summary.column_dimensions["C"].width = 18
    for row in summary.iter_rows(min_row=1, max_row=8, min_col=1, max_col=7):
        for cell in row:
            if cell.row != 1:
                cell.font = Font(name="Arial", size=10, bold=cell.font.bold, color=cell.font.color)
    if ed1:
        _write_ed1_xlsx_summary(summary, conflict, render_date)

    headers = ("fact_id", "raw_value", "unit", "period", "basis", "source_id", "status")
    facts.append(headers)
    for fact in manifest.facts.values():
        facts.append((fact.fact_id, fact.raw_value, fact.unit, fact.period, fact.basis, fact.source_id, fact.status))
    for column, value in enumerate(("Unit", "Source", "As of"), start=9):
        facts.cell(1, column, value)
    as_of = _locale("en")["ed1"]["cover"]["as_of"] if ed1 else ("2026-09-25 DRYRUN" if dryrun else "2099-01-01 FIXTURE")
    for column, value in enumerate(("Per fact", "Canonical fact manifest", as_of), start=9):
        facts.cell(2, column, value)
    for cell in facts[1][:7]:
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.font = Font(name="Arial", size=9, bold=True, color=white)
    for row in facts.iter_rows(min_row=2, max_row=facts.max_row, min_col=1, max_col=7):
        for cell in row:
            cell.font = Font(name="Arial", size=9)
            cell.border = Border(bottom=line)
    widths = {"A": 48, "B": 16, "C": 18, "D": 16, "E": 24, "F": 18, "G": 28, "I": 14, "J": 28, "K": 22}
    for column, width in widths.items():
        facts.column_dimensions[column].width = width
    if ed1:
        facts.column_dimensions["K"].width = 60
        facts["K2"].alignment = Alignment(wrap_text=True)
        facts.row_dimensions[2].height = 30
    facts.freeze_panes = "A2"
    workbook.properties.creator = "MU report builder"
    workbook.properties.created = datetime(2000, 1, 1)
    workbook.properties.modified = datetime(2000, 1, 1)

    raw_temporary = output_path.with_name(f".{output_path.name}.raw.tmp.xlsx")
    temporary = output_path.with_name(f".{output_path.name}.tmp.xlsx")
    try:
        workbook.save(raw_temporary)
        _normalize_xlsx_archive(raw_temporary, temporary)
        os.replace(temporary, output_path)
    finally:
        workbook.close()
        for candidate in (raw_temporary, temporary):
            if candidate.exists():
                candidate.unlink()


def _chart_references(chart_manifest: dict[str, dict], manifest: Manifest) -> list[dict]:
    references = []
    for chart_id in sorted(chart_manifest):
        for point_index, fact_id in enumerate(chart_manifest[chart_id]["fact_ids"]):
            references.append(
                _reference(
                    fact_id,
                    f"chart:{chart_id}:series:0:point:{point_index}",
                    "chart_point",
                    "display",
                    manifest,
                )
            )
    return references


def render_fixture_bundle(manifest: Manifest, output_dir: str | Path) -> dict[str, Path]:
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    assets = output / "assets"
    chart_manifest = render_charts(manifest, fixture_specs(manifest), assets)
    outputs: dict[str, Path] = {}
    logs: dict[str, list[dict]] = {}
    rendered_texts: dict[str, str] = {}
    conflict = load_conflict_confirmation(FIXTURE_CONFLICT, "ed1", FIXTURE_NOW_KST)
    manifest.metadata["internal_configs"] = [
        {"path": conflict.path, "git_blob_sha": conflict.git_blob_sha, "kind": conflict.kind}
    ]
    runtime_path = output / "internal_config_runtime_log.json"
    atomic_write(
        runtime_path,
        (
            json.dumps(
                {
                    "path": conflict.path,
                    "git_blob_sha": conflict.git_blob_sha,
                    "kind": conflict.kind,
                    "read_at_kst": conflict.read_at_kst,
                },
                ensure_ascii=False,
                indent=2,
                sort_keys=True,
            )
            + "\n"
        ).encode("utf-8"),
    )
    toc_entries = {locale: _toc_entries(locale) for locale in ("ko", "en")}
    from .gates import (
        gate_g12c_conflict,
        gate_g15_toc,
        gate_inventory_days,
        gate_market_data_box,
        gate_trailing_pb,
    )

    gate_g15_toc(toc_entries)
    gate_inventory_days(manifest)
    for locale in ("ko", "en"):
        gate_market_data_box(_locale(locale)["market_data"])
    gate_trailing_pb(
        manifest,
        {
            "price_date": str(manifest.fact("meta.price_date.2026-10-01.CITED").raw_value),
            "equity_date": str(manifest.fact("meta.equity_date.FY2026.A-8K").raw_value),
            "shares_date": str(manifest.fact("meta.shares_date.CITED").raw_value),
        },
    )
    rendered_conflict: dict[str, dict] = {}
    for locale in ("ko", "en"):
        strings = _locale(locale)
        conflict_full = _conflict_text(strings, conflict, "conflict")
        conflict_short = _conflict_text(strings, conflict, "conflict_short")
        markdown, refs = render_markdown(manifest, locale, conflict)
        refs.extend(_chart_references(chart_manifest, manifest))
        from .gates import gate_g17_labels, gate_g17b_consensus

        gate_g17_labels(markdown, locale)
        gate_g17b_consensus(refs, manifest)
        html = _html_from_markdown(markdown, locale, conflict_short)
        md_path = output / f"mu_fixture_{locale}.md"
        html_path = output / f"mu_fixture_{locale}.html"
        pdf_path = output / f"mu_fixture_{locale}.pdf"
        atomic_write(md_path, markdown.encode("utf-8"))
        atomic_write(html_path, html.encode("utf-8"))
        pdf_temporary = pdf_path.with_name(f".{pdf_path.name}.tmp")
        try:
            HTML(string=html, base_url=str(output)).write_pdf(pdf_temporary)
            os.replace(pdf_temporary, pdf_path)
        finally:
            if pdf_temporary.exists():
                pdf_temporary.unlink()
        outputs.update({f"md_{locale}": md_path, f"html_{locale}": html_path, f"pdf_{locale}": pdf_path})
        logs[locale] = refs
        rendered_texts[locale] = markdown
        rendered_conflict[locale] = {
            "cover": conflict_full,
            "ending": conflict_full,
            "pdf_pages": [page.extract_text() or "" for page in PdfReader(pdf_path).pages],
            "record_blob_sha": conflict.git_blob_sha,
            "conflict_short": conflict_short,
        }
    from .gates import gate_g15_parity

    gate_g15_parity(logs, manifest, rendered_texts)
    gate_g12c_conflict(conflict, "ed1", FIXTURE_NOW_KST, rendered_conflict)
    ref_path = output / "fact_references.json"
    atomic_write(ref_path, (json.dumps(logs, indent=2, sort_keys=True) + "\n").encode("utf-8"))
    chart_path = assets / "chart_manifest.json"
    xlsx_path = output / "mu_fixture_data.xlsx"
    _render_xlsx(manifest, xlsx_path)
    page_counts = {
        locale: len(PdfReader(outputs[f"pdf_{locale}"]).pages)
        for locale in ("ko", "en")
    }
    manifest.metadata["document"] = {"total_pages": page_counts}
    manifest_path = output / "mu_fixture_manifest.json"
    manifest.write(manifest_path)
    outputs.update(
        {
            "references": ref_path,
            "chart_manifest": chart_path,
            "xlsx": xlsx_path,
            "manifest": manifest_path,
            "runtime_log": runtime_path,
        }
    )
    return outputs


def render_pdf_pages(pdf_paths: dict[str, Path], output_dir: Path) -> list[Path]:
    import fitz

    output_dir.mkdir(parents=True, exist_ok=True)
    written = []
    for locale, pdf_path in pdf_paths.items():
        document = fitz.open(pdf_path)
        try:
            for index, page in enumerate(document, start=1):
                pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                path = output_dir / f"MU_FY2026Q4_DRYRUN_{locale}_page_{index:02d}.png"
                temporary = path.with_name(f".{path.name}.tmp.png")
                try:
                    pixmap.save(temporary)
                    os.replace(temporary, path)
                finally:
                    if temporary.exists():
                        temporary.unlink()
                written.append(path)
        finally:
            document.close()
    return written


def render_dryrun_bundle(manifest: Manifest, output_dir: str | Path, audit_payload: bytes) -> dict[str, Path]:
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    assets = output / "assets"
    audit_path = output / "MU_FY2026Q4_DRYRUN_input_audit.json"
    atomic_write(audit_path, audit_payload)
    outputs: dict[str, Path] = {"audit": audit_path}
    logs: dict[str, list[dict]] = {}
    rendered_texts: dict[str, str] = {}
    rendered_conflict: dict[str, dict] = {}
    toc_entries = {locale: _toc_entries(locale) for locale in ("ko", "en")}
    from .gates import (
        DRYRUN_CONFLICT,
        gate_g12c_dryrun,
        gate_g15_parity,
        gate_g15_toc,
        gate_g17_labels,
        gate_g17b_consensus,
        gate_inventory_days,
        gate_market_data_box,
        gate_theme,
        gate_trailing_pb,
    )

    gate_g15_toc(toc_entries)
    gate_inventory_days(manifest)
    for locale in ("ko", "en"):
        gate_market_data_box(_locale(locale)["market_data"])
    gate_trailing_pb(
        manifest,
        {"price_date": "UNAVAILABLE", "equity_date": "UNAVAILABLE", "shares_date": "UNAVAILABLE"},
    )
    for locale in ("ko", "en"):
        strings = _locale(locale)
        chart_manifest = render_charts(manifest, dryrun_specs(manifest, locale, strings), assets, locale)
        disclosure = DRYRUN_CONFLICT[locale]
        markdown, refs = render_markdown(manifest, locale, dryrun=True, chart_manifest=chart_manifest)
        refs.extend(_chart_references(chart_manifest, manifest))
        gate_g17_labels(markdown, locale)
        gate_g17b_consensus(refs, manifest)
        html = _html_from_markdown(markdown, locale, disclosure, dryrun=True)
        gate_theme(
            html,
            [
                (PACKAGE_DIR / "i18n" / "ko.yaml").read_text(encoding="utf-8"),
                (PACKAGE_DIR / "i18n" / "en.yaml").read_text(encoding="utf-8"),
            ],
        )
        md_path = output / f"MU_FY2026Q4_DRYRUN_{locale}.md"
        html_path = output / f"MU_FY2026Q4_DRYRUN_{locale}.html"
        pdf_path = output / f"MU_FY2026Q4_DRYRUN_{locale}.pdf"
        atomic_write(md_path, markdown.encode("utf-8"))
        atomic_write(html_path, html.encode("utf-8"))
        pdf_temporary = pdf_path.with_name(f".{pdf_path.name}.tmp")
        try:
            HTML(string=html, base_url=str(output)).write_pdf(pdf_temporary)
            os.replace(pdf_temporary, pdf_path)
        finally:
            if pdf_temporary.exists():
                pdf_temporary.unlink()
        page_texts = [page.extract_text() or "" for page in PdfReader(pdf_path).pages]
        outputs.update({f"md_{locale}": md_path, f"html_{locale}": html_path, f"pdf_{locale}": pdf_path})
        outputs[f"chart_manifest_{locale}"] = assets / f"chart_manifest_{locale}.json"
        logs[locale] = refs
        rendered_texts[locale] = markdown
        rendered_conflict[locale] = {
            "cover": disclosure,
            "ending": disclosure,
            "pdf_pages": page_texts,
        }
    gate_g15_parity(logs, manifest, rendered_texts)
    gate_g12c_dryrun(rendered_conflict)
    reference_path = output / "MU_FY2026Q4_DRYRUN_fact_references.json"
    atomic_write(reference_path, (json.dumps(logs, indent=2, sort_keys=True) + "\n").encode("utf-8"))
    xlsx_path = output / "MU_FY2026Q4_DRYRUN_data.xlsx"
    _render_xlsx(manifest, xlsx_path, dryrun=True)
    page_counts = {locale: len(PdfReader(outputs[f"pdf_{locale}"]).pages) for locale in ("ko", "en")}
    manifest.metadata["document"] = {"total_pages": page_counts, "mode": "DRYRUN"}
    manifest_path = output / "MU_FY2026Q4_DRYRUN_manifest.json"
    manifest.write(manifest_path)
    outputs.update(
        {
            "references": reference_path,
            "chart_manifest": assets / "chart_manifest_en.json",
            "xlsx": xlsx_path,
            "manifest": manifest_path,
        }
    )
    render_pdf_pages({locale: outputs[f"pdf_{locale}"] for locale in ("ko", "en")}, output / "pages")
    return outputs


def render_e2b_bundle(
    manifest: Manifest,
    narrative: dict,
    assumptions: ReportLayerAssumptions,
    conflict: ConflictConfirmation,
    output_dir: str | Path,
    audit_payload: bytes,
    render_started_at_kst: datetime,
) -> tuple[dict[str, Path], dict[str, object]]:
    """Render the ed1 deliverables without running post-render gates."""
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    stem = "mu_report_fy2026q4_ed1"
    assets = output / f"{stem}_assets"
    assets.mkdir(parents=True, exist_ok=True)
    outputs: dict[str, Path] = {}
    references: dict[str, list[dict]] = {}
    rendered_texts: dict[str, str] = {}
    chart_manifests: dict[str, dict] = {}
    rendered_conflict: dict[str, dict] = {}

    manifest.metadata["internal_configs"] = [
        {"path": conflict.path, "git_blob_sha": conflict.git_blob_sha, "kind": conflict.kind}
    ]
    audit = json.loads(audit_payload)
    audit["internal_configs"] = manifest.metadata["internal_configs"]
    input_manifest_path = output / f"{stem}_input_manifest.json"
    atomic_write(
        input_manifest_path,
        (json.dumps(audit, ensure_ascii=False, indent=2, sort_keys=True) + "\n").encode("utf-8"),
    )
    outputs["input_manifest"] = input_manifest_path

    runtime = {
        "phase": "E2-B",
        "render_started_at_kst": render_started_at_kst.isoformat(),
        "conflict_confirmation": {
            "path": conflict.path,
            "git_blob_sha": conflict.git_blob_sha,
            "read_at_kst": conflict.read_at_kst,
            "kind": conflict.kind,
        },
        "reads": audit["reads"],
    }
    runtime_dir = PACKAGE_DIR.parents[2] / "logs" / "_mu_report_runs"
    runtime_name = render_started_at_kst.strftime("e2b_%Y%m%dT%H%M%S%z.json")
    runtime_path = runtime_dir / runtime_name
    atomic_write(
        runtime_path,
        (json.dumps(runtime, ensure_ascii=False, indent=2, sort_keys=True) + "\n").encode("utf-8"),
    )

    for locale in ("ko", "en"):
        strings = _locale(locale)
        chart_manifest = render_charts(
            manifest,
            e2b_specs(manifest, locale, strings),
            assets,
            locale,
            write_manifest=False,
        )
        conflict_full = _conflict_text(strings, conflict, "conflict")
        conflict_short = _conflict_text(strings, conflict, "conflict_short")
        markdown, refs = render_markdown(
            manifest,
            locale,
            conflict,
            chart_manifest=chart_manifest,
            narrative=narrative,
            assumptions=assumptions,
            asset_prefix=assets.name,
            render_date=render_started_at_kst.date().isoformat(),
        )
        refs.extend(_chart_references(chart_manifest, manifest))
        html_text = _html_from_markdown(markdown, locale, conflict_short)
        md_path = output / f"{stem}_{locale}.md"
        html_path = output / f"{stem}_{locale}.html"
        pdf_path = output / f"{stem}_{locale}.pdf"
        atomic_write(md_path, markdown.encode("utf-8"))
        atomic_write(html_path, html_text.encode("utf-8"))
        pdf_temporary = pdf_path.with_name(f".{pdf_path.name}.tmp")
        try:
            HTML(string=html_text, base_url=str(output)).write_pdf(pdf_temporary)
            os.replace(pdf_temporary, pdf_path)
        finally:
            if pdf_temporary.exists():
                pdf_temporary.unlink()
        pdf_pages = [page.extract_text() or "" for page in PdfReader(pdf_path).pages]
        outputs.update({f"md_{locale}": md_path, f"html_{locale}": html_path, f"pdf_{locale}": pdf_path})
        references[locale] = refs
        rendered_texts[locale] = markdown
        chart_manifests[locale] = chart_manifest
        rendered_conflict[locale] = {
            "cover": conflict_full,
            "ending": conflict_full,
            "pdf_pages": pdf_pages,
            "record_blob_sha": conflict.git_blob_sha,
            "conflict_short": conflict_short,
        }

    xlsx_path = output / f"{stem}_data.xlsx"
    _render_xlsx(manifest, xlsx_path, conflict=conflict, render_date=render_started_at_kst.date().isoformat())
    outputs["xlsx"] = xlsx_path
    page_counts = {
        locale: len(PdfReader(outputs[f"pdf_{locale}"]).pages)
        for locale in ("ko", "en")
    }
    manifest.metadata["document"] = {"total_pages": page_counts, "mode": "E2-B", "edition": "ed1"}
    manifest_path = output / f"{stem}_manifest.json"
    manifest.write(manifest_path)
    outputs["manifest"] = manifest_path
    state: dict[str, object] = {
        "references": references,
        "rendered_texts": rendered_texts,
        "chart_manifests": chart_manifests,
        "rendered_conflict": rendered_conflict,
        "runtime_log": runtime_path,
    }
    return outputs, state
