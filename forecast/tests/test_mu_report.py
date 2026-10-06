"""E2-A gates: every check has a pass case and a deliberate violation."""

from __future__ import annotations

import copy
import json
from dataclasses import replace
from io import BytesIO
from pathlib import Path

import pytest
from pypdf import PdfReader

from forecast.engine.generic_forecast import run_generic_forecast
from forecast.schemas.generic import GenericProfile
from forecast.scripts.mu_report.charts import CHART_NAMES, e2b_specs, fixture_specs, render_charts
from forecast.scripts.mu_report.extract import normalize_pdf_text
from forecast.scripts.mu_report.facts import fixture_manifest, historical_manifest_from_extracts
from forecast.scripts.mu_report.gates import (
    GateError,
    gate_categorical_verbatim,
    gate_g1_inputs,
    gate_g2_freeze_display,
    gate_g3_financial_identities,
    gate_g3b_balance_sheet,
    gate_g3c_cash_flow,
    gate_g3d_company_fcf,
    gate_g3f_scored_source,
    gate_g7_templates_and_io,
    gate_g9_cutoff,
    gate_g12_manifest,
    gate_g12b_disclaimers,
    gate_g13_charts,
    gate_g13b_chart_identity,
    gate_g13c_chart_semantics,
    gate_g14_provenance,
    gate_g14_template_numbers,
    gate_g15_parity,
    gate_g15b_localized_ui,
    gate_g16_caption,
    gate_g17_labels,
    gate_g17b_consensus,
    gate_g18_hygiene,
    gate_g19_bridge,
    gate_g20_audit,
    gate_g21_formats,
    gate_g23_availability,
    gate_g24_raw_markup,
    gate_g25_cover_dates,
    gate_ed1_rendered,
    gate_narrative_contract,
    qa_html,
    qa_markdown,
    qa_pdf,
    qa_pdf_texts,
    qa_xlsx_rows,
)
from forecast.scripts.mu_report.inputs import AuditEntry, EvidenceReader, InputGateError, load_pins
from forecast.scripts.mu_report.render import (
    _html_from_markdown,
    _render_xlsx,
    render_markdown,
    render_fixture_bundle,
)
from forecast.scripts.mu_report.narrative import e2b_manifest_from_inputs, recompute_fact_bindings
from forecast.scripts.mu_report.rle import (
    CashAssumptions,
    cash_flow,
    load_assumptions_text,
    margin_bridge,
    project_generic,
    quarterly_da_roll_forward,
    raw_growth_from_weekly,
    roll_forward_net_cash,
    fy2027_scenario_results,
    annual_rle_results,
)
from forecast.scripts.mu_report.valuation import (
    implied_valuation,
    normalize_52_from_53,
    sensitivity_heatmap,
)


@pytest.fixture(scope="session")
def rendered_bundle(tmp_path_factory):
    output = tmp_path_factory.mktemp("mu-report-fixture")
    manifest = fixture_manifest()
    outputs = render_fixture_bundle(manifest, output)
    return manifest, outputs


def test_g1_input_pins_pass_and_violation():
    gate_g1_inputs(EvidenceReader("E2-A"))
    pins = copy.deepcopy(load_pins())
    pins["E2-A"][0]["sha256"] = "f" * 64
    with pytest.raises(InputGateError, match="SHA-256 mismatch"):
        EvidenceReader("E2-A", pins).read_bytes(pins["E2-A"][0]["path"])


def test_g2_freeze_display_pass_and_violation():
    expected = {"revenue": "$52.86B", "eps": "$37.41"}
    gate_g2_freeze_display(expected, dict(expected))
    with pytest.raises(GateError):
        gate_g2_freeze_display(expected, {**expected, "eps": "$37.42"})


def _income_statement():
    return {"revenue": 37378, "cogs": 22505, "gross_profit": 14873, "r_and_d": 3798, "sg_and_a": 1205, "restructuring": 39, "other_operating": 61, "operating_income": 9770, "interest_income": 496, "interest_expense": -477, "other_nonoperating": -135, "pretax_income": 9654, "tax": -1124, "equity_method": 9, "net_income": 8539, "diluted_eps": 7.59, "diluted_shares": 1125}


def test_g3_financial_identity_pass_tolerance_and_violation():
    differences = gate_g3_financial_identities(_income_statement())
    assert differences
    broken = _income_statement()
    broken["gross_profit"] += 2
    with pytest.raises(GateError):
        gate_g3_financial_identities(broken)


def test_g3b_balance_sheet_pass_and_violation():
    values = {"total_assets": 82798, "total_liabilities": 28633, "total_equity": 54165}
    assert gate_g3b_balance_sheet(values) == []
    with pytest.raises(GateError):
        gate_g3b_balance_sheet({**values, "total_assets": 82800})


def test_g3c_cash_flow_pass_and_violation():
    values = {"operating": 10, "investing": -3, "financing": -2, "fx": 1, "cash_change": 6}
    assert gate_g3c_cash_flow(values) == []
    with pytest.raises(GateError):
        gate_g3c_cash_flow({**values, "cash_change": 9})


def test_g3d_company_fcf_pass_and_violation():
    values = {"operating_cash_flow": 25388, "ppe_expenditures": 7826, "ppe_disposal_proceeds": 9, "government_incentives": 733, "net_capex": 7084, "adjusted_free_cash_flow": 18304}
    gate_g3d_company_fcf(values)
    with pytest.raises(GateError):
        gate_g3d_company_fcf({**values, "adjusted_free_cash_flow": 18306})


def test_g7_hardcoding_and_direct_open_pass_and_violation(tmp_path):
    good = tmp_path / "good.py"
    good.write_text("from pathlib import Path\nvalue = Path('x').read_text()\n", encoding="utf-8", newline="\n")
    gate_g7_templates_and_io([good])
    bad = tmp_path / "bad.py"
    bad.write_text("value = open('x').read()\n", encoding="utf-8", newline="\n")
    with pytest.raises(GateError, match="direct open"):
        gate_g7_templates_and_io([bad])
    gate_g7_templates_and_io([Path("forecast/scripts/mu_report/i18n/ko.yaml"), Path("forecast/scripts/mu_report/i18n/en.yaml")])
    numeric = tmp_path / "i18n" / "bad.yaml"
    numeric.parent.mkdir()
    numeric.write_text("claim: Revenue is 12345\n", encoding="utf-8")
    with pytest.raises(GateError, match="unreferenced"):
        gate_g7_templates_and_io([numeric])
    dated = tmp_path / "i18n" / "dated.yaml"
    dated.write_text("value: '{{fact:market.shares_outstanding.2026-09-01.CITED}}'\n", encoding="utf-8")
    with pytest.raises(GateError, match="date-dependent"):
        gate_g7_templates_and_io([dated])


def test_g9_cutoff_pass_and_violation():
    gate_g9_cutoff([{"section": "PREREG_A", "information_class": "PRE_PRINT"}, {"section": "RLE", "basis": "A-8K"}])
    with pytest.raises(GateError):
        gate_g9_cutoff([{"section": "PREREG_A", "information_class": "POST_PRINT"}])


def test_g12_and_g19_manifest_pass_and_bridge_violation():
    manifest = fixture_manifest()
    gate_g12_manifest(manifest)
    gate_g19_bridge(manifest)
    fact = manifest.facts["bridge.nongaap_fixed_027.FY2026Q4"]
    manifest.facts[fact.fact_id] = replace(fact, period="FY2027Q1")
    with pytest.raises(ValueError, match="scoped"):
        gate_g19_bridge(manifest)


def test_g12_forbidden_field_violation():
    manifest = fixture_manifest()
    manifest.metadata["target_price"] = 123
    with pytest.raises(ValueError, match="forbidden"):
        gate_g12_manifest(manifest)


def test_g12b_disclaimer_pass_and_violation():
    good = {
        "ko": "본 문서는 투자 자문이 아닙니다. 공개 자료 기반 제3자 분석이며 Micron 및 계열사가 작성·검토·승인한 자료가 아닙니다.",
        "en": "This document is not investment advice. This is third-party analysis based on public information and has not been prepared, reviewed, or approved by Micron or its affiliates.",
    }
    gate_g12b_disclaimers(good)
    with pytest.raises(GateError):
        gate_g12b_disclaimers({**good, "en": "No advice"})
    with pytest.raises(GateError, match="disclaimer text"):
        gate_g12b_disclaimers({**good, "ko": good["ko"].replace("제3자", "독립")})


def test_sca_dependencies_and_debt_bases_violation():
    manifest = fixture_manifest()
    manifest.validate()
    debt = manifest.facts["note.debt_prepayment_loss.FY2026Q3.A"]
    manifest.facts[debt.fact_id] = replace(debt, raw_value=325)
    with pytest.raises(ValueError, match="bases must remain separate"):
        manifest.validate()


def test_g13_and_g13b_chart_pass_and_violation(tmp_path):
    manifest = fixture_manifest()
    charts = render_charts(manifest, fixture_specs(manifest), tmp_path)
    gate_g13_charts(charts)
    gate_g13b_chart_identity(manifest, charts, 7.0, 7.0)
    broken = copy.deepcopy(charts)
    broken[CHART_NAMES[0]]["values"][0] += 1
    with pytest.raises(GateError):
        gate_g13b_chart_identity(manifest, broken)
    with pytest.raises(GateError):
        gate_g13b_chart_identity(manifest, charts, 7.0, 7.1)
    with pytest.raises(GateError):
        gate_g13_charts({key: value for key, value in charts.items() if key != CHART_NAMES[-1]})


def test_g14_provenance_pass_and_violation():
    manifest = fixture_manifest()
    gate_g14_provenance(manifest)
    fact = manifest.facts["valuation.ev_adj.FY2026Q4"]
    manifest.facts[fact.fact_id] = replace(fact, lineage=None)
    with pytest.raises(ValueError, match="lineage"):
        gate_g14_provenance(manifest)


def test_g14_template_number_pass_and_violation():
    gate_g14_template_numbers("FY2026Q4 has 14 weeks; {{fact:is.revenue.FY2026Q4.A-8K}}")
    with pytest.raises(GateError, match="unreferenced"):
        gate_g14_template_numbers("Revenue is 12345")


def _parity_reference(manifest, fact_id, position, kind):
    fact = manifest.fact(fact_id)
    return {
        "position": position,
        "kind": kind,
        "role": "display",
        "fact_id": fact_id,
        "raw_value": fact.raw_value,
        "unit": fact.unit,
        "period": fact.period,
        "basis": fact.basis,
        "label": fact.label,
        "status": fact.status,
    }


def test_g15_parity_all_positions_and_backup_numbers():
    manifest = fixture_manifest()
    good = [
        _parity_reference(manifest, "is.revenue.FY2026Q4.A-8K", "table:summary:0:value:0", "table_cell"),
        _parity_reference(manifest, "consensus.revenue.FY2027", "text:methodology:0", "text_placeholder"),
        _parity_reference(manifest, "is.eps_gaap.FY2026Q4.PREREG_A", "chart:01:series:0:point:0", "chart_point"),
    ]
    texts = {"ko": "FY2026 99,999 30.73", "en": "FY2026 99,999 30.73"}
    gate_g15_parity({"ko": good, "en": copy.deepcopy(good)}, manifest, texts)

    swapped = copy.deepcopy(good)
    swapped[0]["fact_id"], swapped[1]["fact_id"] = swapped[1]["fact_id"], swapped[0]["fact_id"]
    with pytest.raises(GateError, match="fact_id"):
        gate_g15_parity({"ko": good, "en": swapped}, manifest, texts)
    with pytest.raises(GateError, match="positions"):
        gate_g15_parity({"ko": good, "en": copy.deepcopy(good[:-1])}, manifest, texts)
    with pytest.raises(GateError, match="number multisets"):
        gate_g15_parity({"ko": good, "en": copy.deepcopy(good)}, manifest, {**texts, "en": "FY2026 99,998 30.73"})


def test_g16_caption_pass_and_violation():
    good = {"number": "Figure 1.", "title": "Quarterly revenue", "unit": "USD million / %", "source": "Micron 10-K [SRC-10K-FY25]", "as_of": "2026-09-25", "basis": "GAAP actual (A)"}
    gate_g16_caption(good)
    with pytest.raises(GateError):
        gate_g16_caption({**good, "source": ""})
    with pytest.raises(GateError, match="generic"):
        gate_g16_caption({**good, "unit": "per chart axis"})
    with pytest.raises(GateError, match="accounting basis"):
        gate_g16_caption({**good, "basis": "A / PREREG_A"})


def test_g17_labels_each_language_and_fact_based_consensus():
    gate_g17_labels("FY2026Q4 calendar Q3 14 weeks 53 weeks PREREG_A A-8K RLE UNAVAILABLE", "en")
    gate_g17_labels("FY2026Q4 달력 Q3 14주 53주 PREREG_A A-8K RLE UNAVAILABLE", "ko")
    manifest = fixture_manifest()
    reference = _parity_reference(manifest, "consensus.revenue.FY2027", "text:methodology:0", "text_placeholder")
    gate_g17b_consensus([reference], manifest)
    with pytest.raises(GateError):
        gate_g17_labels("FY2026Q4 14 weeks 53 weeks PREREG_A A-8K RLE UNAVAILABLE", "ko")
    comparison = {**reference, "role": "comparison"}
    with pytest.raises(GateError, match="used as comparison"):
        gate_g17b_consensus([comparison], manifest)
    derived = manifest.facts["valuation.ev_adj.FY2026Q4"]
    manifest.facts[derived.fact_id] = replace(
        derived,
        lineage={"formula": "forbidden", "inputs": ["bs.sca_customer_deposits.FY2026Q4", "consensus.revenue.FY2027"]},
    )
    with pytest.raises(GateError, match="consumes unavailable consensus"):
        gate_g17b_consensus([reference], manifest)


def test_g18_hygiene_pass_and_violations(tmp_path):
    good = tmp_path / "good.md"
    good.write_bytes(b"clean\n")
    gate_g18_hygiene([good])
    for name, payload in (("crlf.md", b"bad\r\n"), ("nul.md", b"bad\x00\n"), ("space.md", b"bad \n")):
        bad = tmp_path / name
        bad.write_bytes(payload)
        with pytest.raises(GateError):
            gate_g18_hygiene([bad])


def test_g20_audit_pass_and_reserved_violation():
    reader = EvidenceReader("E2-A")
    reader.read_bytes(load_pins()["E2-A"][0]["path"])
    gate_g20_audit(reader)
    reader.audit.append(AuditEntry("logs/mu/fy2026q4/postprint/ex991.htm", "0" * 64, "E2-A"))
    with pytest.raises(GateError, match="reserved"):
        gate_g20_audit(reader)

    class NondeterministicReader:
        calls = 0
        def audit_payload(self):
            self.calls += 1
            return json.dumps({"phase": "E2-A", "reads": [], "run": self.calls}).encode()
    with pytest.raises(GateError, match="non-deterministic"):
        gate_g20_audit(NondeterministicReader())


def test_g21_rendered_formats_pass(rendered_bundle):
    _, outputs = rendered_bundle
    gate_g21_formats(outputs)


def test_g21_pdf_html_md_xlsx_violations(tmp_path, rendered_bundle):
    _, outputs = rendered_bundle
    bad_pdf = tmp_path / "bad.pdf"
    bad_pdf.write_bytes(outputs["pdf_en"].read_bytes()[:-6])
    with pytest.raises(GateError):
        qa_pdf(bad_pdf, "en")
    with pytest.raises(GateError, match="file footer"):
        qa_pdf_texts(["This document is not investment advice. file:///tmp/report"], "This document is not investment advice.")
    bad_html = tmp_path / "bad.html"
    bad_html.write_text('<script src="https://example.com/x.js"></script>', encoding="utf-8")
    with pytest.raises(GateError):
        qa_html(bad_html)
    bad_md = tmp_path / "bad.md"
    bad_md.write_text("| malformed\n", encoding="utf-8")
    with pytest.raises(GateError):
        qa_markdown(bad_md)
    with pytest.raises(GateError):
        qa_xlsx_rows({"Sheet1": ["Source", "As of", "Unit", "#REF!"]})


def test_rle_and_valuation_formulas_pass_and_violations():
    assert raw_growth_from_weekly(0.05, 14, 13) == pytest.approx(-0.025)
    bridge = margin_bridge(100, 0.6, 20, -1, 0.2, 0, 10)
    assert bridge["gross_profit"] - bridge["opex"] == bridge["operating_income"]
    cash = cash_flow(10, CashAssumptions(3, 1, 2, 0, 8, 1, 2, 1))
    assert cash == {"operating_cash_flow": 12, "net_capex": 5, "adjusted_fcf": 7}
    assert roll_forward_net_cash(20, 7, 1) == 26
    result = implied_valuation(100, 1000, 10, 5, 20, 100, 120, 80, None)
    assert result.ev_adj is None and result.pe == 1000
    assert sensitivity_heatmap([100], [100], 1000)[0][0] == result.pe
    assert normalize_52_from_53(53, "revenue") == 52
    with pytest.raises(ValueError):
        normalize_52_from_53(53, "market_cap")


def test_rle_fixture_schema_and_existing_engine_path():
    fixture_text = Path("forecast/tests/fixtures/mu_report/fake_postprint.yaml").read_text(encoding="utf-8")
    assumptions = load_assumptions_text(fixture_text)
    assert assumptions.scope == "report_layer" and assumptions.confidence == "low"
    assert "A11_net_capex_roll_forward" in assumptions.interpretation_note
    with pytest.raises(ValueError):
        load_assumptions_text(fixture_text.replace("scope: report_layer", "scope: valuation"))
    with pytest.raises(ValueError):
        load_assumptions_text(
            "scope: report_layer\ninformation_cutoff: 2099-01-01\nsources: [SRC-FIXTURE]\n"
            "units: USD_million\nweek_basis: {FY2027Q1: 13}\nconfidence: low\n"
            f"input_sha256: {'0' * 64}\nfq4_actual: {{revenue: 99999}}\nrle: {{}}\n"
        )
    rules = assumptions.rules
    fq4_revenue = assumptions.post_print_inputs["fq4_actual_revenue"].value
    profile = GenericProfile.model_validate(
        {
            "name": "FIXTURE — NOT REAL DATA",
            "name_kr": "FIXTURE — NOT REAL DATA",
            "ticker": "FAKE",
            "currency": "USD",
            "reporting_unit": "USD_million",
            "fiscal_year_end_month": 8,
            "weighted_avg_diluted": 1_000_000_000,
            "seed": {"quarter_label": "2026Q4", "revenue_total": 99999, "net_profit": 77777},
            "window": {"start_quarter": "2027Q1", "n_quarters": 4},
            "bear": {"probability": 0.2, "revenue_growth_qoq": [0.01] * 4, "op_margin": 0.5, "effective_tax_rate": 0.2},
            "base": {"probability": 0.5, "revenue_growth_qoq": [0.02] * 4, "op_margin": 0.6, "effective_tax_rate": 0.2},
            "bull": {"probability": 0.3, "revenue_growth_qoq": [0.03] * 4, "op_margin": 0.7, "effective_tax_rate": 0.2},
        }
    )
    overrides = {
        scenario: {
            "revenue_growth_qoq": [
                rules["A1_fq1_revenue"][scenario]["value"] / fq4_revenue - 1.0,
                *rules["A2_fq2_to_fq4_weekly_revenue_growth"][scenario]["value"],
            ]
        }
        for scenario in ("bear", "base", "bull")
    }
    forecast = project_generic(profile, overrides)
    expected = rules["A1_fq1_revenue"]["base"]["value"]
    assert forecast.scenarios_quarterly["base"][0].revenue_total == pytest.approx(expected)


def test_r13_actual_yaml_reproduces_r12_and_cash_flow_values():
    assumptions = load_assumptions_text(
        Path("forecast/inputs/mu_fy2026q4_report_assumptions.yaml").read_text(encoding="utf-8")
    )
    rules = assumptions.rules
    inputs = assumptions.post_print_inputs
    fq4_revenue = inputs["A4_fq4_fy2026_actual_revenue_14w"].value
    shares_million = rules["A7_diluted_shares"]["quarterly_all_scenarios"]["value"][0]
    scenario_probabilities = {"bear": 0.333333, "base": 0.333334, "bull": 0.333333}
    profile_data = {
        "name": "MU RLE CONTRACT TEST",
        "name_kr": "MU RLE CONTRACT TEST",
        "ticker": "MU",
        "currency": "USD",
        "reporting_unit": "USD_million",
        "fiscal_year_end_month": 8,
        "weighted_avg_diluted": shares_million * 1_000_000,
        "seed": {"quarter_label": "2026Q4", "revenue_total": fq4_revenue, "net_profit": 37704},
        "window": {"start_quarter": "2027Q1", "n_quarters": 4},
    }
    opex = rules["A4_gaap_opex"]["quarterly_all_scenarios"]["value"]
    below_op_pct = rules["A5_below_operating_pct_of_revenue"]["all_scenarios"]["value"]
    for scenario in ("bear", "base", "bull"):
        fq1_revenue = rules["A1_fq1_revenue"][scenario]["value"]
        growth = [
            fq1_revenue / fq4_revenue - 1.0,
            *rules["A2_fq2_to_fq4_weekly_revenue_growth"][scenario]["value"],
        ]
        revenue = [fq1_revenue]
        for rate in growth[1:]:
            revenue.append(revenue[-1] * (1.0 + rate))
        gross_margin = rules["A3_gaap_gross_margin"][scenario]["value"]
        profile_data[scenario] = {
            "probability": scenario_probabilities[scenario],
            "revenue_growth_qoq": growth,
            "op_margin": [margin - expense / sales for margin, expense, sales in zip(gross_margin, opex, revenue)],
            "effective_tax_rate": rules["A6_gaap_effective_tax_rate"][scenario]["value"],
            "net_interest_pct_of_revenue": below_op_pct,
        }

    forecast = run_generic_forecast(GenericProfile.model_validate(profile_data))
    base_rows = forecast.scenarios_quarterly["base"]
    base_revenue = sum(row.revenue_total for row in base_rows)
    base_net_income = sum(row.net_profit for row in base_rows)
    assert base_revenue == pytest.approx(274784.139585)
    assert base_net_income / shares_million == pytest.approx(169.05, abs=0.01)

    da_rows = quarterly_da_roll_forward(
        inputs["fy2026_ending_ppe"].value,
        rules["A14_net_capex"]["quarterly_all_scenarios"]["value"],
        rules["A11_da"]["fy2026_rate_d"]["value"],
    )
    assert sum(row["da"] for row in da_rows) == pytest.approx(
        rules["A11_da"]["fy2027_total"]["value"]
    )
    assert da_rows[-1]["ppe_end"] == pytest.approx(rules["A11_da"]["fy2027_ending_ppe"]["value"])

    k = rules["A13_working_capital"]["median_k"]["value"]
    sbc = rules["A12_sbc"]["fy2027_total"]["value"]
    net_capex = rules["A14_net_capex"]["fy2027_total"]["value"]
    dividends = rules["A15_dividends"]["fy2027_total"]["value"]
    opening_net_cash = rules["opening_net_cash_unadjusted"]["amount"]["value"]
    working_capital_investment = k * (base_revenue - inputs["fy2026_revenue"].value)
    cfo = base_net_income + sum(row["da"] for row in da_rows) + sbc - working_capital_investment
    fcf = cfo - net_capex
    assert fcf == pytest.approx(155589.2775406485)
    assert roll_forward_net_cash(opening_net_cash, fcf, dividends) == pytest.approx(223173.2775406485)


def _e2b_narrative_inputs():
    narrative_text = Path("forecast/inputs/mu_fy2026q4_narrative_ed1.yaml").read_text(encoding="utf-8")
    assumptions_text = Path("forecast/inputs/mu_fy2026q4_report_assumptions.yaml").read_text(encoding="utf-8")
    ex991_text = Path("logs/mu/fy2026q4/postprint/ex991.htm").read_text(encoding="utf-8")
    remarks_pdf = PdfReader("logs/mu/fy2026q4/postprint/remarks.pdf")
    remarks_text = normalize_pdf_text(" ".join(page.extract_text() or "" for page in remarks_pdf.pages))
    scored_text = Path("forecast/reports/mu_fy2026q4_SCORED.md").read_text(encoding="utf-8")
    manifest, narrative, assumptions, values = e2b_manifest_from_inputs(
        "forecast/scripts/mu_report/sources",
        narrative_text,
        assumptions_text,
        ex991_text,
        remarks_text,
        scored_text,
        {
            "SRC-EX991-FQ4FY26": "logs/mu/fy2026q4/postprint/ex991.htm",
            "SRC-REMARKS-FQ4FY26": "logs/mu/fy2026q4/postprint/remarks.pdf",
            "SRC-SCORED-FQ4FY26": "forecast/reports/mu_fy2026q4_SCORED.md",
            "SRC-RLE-FY27": "forecast/inputs/mu_fy2026q4_report_assumptions.yaml",
        },
    )
    return manifest, narrative, assumptions, values


def test_r14_fact_bindings_recompute_from_original_sources_and_rle():
    manifest, narrative, assumptions, values = _e2b_narrative_inputs()
    gate_narrative_contract(narrative, manifest)
    assert len(values) == 19
    assert values["guidance.revenue_mid.FQ1FY27"] == 61500
    assert values["derived.weekly_growth.FQ1FY27_guide"] == pytest.approx(0.2213164401108121)
    assert values["scored.opex_gap_pct.FQ4FY26"] == pytest.approx(3296 / 1860 - 1)
    assert values["scored.eps_error.FQ4FY26"] == pytest.approx(32.87 - 32.57)
    results = fy2027_scenario_results(assumptions)
    assert results["base"]["revenue"] == pytest.approx(274784.139585)
    assert results["bear"]["eps_gaap"] == pytest.approx(118.95, abs=0.005)
    assert results["base"]["eps_gaap"] == pytest.approx(169.05, abs=0.005)
    assert results["bull"]["eps_gaap"] == pytest.approx(199.52, abs=0.005)


def test_r14_fact_binding_mismatch_and_structure_fail_closed():
    manifest, narrative, assumptions, _ = _e2b_narrative_inputs()
    broken_value = yaml.safe_load(Path("forecast/inputs/mu_fy2026q4_narrative_ed1.yaml").read_text(encoding="utf-8"))
    broken_value["fact_bindings"][0]["value"] = 1
    with pytest.raises(ValueError, match="proposed value mismatch"):
        recompute_fact_bindings(
            broken_value,
            Path("logs/mu/fy2026q4/postprint/ex991.htm").read_text(encoding="utf-8"),
            Path("forecast/reports/mu_fy2026q4_SCORED.md").read_text(encoding="utf-8"),
            assumptions,
        )
    broken_shape = copy.deepcopy(narrative)
    broken_shape["risks"]["en"].pop()
    with pytest.raises(GateError, match="structure differs"):
        gate_narrative_contract(broken_shape, manifest)


def test_r14_ed1_sections_appendix_and_net_cash_label_are_ready_without_writes():
    manifest, narrative, assumptions, _ = _e2b_narrative_inputs()
    rendered = {}
    for locale in ("ko", "en"):
        rendered[locale], _ = render_markdown(
            manifest,
            locale,
            narrative=narrative,
            assumptions=assumptions,
        )
    gate_ed1_rendered(rendered)
    assert "공급 부족이 2026년보다 2027·2028년에 더 심해진다고 회사가 밝혔다." in rendered["ko"]
    assert "The company said supply will be tighter in 2027 and 2028 than in 2026." in rendered["en"]
    assert "순현금(SCA 예치금 미조정)" in rendered["ko"]
    assert "Net cash (not adjusted for SCA deposits)" in rendered["en"]
    assert "조정 전 순현금" not in rendered["ko"]
    assert "Unadjusted net cash" not in rendered["en"]
    assert "사전등록 증거는 대화 기록뿐이다" in rendered["ko"]
    assert "only evidence of pre-registration" in rendered["en"]


def test_renderer_parity_reference_contract(rendered_bundle):
    manifest, outputs = rendered_bundle
    logs = json.loads(outputs["references"].read_text(encoding="utf-8"))
    texts = {locale: outputs[f"md_{locale}"].read_text(encoding="utf-8") for locale in ("ko", "en")}
    gate_g15_parity(logs, manifest, texts)
    assert {entry["kind"] for entry in logs["ko"]} == {"table_cell", "text_placeholder", "chart_point"}


def test_xlsx_builder_uses_python_only():
    renderer = Path("forecast/scripts/mu_report/render.py").read_text(encoding="utf-8")
    assert "openpyxl" in renderer
    assert "codex-runtimes" not in renderer
    assert "subprocess" not in renderer
    assert not Path("forecast/scripts/mu_report/xlsx_builder.mjs").exists()


def test_xlsx_rebuild_is_byte_deterministic(tmp_path, monkeypatch):
    from datetime import datetime, timezone
    from types import SimpleNamespace

    class SaveClock(datetime):
        current = datetime(2026, 10, 5, 2, 0, 0)

        @classmethod
        def now(cls, tz=None):
            return cls.current.replace(tzinfo=tz)

    monkeypatch.setattr("openpyxl.writer.excel.datetime", SimpleNamespace(datetime=SaveClock, timezone=timezone))
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    manifest = fixture_manifest()
    _render_xlsx(manifest, first)
    SaveClock.current = datetime(2026, 10, 5, 2, 0, 5)
    _render_xlsx(manifest, second)
    assert first.read_bytes() == second.read_bytes()


def test_checked_in_historical_extracts_tie_to_companyfacts():
    source_dir = Path("forecast/scripts/mu_report/sources")
    annual = json.loads((source_dir / "annual_financials.json").read_text(encoding="utf-8"))["years"]
    companyfacts = json.loads((source_dir / "companyfacts_quarterly.json").read_text(encoding="utf-8"))["annual_crosscheck"]
    for year, statements in annual.items():
        gate_g3_financial_identities(statements["income_statement"])
        gate_g3b_balance_sheet(statements["balance_sheet"])
        assert all(len(statements["as_filed_rows"][name]) >= minimum for name, minimum in {"income_statement": 16, "balance_sheet": 15, "cash_flow": 25}.items())
        for metric in ("revenue", "operating_income", "net_income", "diluted_eps"):
            scale = 1 if metric == "diluted_eps" else 1_000_000
            assert statements["income_statement"][metric] == companyfacts[year][metric]["val"] / scale
        for metric in ("total_assets", "total_liabilities", "total_equity"):
            assert statements["balance_sheet"][metric] == companyfacts[year][metric]["val"] / 1_000_000
        assert statements["cash_flow"]["operating_cash_flow"] == companyfacts[year]["operating_cash_flow"]["val"] / 1_000_000
    business_units = json.loads((source_dir / "business_units.json").read_text(encoding="utf-8"))
    assert len(business_units["periods"]) == 8
    assert business_units["rewrites"] == []
    guidance = json.loads((source_dir / "guidance_history.json").read_text(encoding="utf-8"))
    assert guidance["count"] == 11
    freeze = json.loads((source_dir / "freeze_prereg.json").read_text(encoding="utf-8"))
    base_q4 = freeze["profile_path"]["base"][0]
    frozen_rows = freeze["prereg_a_display_rows"]
    gate_g2_freeze_display(
        {"revenue": frozen_rows[0]["display_text"]["base"], "operating_income": frozen_rows[8]["display_text"]["base"], "gaap_eps": frozen_rows[15]["display_text"]["base"]},
        {"revenue": f"{base_q4['revenue']:,.0f}", "operating_income": f"{base_q4['operating_income']:,.0f}", "gaap_eps": f"{base_q4['gaap_eps']:.2f}"},
    )
    historical = historical_manifest_from_extracts(source_dir)
    assert historical.fact("is.revenue.FY2025A").raw_value == 37378
    assert historical.fact("bs.total_assets.FY2023A").source_id == "SRC-10K-FY23"


def test_html_closes_table_before_following_heading():
    html = _html_from_markdown("| Metric | Value |\n|---|---|\n| A | B |\n\n## Method\n", "en")
    assert html.index("</table>") < html.index("<h2>Method</h2>")


# R4 E2-A' additions. Existing E2-A tests above remain unchanged.
from datetime import datetime, timedelta, timezone

import yaml

from forecast.scripts.mu_report.facts import Fact
from forecast.scripts.mu_report.gates import (
    gate_g12c_conflict,
    gate_g15_toc,
    gate_inventory_days,
    gate_market_data_box,
    gate_trailing_pb,
    qa_pdf_page_metadata,
)
from forecast.scripts.mu_report.inputs import InputGateError, load_conflict_confirmation
from forecast.scripts.mu_report.render import render_markdown
from forecast.scripts.mu_report.rle import inventory_days
from forecast.scripts.mu_report.valuation import trailing_pb


R4_FIXTURES = Path("forecast/tests/fixtures/mu_report")
R4_NOW = datetime(2099, 1, 1, 12, 30, tzinfo=timezone(timedelta(hours=9)))


def _conflict_rendered(record):
    rendered = {}
    for locale in ("ko", "en"):
        strings = yaml.safe_load(Path(f"forecast/scripts/mu_report/i18n/{locale}.yaml").read_text(encoding="utf-8"))
        stamp = record.confirmed_at_kst.isoformat()
        full = strings["disclaimer"]["conflict"].replace("{{conflict.confirmed_at_kst}}", stamp)
        short = strings["disclaimer"]["conflict_short"]
        rendered[locale] = {
            "cover": full,
            "ending": full,
            "pdf_pages": [short, short],
            "record_blob_sha": record.git_blob_sha,
            "conflict_short": short,
        }
    return rendered


def test_r4_inventory_days_formula_and_fail_closed_manifest_rules():
    assert inventory_days(80, 120, 200, 52) == pytest.approx(182.0)
    assert inventory_days(None, 120, 200, 52) == "UNAVAILABLE"
    assert inventory_days(80, 120, 0, 52) == "N/M"
    assert inventory_days(80, 120, 200, 52, same_scope=False) == "UNAVAILABLE"

    manifest = fixture_manifest()
    gate_inventory_days(manifest)
    weeks = manifest.facts["meta.period_weeks.FY2026.A-8K"]
    manifest.facts[weeks.fact_id] = replace(weeks, raw_value=52, display={"ko": "52", "en": "52"})
    with pytest.raises(GateError, match="53 weeks"):
        gate_inventory_days(manifest)
    manifest.facts[weeks.fact_id] = weeks

    fy27 = manifest.facts["ratio.inventory_days.FY2027.RLE"]
    manifest.facts[fy27.fact_id] = replace(fy27, raw_value=100, status="AVAILABLE", display={"ko": "100", "en": "100"})
    with pytest.raises(GateError, match="must be unavailable"):
        gate_inventory_days(manifest)
    manifest.facts[fy27.fact_id] = fy27

    manifest.facts["ratio.industry_inventory.FY2026.A"] = replace(
        manifest.facts["ratio.inventory_days.FY2026.A-8K"],
        fact_id="ratio.industry_inventory.FY2026.A",
    )
    with pytest.raises(GateError, match="industry inventory"):
        gate_inventory_days(manifest)


def test_r4_inventory_extract_coordinates_and_lineage():
    source_dir = Path("forecast/scripts/mu_report/sources")
    annual = json.loads((source_dir / "annual_financials.json").read_text(encoding="utf-8"))
    expected = {
        "FY2022A": (6663, "logs/_claude_scratch/mu-20230831.htm", [6, 8]),
        "FY2023A": (8387, "logs/_claude_scratch/mu-20230831.htm", [3, 5]),
        "FY2024A": (8875, "logs/_claude_scratch/mu-20250828.htm", [6, 8]),
        "FY2025A": (8355, "logs/_claude_scratch/mu-20250828.htm", [3, 5]),
    }
    for period, (value, path, columns) in expected.items():
        item = annual["inventory_balances"][period]
        assert item["caption"] == "Inventories"
        assert (item["value"], item["source_path"], item["coordinate"]["columns"]) == (value, path, columns)
        assert item["coordinate"]["table"] == 20 and item["coordinate"]["row"] == 7
    manifest = historical_manifest_from_extracts(source_dir)
    gate_inventory_days(manifest)
    ratio = manifest.facts["ratio.inventory_days.FY2025.A"]
    assert ratio.lineage["formula"] == "inventory_days_v1" and len(ratio.lineage["inputs"]) == 4
    manifest.facts[ratio.fact_id] = replace(ratio, lineage=None)
    with pytest.raises(ValueError, match="lineage"):
        gate_g14_provenance(manifest)


def test_r4_trailing_pb_and_market_data_contracts():
    assert trailing_pb(100, 500, 10) == 2
    assert trailing_pb("UNAVAILABLE", 500, 10) == "UNAVAILABLE"
    assert trailing_pb(None, 500, 10) == "UNAVAILABLE"
    assert trailing_pb(100, -500, 10) == "N/M"
    manifest = fixture_manifest()
    caption = {"price_date": "2099-01-01", "equity_date": "2098-12-31", "shares_date": "2098-12-30"}
    gate_trailing_pb(manifest, caption)
    with pytest.raises(GateError, match="shares_date"):
        gate_trailing_pb(manifest, {key: value for key, value in caption.items() if key != "shares_date"})
    pb = manifest.facts.pop("val.pb_trailing.FY2026.A-8K")
    manifest.facts["val.pb_trailing.FY2027.RLE"] = replace(pb, fact_id="val.pb_trailing.FY2027.RLE")
    with pytest.raises(GateError, match="RLE"):
        gate_trailing_pb(manifest, caption)

    for locale in ("ko", "en"):
        config = yaml.safe_load(Path(f"forecast/scripts/mu_report/i18n/{locale}.yaml").read_text(encoding="utf-8"))["market_data"]
        gate_market_data_box(config)
        with pytest.raises(GateError, match="exactly four"):
            gate_market_data_box({**config, "rows": [*config["rows"], {"label": "Volume", "value": "x"}]})
        forbidden = copy.deepcopy(config)
        forbidden["rows"][0]["label"] = "52-week volume"
        with pytest.raises(GateError, match="forbidden market-data"):
            gate_market_data_box(forbidden)


def test_r4_market_data_missing_fact_renders_only_that_cell_unavailable():
    manifest = fixture_manifest()
    shares = manifest.facts["market.shares_outstanding.CITED"]
    manifest.facts[shares.fact_id] = replace(
        shares,
        raw_value=None,
        display={"ko": "UNAVAILABLE", "en": "UNAVAILABLE"},
        status="UNAVAILABLE",
    )
    for locale in ("ko", "en"):
        markdown, _ = render_markdown(manifest, locale)
        assert "UNAVAILABLE" in markdown
        assert manifest.facts["market.price.2026-10-01.CITED"].display[locale] in markdown


@pytest.mark.parametrize(
    ("name", "loader_error", "gate_error"),
    [
        ("ok", False, False),
        ("held", False, True),
        ("expired", False, True),
        ("future", False, True),
        ("edition_mismatch", False, True),
        ("missing_key", True, False),
    ],
)
def test_r4_g12c_fixture_matrix(name, loader_error, gate_error):
    path = R4_FIXTURES / f"conflict_confirmation_{name}.yaml"
    if loader_error:
        with pytest.raises(InputGateError):
            load_conflict_confirmation(path, "ed1", R4_NOW)
        return
    record = load_conflict_confirmation(path, "ed1", R4_NOW)
    rendered = _conflict_rendered(record)
    if gate_error:
        with pytest.raises(GateError):
            gate_g12c_conflict(record, "ed1", R4_NOW, rendered)
    else:
        gate_g12c_conflict(record, "ed1", R4_NOW, rendered)


def test_r4_g12c_footer_and_shared_record_violations():
    record = load_conflict_confirmation(R4_FIXTURES / "conflict_confirmation_ok.yaml", "ed1", R4_NOW)
    rendered = _conflict_rendered(record)
    gate_g12c_conflict(record, "ed1", R4_NOW, rendered)
    missing_footer = copy.deepcopy(rendered)
    missing_footer["en"]["pdf_pages"][0] = "removed"
    with pytest.raises(GateError, match="conflict_short"):
        gate_g12c_conflict(record, "ed1", R4_NOW, missing_footer)
    different_record = copy.deepcopy(rendered)
    different_record["en"]["record_blob_sha"] = "f" * 40
    with pytest.raises(GateError, match="different confirmation"):
        gate_g12c_conflict(record, "ed1", R4_NOW, different_record)


def test_r4_g12c_rejects_timestamp_without_full_disclosure():
    record = load_conflict_confirmation(R4_FIXTURES / "conflict_confirmation_ok.yaml", "ed1", R4_NOW)
    rendered = _conflict_rendered(record)
    rendered["ko"]["cover"] = f"random text {record.confirmed_at_kst.isoformat()}"
    with pytest.raises(GateError, match="cover conflict disclosure"):
        gate_g12c_conflict(record, "ed1", R4_NOW, rendered)


def test_r4_g12c_rejects_caller_supplied_footer_text():
    record = load_conflict_confirmation(R4_FIXTURES / "conflict_confirmation_ok.yaml", "ed1", R4_NOW)
    rendered = _conflict_rendered(record)
    rendered["en"]["conflict_short"] = "x"
    rendered["en"]["pdf_pages"] = ["x", "x"]
    with pytest.raises(GateError, match="conflict_short"):
        gate_g12c_conflict(record, "ed1", R4_NOW, rendered)


def test_r4_toc_parity_and_local_anchor_violations(tmp_path):
    entries = {}
    for locale in ("ko", "en"):
        entries[locale] = yaml.safe_load(Path(f"forecast/scripts/mu_report/i18n/{locale}.yaml").read_text(encoding="utf-8"))["sections"]
    gate_g15_toc(entries)
    swapped = copy.deepcopy(entries)
    swapped["en"][0], swapped["en"][1] = swapped["en"][1], swapped["en"][0]
    with pytest.raises(GateError, match="order"):
        gate_g15_toc(swapped)
    bad_html = tmp_path / "bad_anchor.html"
    bad_html.write_text("<a href='#missing'>broken</a>", encoding="utf-8", newline="\n")
    with pytest.raises(GateError, match="anchor"):
        qa_html(bad_html)
    bad_md = tmp_path / "bad_anchor.md"
    bad_md.write_text("[broken](#missing)\n", encoding="utf-8", newline="\n")
    with pytest.raises(GateError, match="anchor"):
        qa_markdown(bad_md)


def test_r4_pdf_page_metadata_and_toc_violations():
    short = "Author does not hold MU shares - reconfirmed before publication."
    entries = [{"id": "alpha", "title": "Alpha"}, {"id": "beta", "title": "Beta"}]
    texts = [
        f"Cover {short} 1/3",
        f"Contents 1. Alpha ... 3 2. Beta ... 3 {short} 2/3",
        f"1. Alpha 2. Beta {short} 3/3",
    ]
    qa_pdf_page_metadata(texts, 3, short, entries)
    with pytest.raises(GateError, match="footer number"):
        qa_pdf_page_metadata([texts[0].replace("1/3", "1/4"), *texts[1:]], 3, short, entries)
    with pytest.raises(GateError, match="manifest"):
        qa_pdf_page_metadata(texts, 4, short, entries)
    with pytest.raises(GateError, match="target page"):
        qa_pdf_page_metadata([texts[0], texts[1].replace("1. Alpha ... 3", "1. Alpha ... 2"), texts[2]], 3, short, entries)


def test_r4_rendered_bundle_pages_toc_and_conflict_footer(rendered_bundle):
    _, outputs = rendered_bundle
    manifest = json.loads(outputs["manifest"].read_text(encoding="utf-8"))
    totals = manifest["metadata"]["document"]["total_pages"]
    assert totals == {"ko": 4, "en": 4}
    for locale in ("ko", "en"):
        texts = [page.extract_text() or "" for page in __import__("pypdf").PdfReader(outputs[f"pdf_{locale}"]).pages]
        short = "작성자 MU 주식 미보유 - 발행 직전 재확인." if locale == "ko" else "Author does not hold MU shares - reconfirmed before publication."
        assert sum(short in text for text in texts) == totals[locale]


# R5 DRYRUN additions. Existing E2-A and R4 tests above remain unchanged.
from forecast.scripts.mu_report.build import main as mu_report_main
from forecast.scripts.mu_report.facts import dryrun_manifest_from_extracts
from forecast.scripts.mu_report.gates import (
    DRYRUN_CONFLICT,
    gate_dryrun_boundary,
    gate_dryrun_placeholders,
    gate_dryrun_watermark,
    gate_g12c_dryrun,
    gate_theme,
)
from forecast.scripts.mu_report.theme import COLORS, contrast_ratio


def _dryrun_rendered():
    return {
        locale: {
            "cover": DRYRUN_CONFLICT[locale],
            "ending": DRYRUN_CONFLICT[locale],
            "pdf_pages": [DRYRUN_CONFLICT[locale], DRYRUN_CONFLICT[locale]],
        }
        for locale in ("ko", "en")
    }


def test_r5_dryrun_conflict_branch_and_violations():
    rendered = _dryrun_rendered()
    gate_g12c_dryrun(rendered)
    missing = copy.deepcopy(rendered)
    missing["en"]["pdf_pages"][0] = "removed"
    with pytest.raises(GateError, match="absent"):
        gate_g12c_dryrun(missing)
    real = copy.deepcopy(rendered)
    real["ko"]["pdf_pages"][0] += " 작성자 MU 주식 미보유 - 발행 직전 재확인."
    with pytest.raises(GateError, match="real conflict_short"):
        gate_g12c_dryrun(real)
    record = load_conflict_confirmation(R4_FIXTURES / "conflict_confirmation_ok.yaml", "ed1", R4_NOW)
    with pytest.raises(GateError, match="must not receive"):
        gate_g12c_dryrun(rendered, record)


def test_r5_watermark_and_theme_violations(monkeypatch):
    gate_dryrun_watermark(["PRE-PRINT DRY RUN", "PRE-PRINT DRY RUN"], "PRE-PRINT DRY RUN")
    with pytest.raises(GateError, match="watermark"):
        gate_dryrun_watermark(["PRE-PRINT DRY RUN", "removed"], "PRE-PRINT DRY RUN")
    ratios = gate_theme("<html></html>", ["plain text"])
    assert min(ratios.values()) >= 4.5
    with pytest.raises(GateError, match="color value"):
        gate_theme("<html></html>", ["accent: '#164E87'"])
    with pytest.raises(GateError, match="logo"):
        gate_theme("<img src='micron-logo.png'>", ["plain text"])
    monkeypatch.setitem(COLORS, "primary", "#DCE9F5")
    with pytest.raises(GateError, match="contrast"):
        gate_theme("<html></html>", ["plain text"])


def test_r5_theme_contrast_calculation():
    assert contrast_ratio(COLORS["ink"], COLORS["paper"]) > 4.5
    assert contrast_ratio("#FFFFFF", "#FFFFFF") == 1.0


def test_r5_dryrun_manifest_has_actuals_and_no_fixture_values():
    manifest = dryrun_manifest_from_extracts("forecast/scripts/mu_report/sources")
    assert manifest.metadata["layer"] == "DRYRUN"
    assert manifest.metadata["fixture_values"] is False
    assert manifest.fact("is.revenue.FY2025A").status == "AVAILABLE"
    assert manifest.fact("is.revenue.FY2026Q4.A-8K").status == "UNAVAILABLE"
    assert manifest.fact("market.price.2026-10-01.CITED").status == "UNAVAILABLE"


def test_r5_dryrun_placeholder_value_injections():
    manifest = dryrun_manifest_from_extracts("forecast/scripts/mu_report/sources")
    heatmap = {"07_valuation_heatmap": {"fact_ids": [], "values": []}}
    gate_dryrun_placeholders(manifest, heatmap)
    price = manifest.facts["market.price.2026-10-01.CITED"]
    manifest.facts[price.fact_id] = replace(price, raw_value=123.0, status="AVAILABLE", display={"ko": "123", "en": "123"})
    with pytest.raises(GateError, match="placeholder contains"):
        gate_dryrun_placeholders(manifest, heatmap)
    manifest.facts[price.fact_id] = price
    with pytest.raises(GateError, match="heatmap"):
        gate_dryrun_placeholders(manifest, {"07_valuation_heatmap": {"fact_ids": ["x"], "values": [1]}})


def test_r5_dryrun_boundary_and_fixture_injection(tmp_path):
    manifest = dryrun_manifest_from_extracts("forecast/scripts/mu_report/sources")
    output = Path("logs/_mu_report_runs/dryrun_test").resolve()
    reader = EvidenceReader("E2-A")
    for pin in reader.pins["E2-A"]:
        reader.read_bytes(pin["path"])
    gate_dryrun_boundary(output, manifest, reader.audit_payload())
    fixture = fixture_manifest()
    with pytest.raises(GateError, match="layer"):
        gate_dryrun_boundary(output, fixture, reader.audit_payload())
    with pytest.raises(GateError, match="escaped"):
        gate_dryrun_boundary(tmp_path / "dryrun_test", manifest, reader.audit_payload())


def test_r5_dryrun_cli_rejects_wrong_edition(monkeypatch):
    monkeypatch.setattr("sys.argv", ["build.py", "--phase", "DRYRUN", "--edition", "1"])
    with pytest.raises(SystemExit):
        mu_report_main()


def test_r5_g13c_semantic_contract_and_violation(tmp_path):
    from forecast.scripts.mu_report.charts import dryrun_specs
    import yaml

    manifest = dryrun_manifest_from_extracts("forecast/scripts/mu_report/sources")
    strings = yaml.safe_load(Path("forecast/scripts/mu_report/i18n/en.yaml").read_text(encoding="utf-8"))
    charts = render_charts(manifest, dryrun_specs(manifest, "en", strings), tmp_path, "en")
    gate_g13c_chart_semantics(manifest, charts)
    broken = copy.deepcopy(charts)
    broken["04_beat_history"]["chart_type"] = "line"
    with pytest.raises(GateError, match="chart type"):
        gate_g13c_chart_semantics(manifest, broken)
    broken = copy.deepcopy(charts)
    broken["08_cash_flow_capex_net_cash"]["fact_ids"][0] = "is.revenue.FY2025A"
    with pytest.raises(GateError, match="forbidden fact family"):
        gate_g13c_chart_semantics(manifest, broken)


def test_r6_g13c_rejects_mixed_period_units_reverse_order_and_omission(tmp_path):
    from forecast.scripts.mu_report.charts import dryrun_specs
    import yaml

    manifest = dryrun_manifest_from_extracts("forecast/scripts/mu_report/sources")
    strings = yaml.safe_load(Path("forecast/scripts/mu_report/i18n/en.yaml").read_text(encoding="utf-8"))
    charts = render_charts(manifest, dryrun_specs(manifest, "en", strings), tmp_path, "en")

    mixed = copy.deepcopy(charts)
    mixed["08_cash_flow_capex_net_cash"]["periods"][0] = "FQ3-26"
    with pytest.raises(GateError, match="mixes annual and quarterly"):
        gate_g13c_chart_semantics(manifest, mixed)

    reversed_series = copy.deepcopy(charts)
    reversed_series["08_cash_flow_capex_net_cash"]["series_periods"]["cf.fcf_adjusted."].reverse()
    with pytest.raises(GateError, match="not chronological"):
        gate_g13c_chart_semantics(manifest, reversed_series)

    omitted = copy.deepcopy(charts)
    chart = omitted["08_cash_flow_capex_net_cash"]
    missing_id = "cf.net_capex.FY2024A"
    index = chart["fact_ids"].index(missing_id)
    for key in ("fact_ids", "values", "periods"):
        chart[key].pop(index)
    chart["series_periods"]["cf.net_capex."].remove("FY2024A")
    with pytest.raises(GateError, match="omits manifest values"):
        gate_g13c_chart_semantics(manifest, omitted)


def test_r6_prepared_remarks_categories_are_exact_source_substrings():
    remarks = json.loads(Path("forecast/scripts/mu_report/sources/prepared_remarks.json").read_text(encoding="utf-8"))
    reader = EvidenceReader("E2-A")
    pdf = PdfReader(BytesIO(reader.read_bytes(remarks["source"]["path"])))
    normalized_source = normalize_pdf_text(" ".join(page.extract_text() or "" for page in pdf.pages))
    values = [value for metrics in remarks["product_metrics"].values() for value in metrics.values()]
    gate_categorical_verbatim(values, normalized_source)

    injected = [value.replace("mid-single", "mid -single") for value in values]
    assert injected != values
    with pytest.raises(GateError, match="exact source substring"):
        gate_categorical_verbatim(injected, normalized_source)


def test_r5_g15b_localized_assets_and_ui_violations():
    ko = {"01": {"path": "01_ko.png", "ui_text": ["분기 매출", "회계 분기", "USD million", "GAAP"]}}
    en = {"01": {"path": "01_en.png", "ui_text": ["Quarterly revenue"]}}
    texts = {"ko": "| 항목 | 값 | 기준 |", "en": "| Metric | Value | Basis |"}
    gate_g15b_localized_ui({"ko": ko, "en": en}, texts)
    bad = copy.deepcopy(ko)
    bad["01"]["ui_text"].append("Revenue")
    with pytest.raises(GateError, match="English UI"):
        gate_g15b_localized_ui({"ko": bad, "en": en}, texts)
    bad = copy.deepcopy(ko)
    bad["01"]["path"] = "01_en.png"
    with pytest.raises(GateError, match="locale-specific"):
        gate_g15b_localized_ui({"ko": bad, "en": en}, texts)


def test_r16_scored_source_and_annual_rle_are_bound_to_inputs():
    manifest, _, assumptions, _ = _e2b_narrative_inputs()
    scored_text = Path("forecast/reports/mu_fy2026q4_SCORED.md").read_text(encoding="utf-8")
    gate_g3f_scored_source(manifest, scored_text)
    changed = copy.deepcopy(manifest)
    fact = changed.facts["scored.lever.total.FQ4FY26"]
    changed.facts[fact.fact_id] = replace(fact, raw_value=0.301)
    with pytest.raises(GateError, match="G-3f value mismatch"):
        gate_g3f_scored_source(changed, scored_text)
    annual = annual_rle_results(assumptions)
    assert set(annual) == {"bear", "base", "bull"}
    assert annual["base"][2028]["net_cash_unadjusted"] > annual["base"][2027]["net_cash_unadjusted"]


def test_r17_f1_scored_values_match_signed_displays():
    manifest, narrative, assumptions, _ = _e2b_narrative_inputs()
    expected = {
        "revenue": (0.843, "+$0.84"), "op_margin": (-0.730, "−$0.73"),
        "op_to_ni": (0.103, "+$0.10"), "shares": (0.086, "+$0.09"), "total": (0.302, "+$0.30"),
    }
    for name, (value, display) in expected.items():
        fact = manifest.fact(f"scored.lever.{name}.FQ4FY26")
        assert fact.raw_value == pytest.approx(value)
        assert fact.unit == "USD_per_share"
        assert fact.display == {"ko": display, "en": display}
    assert manifest.fact("scored.labels_hit.FQ4FY26").raw_value == 4
    assert manifest.fact("scored.labels_hit.FQ4FY26").display == {"ko": "4", "en": "4"}
    for locale in ("ko", "en"):
        text, _ = render_markdown(manifest, locale, narrative=narrative, assumptions=assumptions)
        assert "+$0.84 / −$0.73 / +$0.10 / +$0.09 = +$0.30" in text
        assert "4/4 " + ("적중" if locale == "ko" else "hit") in text
        assert "3,296" in text and "1,860" in text


def test_r17_f2_rle_cash_flow_signs_match_displays_and_totals():
    manifest, _, assumptions, _ = _e2b_narrative_inputs()
    annual = annual_rle_results(assumptions)
    for scenario in ("bear", "base", "bull"):
        for year in (2027, 2028):
            values = annual[scenario][year]
            for family, metric in (("is", "tax"), ("cf", "working_capital_investment"), ("cf", "dividends")):
                fact = manifest.fact(f"{family}.{metric}.FY{year}E.{scenario}")
                assert fact.raw_value == pytest.approx(-values[metric])
                assert fact.unit == "USD_million"
                expected = "0" if values[metric] == 0 else f"{-values[metric]:,.0f}"
                assert fact.display == {"ko": expected, "en": expected}
            cfo = values["net_income"] + values["d_and_a"] + values["sbc"] + manifest.fact(f"cf.working_capital_investment.FY{year}E.{scenario}").raw_value
            assert cfo == pytest.approx(manifest.fact(f"cf.operating_cash_flow.FY{year}E.{scenario}").raw_value)
    for fact_id, display in {
        "is.tax.FY2027E.base": "-30,794", "is.tax.FY2028E.base": "-29,553",
        "cf.working_capital_investment.FY2027E.base": "-4,882", "cf.working_capital_investment.FY2028E.base": "0",
        "cf.dividends.FY2027E.base": "-690", "cf.dividends.FY2028E.base": "-690",
    }.items():
        assert manifest.fact(fact_id).display["ko"] == display


def test_r17_f5_market_data_labels_units_and_formula_match_facts():
    manifest, narrative, assumptions, _ = _e2b_narrative_inputs()
    shares = manifest.fact("market.shares_outstanding.CITED")
    price = manifest.fact("market.price.2026-10-01.CITED")
    cap = manifest.fact("market.market_cap.2026-10-01.CALCULATED")
    assert shares.raw_value == 1147
    assert shares.period == "FY2026Q4"
    assert shares.label == "FQ4 diluted weighted-average shares"
    assert shares.unit == "million_shares"
    assert shares.display == {"ko": "1,147M", "en": "1,147M"}
    assert cap.raw_value == pytest.approx(price.raw_value * shares.raw_value)
    assert cap.unit == "USD_million"
    assert cap.display == {"ko": "USD 1,258,706 million", "en": "USD 1,258,706 million"}
    for locale in ("ko", "en"):
        text, _ = render_markdown(manifest, locale, narrative=narrative, assumptions=assumptions)
        assert "| 1,258,706 |" in text
        assert "USD million" in text
        assert ("희석 가중평균 주식수(FQ4)" if locale == "ko" else "Diluted weighted-average shares (FQ4)") in text
        assert ("USD million · 주가 × 주식수" if locale == "ko" else "USD million · price × shares") in text
        assert "million shares · FQ4 A-8K" in text


def test_r17_table_cell_english_and_python_tuple_are_rejected():
    charts = {"ko": {}, "en": {}}
    for cell in (
        "첫 분기 = M1; 이후 -0.5 percentage point 경로 retain point path as preregistered sensitivity",
        "FQ4 actual net capex * 1.05, flat; replace if numerical FY2027 guidance is issued",
    ):
        with pytest.raises(GateError, match="KO table cell"):
            gate_g15b_localized_ui(charts, {"ko": f"| 원안 | {cell} |", "en": "| Original | English |"})
    with pytest.raises(GateError, match="python_tuple"):
        gate_g24_raw_markup({"ko": ["잔액 (786,)"], "en": ["clean"]})


def test_r17_fq4_categories_guidance_and_bridge_are_source_bound():
    manifest, narrative, assumptions, _ = _e2b_narrative_inputs()
    source = PdfReader("logs/mu/fy2026q4/postprint/remarks.pdf").pages[6].extract_text()
    normalized = normalize_pdf_text(source)
    values = [fact.raw_value for key, fact in manifest.facts.items() if key.startswith("qual.") and ".FQ4-26." in key]
    gate_categorical_verbatim(values, normalized)
    assert manifest.fact("guidance.gm_gaap.FQ1FY27").display["en"] == "85.95%"
    assert manifest.fact("guidance.gm_nongaap.FQ1FY27").display["en"] == "86.25%"
    assert manifest.fact("guidance.opex_gaap.FQ1FY27").raw_value == 2310
    assert manifest.fact("guidance.opex_nongaap.FQ1FY27").raw_value == 2060
    assert manifest.fact("bridge.opex.delta.FY2026Q3_to_Q4").raw_value == pytest.approx(3296 - 1738)
    contribution = manifest.fact("bridge.gm.delta.FY2026Q3_to_Q4").raw_value + manifest.fact("bridge.opex_ratio.contribution.FY2026Q3_to_Q4").raw_value
    actual_change = manifest.fact("ratio.operating_margin.FY2026Q4.A").raw_value - manifest.fact("ratio.operating_margin.FY2026Q3.A").raw_value
    assert contribution == pytest.approx(actual_change)
    for locale in ("ko", "en"):
        text, _ = render_markdown(manifest, locale, narrative=narrative, assumptions=assumptions)
        assert "NOT_IN_SOURCE" not in text
        specs = e2b_specs(manifest, locale, yaml.safe_load(Path(f"forecast/scripts/mu_report/i18n/{locale}.yaml").read_text(encoding="utf-8")))
        price_bit = next(spec for spec in specs if spec.chart_id == "10_price_bit_ranges")
        assert {point.label for point in price_bit.points} == {"FQ3-26", "FQ4-26"}
        assert "(786,)" not in text
        assert "E·D / J" in text
        assert "85.95%" in text and "86.25%" in text


@pytest.mark.parametrize("locale", ["ko", "en"])
def test_r18_required_rle_cells_and_appendix(locale):
    manifest, narrative, assumptions, _ = _e2b_narrative_inputs()
    text, refs = render_markdown(manifest, locale, narrative=narrative, assumptions=assumptions)
    for metric, family, row in (("pretax", "is", "pretax_income"), ("net_income", "cf", "net_income")):
        for year in (2027, 2028):
            fact_id = f"is.{metric}.FY{year}E.base"
            position = f"table:{family}:{row}:fy{str(year)[-2:]}:0"
            assert any(ref["position"] == position and ref["fact_id"] == fact_id for ref in refs)
    gate_g23_availability(manifest, {}, {locale: text})
    labels = ("세전이익", "순이익") if locale == "ko" else ("Pretax income", "Net income")
    for label in labels:
        lines = text.splitlines()
        candidates = [i for i, line in enumerate(lines) if line.startswith(f"| {label} |")]
        index = candidates[0 if label == labels[0] else 1]
        for cell_index in (-3, -2):
            broken_lines = lines.copy()
            cells = broken_lines[index].split("|")
            cells[cell_index] = " UNAVAILABLE "
            broken_lines[index] = "|".join(cells)
            with pytest.raises(GateError, match="required rendered cell"):
                gate_g23_availability(manifest, {}, {locale: "\n".join(broken_lines)})
    dividend_rule = "분기 주당 배당×4분기×희석주식수" if locale == "ko" else "quarterly dividend per share × 4 quarters × diluted shares"
    assert dividend_rule in text
    assert "quarterly_dps * 4 * S1" not in text
    assert ("분기 주당 배당: " if locale == "ko" else "Quarterly dividend per share: ") in text
    assert "quarterly_dps.value=" not in text
    html = _html_from_markdown(text, locale)
    assert dividend_rule in html
    assert "quarterly_dps.value=" not in html


@pytest.mark.parametrize("metric,year", [(metric, year) for metric in ("pretax", "net_income") for year in (2027, 2028)])
def test_r18_g23_requires_each_rle_fact(metric, year):
    manifest, _, _, _ = _e2b_narrative_inputs()
    fact_id = f"is.{metric}.FY{year}E.base"
    manifest.facts.pop(fact_id)
    with pytest.raises(GateError, match="required available"):
        gate_g23_availability(manifest, {})


def test_r16_new_gates_red_and_green_contracts():
    manifest, narrative, assumptions, _ = _e2b_narrative_inputs()
    charts = {"ko": {}, "en": {}}
    texts = {locale: render_markdown(manifest, locale, narrative=narrative, assumptions=assumptions)[0] for locale in ("ko", "en")}
    audit = gate_g23_availability(manifest, charts, texts)
    assert audit["unavailable_count"] >= 2
    broken = copy.deepcopy(manifest)
    fact = broken.facts["market.price.2026-10-01.CITED"]
    broken.facts[fact.fact_id] = replace(fact, raw_value=None, status="UNAVAILABLE")
    with pytest.raises(GateError, match="required available"):
        gate_g23_availability(broken, charts)

    gate_g24_raw_markup({"ko": ["정상 PDF"], "en": ["clean PDF"]})
    with pytest.raises(GateError, match="raw markup"):
        gate_g24_raw_markup({"ko": ["#### 원시 표식"], "en": ["clean PDF"]})

    dated = {
        "ko": "발행일: 2026-10-04\n자료 기준일: 2026-10-01(주가) · 2026-09-30(실적)\n정보 컷오프: 2026-10-04 KST",
        "en": "Issue date: 2026-10-04\nData as of: 2026-10-01 (price) · 2026-09-30 (results)\nInformation cutoff: 2026-10-04 KST",
    }
    gate_g25_cover_dates(dated)
    with pytest.raises(GateError, match="cover date"):
        gate_g25_cover_dates({**dated, "ko": dated["ko"].replace("2026-10-04", "UNAVAILABLE", 1)})

    ko = {"01": {"path": "01_ko.png", "ui_text": ["분기 매출"]}}
    en = {"01": {"path": "01_en.png", "ui_text": ["Quarterly revenue"]}}
    with pytest.raises(GateError, match="English sentence"):
        gate_g15b_localized_ui(
            {"ko": ko, "en": en},
            {"ko": "percentage range driven by tight industry conditions", "en": "English"},
        )


def test_r16_e2b_markdown_and_chart_contracts_pass_together(tmp_path):
    import yaml

    manifest, narrative, assumptions, _ = _e2b_narrative_inputs()
    charts, texts, references = {}, {}, {}
    for locale in ("ko", "en"):
        strings = yaml.safe_load(Path(f"forecast/scripts/mu_report/i18n/{locale}.yaml").read_text(encoding="utf-8"))
        charts[locale] = render_charts(manifest, e2b_specs(manifest, locale, strings), tmp_path, locale, write_manifest=False)
        texts[locale], references[locale] = render_markdown(
            manifest,
            locale,
            chart_manifest=charts[locale],
            narrative=narrative,
            assumptions=assumptions,
            render_date="2026-10-04",
        )
        gate_g13c_chart_semantics(manifest, charts[locale])
    gate_g15_parity(references, manifest, texts)
    gate_g15b_localized_ui(charts, texts)
    gate_g23_availability(manifest, charts, texts)
    gate_g25_cover_dates(texts)


@pytest.mark.parametrize("label", ["FIXTURE", "DRY RUN", "E2-A fixture", "prefixFIXTUREsuffix"])
def test_r19_g24_rejects_labels_in_every_xlsx_sheet(tmp_path, label):
    from openpyxl import Workbook

    workbook = Workbook()
    hidden = workbook.create_sheet("Hidden audit")
    hidden.sheet_state = "hidden"
    hidden["D7"] = f"source / {label}"
    path = tmp_path / "labels.xlsx"
    workbook.save(path)
    workbook.close()
    with pytest.raises(GateError, match="Hidden audit!D7"):
        gate_g24_raw_markup({"ko": ["정상 PDF"], "en": ["clean PDF"]}, path)


def test_r19_ed1_xlsx_uses_cover_strings_without_changing_facts(tmp_path):
    import yaml
    from openpyxl import load_workbook

    manifest, _, _, _ = _e2b_narrative_inputs()
    record = load_conflict_confirmation(R4_FIXTURES / "conflict_confirmation_ok.yaml", "ed1", R4_NOW)
    path = tmp_path / "ed1.xlsx"
    _render_xlsx(manifest, path, conflict=record, render_date="2026-10-05")
    workbook = load_workbook(path, data_only=False)
    try:
        for locale, column in (("ko", 2), ("en", 5)):
            strings = yaml.safe_load(Path(f"forecast/scripts/mu_report/i18n/{locale}.yaml").read_text(encoding="utf-8"))
            expected = [
                strings["ed1"]["ribbon"],
                strings["ed1"]["cover"]["author"],
                strings["ed1"]["cover"]["issued"].replace("{{render_date}}", "2026-10-05"),
                strings["ed1"]["cover"]["as_of"],
                strings["ed1"]["cover"]["cutoff"],
                strings["ed1"]["cover"]["completeness"],
                strings["disclaimer"]["not_advice"],
                strings["disclaimer"]["third_party"],
                _conflict_rendered(record)[locale]["cover"],
            ]
            assert [workbook["Summary"].cell(row, column).value for row in range(4, 13)] == expected
        actual = list(workbook["Facts"].iter_rows(min_row=2, max_col=7, values_only=True))
        assert len(actual) == len(manifest.facts)
        for row, fact in zip(actual, manifest.facts.values()):
            assert (row[0], *row[2:]) == (fact.fact_id, fact.unit, fact.period, fact.basis, fact.source_id, fact.status)
            if isinstance(fact.raw_value, (int, float)):
                # XLSX numeric serialization retains 15 significant digits.
                assert row[1] == pytest.approx(fact.raw_value, rel=1e-15, abs=1e-9)
            else:
                assert row[1] == fact.raw_value
        assert workbook["Facts"]["K2"].value == strings["ed1"]["cover"]["as_of"]
        qa_xlsx_rows({
            sheet.title: [str(cell.value) for row in sheet for cell in row if cell.value is not None]
            for sheet in workbook
        })
    finally:
        workbook.close()
    assert gate_g24_raw_markup({"ko": ["정상 PDF"], "en": ["clean PDF"]}, path)["xlsx"] == []


def test_r19_ed1_xlsx_requires_confirmed_cover_context(tmp_path):
    manifest, _, _, _ = _e2b_narrative_inputs()
    with pytest.raises(ValueError, match="confirmed disclosure and render date"):
        _render_xlsx(manifest, tmp_path / "ed1.xlsx")
