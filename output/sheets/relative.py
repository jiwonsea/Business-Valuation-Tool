"""Sheet: Relative Valuation -- diagnostic multiples, peer-median comparison, justified verdicts.

Renders the diagnostic relative-valuation layer (ValuationResult.relative_valuation):
  * Tier 1/2 ratios (P/E, Fwd P/E, P/B, EV/EBITDA, EV/Sales, Div Yield, PEG, PEGY)
  * Peer-median comparison (company EV/EBITDA vs peer_stats median)
  * Justified-multiple verdicts (actual vs fundamental-justified P/E, P/B)

NA/CAUTION rows carry their reason in the 비고 column. This is a diagnostic layer,
not a primary method -- see engine/relative_metrics.py.
"""

from statistics import median

from openpyxl.utils import get_column_letter

from ._ctx import Ctx
from ..excel_styles import (
    YELLOW_FILL,
    GREEN_FILL,
    RED_FILL,
    SECTION_FONT,
    TITLE_FONT,
    MULT_FMT,
    style_header_row,
    write_cell,
)

_PCT_NAMES = {"Div Yield", "PEG", "PEGY"}


def _peer_ev_ebitda_median(ctx: Ctx):
    """Aggregate peer EV/EBITDA median (median of per-segment medians)."""
    vals = [
        ps.ev_ebitda_median
        for ps in ctx.result.peer_stats
        if getattr(ps, "ev_ebitda_median", 0)
    ]
    return median(vals) if vals else None


def sheet_relative(ctx: Ctx):
    rv = ctx.result.relative_valuation
    if not rv or not rv.ratios:
        return

    ws = ctx.wb.create_sheet("Relative Valuation")
    ws.sheet_properties.tabColor = "8E44AD"
    for c, w in enumerate([22, 14, 12, 52], 1):
        ws.column_dimensions[get_column_letter(c)].width = w

    write_cell(ws, 1, 1, "상대가치 진단 (Relative Valuation Diagnostics)", font=TITLE_FONT)

    # ── Section 1: diagnostic ratios ──
    r = 3
    write_cell(ws, r, 1, "진단 배수", font=SECTION_FONT)
    r += 1
    for c, h in enumerate(["지표", "값", "상태", "비고"], 1):
        write_cell(ws, r, c, h)
    style_header_row(ws, r, 4)

    for m in rv.ratios:
        r += 1
        write_cell(ws, r, 1, m.name)
        if m.value is None:
            write_cell(ws, r, 2, "N/A")
        elif m.name in _PCT_NAMES:
            unit = "%" if m.name == "Div Yield" else ""
            write_cell(ws, r, 2, f"{m.value:.2f}{unit}")
        else:
            write_cell(ws, r, 2, m.value, fmt=MULT_FMT)
        status = (m.status or "").upper()
        fill = GREEN_FILL if m.status == "ok" else (YELLOW_FILL if m.status == "caution" else RED_FILL)
        write_cell(ws, r, 3, status, fill=fill)
        write_cell(ws, r, 4, m.note or "")

    # ── Section 2: peer-median comparison ──
    peer_med = _peer_ev_ebitda_median(ctx)
    ev_ebitda = next((m for m in rv.ratios if m.name == "EV/EBITDA"), None)
    if peer_med is not None and ev_ebitda is not None and ev_ebitda.value is not None:
        r += 2
        write_cell(ws, r, 1, "피어 median 대비", font=SECTION_FONT)
        r += 1
        for c, h in enumerate(["지표", "자사", "피어 median", "판정"], 1):
            write_cell(ws, r, c, h)
        style_header_row(ws, r, 4)
        r += 1
        write_cell(ws, r, 1, "EV/EBITDA")
        write_cell(ws, r, 2, ev_ebitda.value, fmt=MULT_FMT)
        write_cell(ws, r, 3, round(peer_med, 2), fmt=MULT_FMT, fill=GREEN_FILL)
        if ev_ebitda.value > peer_med * 1.1:
            verdict, fill = "피어 대비 프리미엄", RED_FILL
        elif ev_ebitda.value < peer_med * 0.9:
            verdict, fill = "피어 대비 할인", GREEN_FILL
        else:
            verdict, fill = "피어 수준", YELLOW_FILL
        write_cell(ws, r, 4, verdict, fill=fill)

    # ── Section 3: justified-multiple verdicts ──
    if rv.verdicts:
        r += 2
        write_cell(ws, r, 1, "정당배수 대비 판정", font=SECTION_FONT)
        r += 1
        for c, h in enumerate(["지표", "실제", "정당", "괴리율", "판정"], 1):
            write_cell(ws, r, c, h)
        style_header_row(ws, r, 5)
        for v in rv.verdicts:
            r += 1
            write_cell(ws, r, 1, v.name)
            write_cell(ws, r, 2, v.actual if v.actual is not None else "N/A",
                       fmt=MULT_FMT if v.actual is not None else None)
            write_cell(ws, r, 3, v.justified if v.justified is not None else "N/A",
                       fmt=MULT_FMT if v.justified is not None else None)
            write_cell(ws, r, 4, f"{v.gap_pct:+.1f}%" if v.gap_pct is not None else "—")
            fill = {
                "저평가": GREEN_FILL,
                "고평가": RED_FILL,
                "적정": YELLOW_FILL,
            }.get(v.verdict)
            write_cell(ws, r, 5, v.verdict, fill=fill)

    # ── Footnote: growth source ──
    if rv.growth_pct is not None:
        r += 2
        write_cell(
            ws, r, 1,
            f"* PEG/정당배수 성장률: {rv.growth_pct:.1f}% ({rv.growth_source})",
        )
    r += 1
    write_cell(ws, r, 1, "* 진단 레이어 — 주가치 산정의 1차 방법론이 아닌 정합성 점검용")
