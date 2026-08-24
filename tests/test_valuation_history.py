"""Phase 1 persisted valuation-history contract tests."""

from __future__ import annotations

from datetime import date
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook

from db import repository
from output.sheets.history import sheet_valuation_history
from schemas.history import Provenance, ValuationHistoryRecord


class _Response:
    def __init__(self, data):
        self.data = data


class _Query:
    def __init__(self, data):
        self.data = data
        self.calls = []

    def _call(self, name, *args, **kwargs):
        self.calls.append((name, args, kwargs))
        return self

    def select(self, *a, **kw):
        return self._call("select", *a, **kw)

    def order(self, *a, **kw):
        return self._call("order", *a, **kw)

    def limit(self, *a, **kw):
        return self._call("limit", *a, **kw)

    def eq(self, *a, **kw):
        return self._call("eq", *a, **kw)

    def insert(self, *a, **kw):
        return self._call("insert", *a, **kw)

    def upsert(self, *a, **kw):
        return self._call("upsert", *a, **kw)

    def execute(self):
        return _Response(self.data)


class _Client:
    def __init__(self, data):
        self.query = _Query(data)

    def table(self, name):
        self.query.calls.append(("table", (name,), {}))
        return self.query


def test_history_query_uses_ticker_and_market_and_maps_saved_values():
    client = _Client(
        [
            {
                "company_name": "Same Name",
                "ticker": "AAA",
                "market": "US",
                "analysis_date": "2026-07-16",
                "weighted_value": 120,
                "market_price": 150,
                "gap_ratio": -0.2,
                "wacc_pct": 9.5,
                "valuation_method": "sotp",
                "result_data": {
                    "quality": {"grade": "B"},
                    "valuation_bucket": "plain_operating",
                },
                "created_at": "2026-07-16T01:00:00Z",
            }
        ]
    )
    with patch.object(repository, "get_client", return_value=client):
        records = repository.list_valuation_history(
            ticker="AAA", market="US", company_name="Same Name"
        )

    eq_calls = [call[1] for call in client.query.calls if call[0] == "eq"]
    assert ("ticker", "AAA") in eq_calls
    assert ("market", "US") in eq_calls
    assert not any(args[0] == "company_name" for args in eq_calls)
    assert records[0].gap_pct == -20.0
    assert records[0].quality_grade == "B"
    assert records[0].provenance["weighted_value"] is Provenance.DERIVED


def test_save_valuation_is_append_first():
    client = _Client([{"id": "valuation-1"}])
    vi = SimpleNamespace(
        draft=False,
        company=SimpleNamespace(
            name="Company",
            ticker="AAA",
            market="US",
            legal_status="listed",
            analysis_date=date(2026, 7, 16),
        ),
        base_year=2025,
        model_dump=lambda **_: {"input": True},
    )
    result = SimpleNamespace(
        draft=False,
        primary_method="sotp",
        total_ev=100,
        weighted_value=90,
        wacc=SimpleNamespace(wacc=10.0),
        market_comparison=None,
        model_dump=lambda **_: {"result": True},
    )
    with patch.object(repository, "get_client", return_value=client):
        assert repository.save_valuation(vi, result) == "valuation-1"
    assert any(call[0] == "insert" for call in client.query.calls)
    assert not any(call[0] == "upsert" for call in client.query.calls)


def _ctx():
    return SimpleNamespace(
        wb=Workbook(),
        vi=SimpleNamespace(
            company=SimpleNamespace(ticker="AAA", market="US", name="Company")
        ),
    )


def test_history_sheet_empty_state_has_no_chart():
    ctx = _ctx()
    with patch("db.repository.list_valuation_history", return_value=[]):
        sheet_valuation_history(ctx)
    ws = ctx.wb["Valuation History"]
    assert ws["A4"].value == "이력 부족(N=0)"
    assert ws._charts == []


def test_history_sheet_renders_persisted_values_without_engine_calls():
    ctx = _ctx()
    records = [
        ValuationHistoryRecord(
            ticker="AAA",
            market="US",
            company_name="Company",
            analysis_date=date(2026, 7, 16),
            weighted_value=120,
            market_price=150,
            gap_pct=-20,
            wacc_pct=9.5,
        )
    ]
    with patch("db.repository.list_valuation_history", return_value=records):
        sheet_valuation_history(ctx)
    ws = ctx.wb["Valuation History"]
    assert ws["B5"].value == 120
    assert ws["C5"].value == 150
    assert len(ws._charts) == 1


def test_all_three_entry_paths_use_shared_save_contract():
    root = Path(__file__).resolve().parents[1]
    cli = (root / "cli.py").read_text(encoding="utf-8")
    orchestrator = (root / "orchestrator.py").read_text(encoding="utf-8")
    auto = (root / "pipeline" / "profile_generator.py").read_text(encoding="utf-8")
    assert "_save_to_db(vi, result, args.profile)" in cli
    assert "_save_to_db(vi, result, profile_path)" in orchestrator
    assert auto.count("_save_to_db(vi, result, yaml_path)") == 2
    assert orchestrator.index(
        "_save_to_db(vi, result, profile_path)"
    ) < orchestrator.index("excel_path = export(vi, result, output_dir)")
    for save_at in [
        pos
        for pos in range(len(auto))
        if auto.startswith("_save_to_db(vi, result, yaml_path)", pos)
    ]:
        assert save_at < auto.index("path = export(vi, result, output_dir)", save_at)


def test_migration_declares_append_policy():
    sql = Path("db/migrations.sql").read_text(encoding="utf-8")
    assert "DROP INDEX IF EXISTS uq_valuations_company_date" in sql
    assert "idx_valuations_ticker_market_date" in sql
